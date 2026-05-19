from __future__ import annotations

import base64
import binascii
import copy
import hashlib
import mimetypes
import os
import re
import subprocess
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from html import escape as html_escape
from pathlib import Path

from flask import current_app

from ..extensions import db
from ..models import (
    Asset,
    ChatSession,
    ChatMessage,
    Character,
    CharacterOutfit,
    CharacterMemoryNote,
    CinemaNovel,
    CinemaNovelChapter,
    CinemaNovelCharacterImpression,
    CinemaNovelLoreEntry,
    CinemaNovelProgress,
    CinemaNovelReview,
    FeedPost,
    Project,
    SessionImage,
    World,
    WorldNewsItem,
)
from ..utils import json_util
from ..clients.image_ai_client import ImageAIClient
from ..clients.text_ai_client import TextAIClient
from .asset_service import AssetService
from .user_setting_service import UserSettingService


class CinemaNovelService:
    VALID_STATUSES = {"draft", "building", "published", "archived"}
    DEFAULT_REFERENCE_SOURCES = ("worldbuilding", "characters", "world")
    VALID_REFERENCE_SOURCES = {
        "worldbuilding",
        "characters",
        "world",
        "feed",
        "news",
        "short_stories",
    }

    def list_novels(self, project_id: int, *, include_unpublished: bool = False, mobile_only: bool = False):
        query = CinemaNovel.query.filter(
            CinemaNovel.project_id == project_id,
            CinemaNovel.deleted_at.is_(None),
        )
        if not include_unpublished:
            query = query.filter(CinemaNovel.status == "published")
        if mobile_only:
            query = query.filter(CinemaNovel.mobile_visible.is_(True))
        return query.order_by(CinemaNovel.sort_order.asc(), CinemaNovel.updated_at.desc(), CinemaNovel.id.desc()).all()

    def list_manual_short_videos(self, project_id: int, *, include_unpublished: bool = False):
        query = CinemaNovel.query.filter(
            CinemaNovel.project_id == project_id,
            CinemaNovel.deleted_at.is_(None),
            CinemaNovel.mode == "short_comic_video",
        )
        if not include_unpublished:
            query = query.filter(CinemaNovel.status == "published")
        novels = query.order_by(CinemaNovel.updated_at.desc(), CinemaNovel.id.desc()).all()
        return [
            novel
            for novel in novels
            if self._load_json(novel.production_json, default={}).get("source_type") == "manual_short_video"
        ]

    def get_novel(self, novel_id: int):
        return CinemaNovel.query.filter(CinemaNovel.id == novel_id, CinemaNovel.deleted_at.is_(None)).first()

    def update_novel_publication(self, novel_id: int, payload: dict | None):
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        payload = dict(payload or {})
        next_status = str(payload.get("status") or "").strip() or "draft"
        if next_status not in {"draft", "published"}:
            raise ValueError("status must be draft or published")
        novel.status = next_status
        novel.mobile_visible = bool(payload.get("mobile_visible", getattr(novel, "mobile_visible", True)))
        db.session.add(novel)
        db.session.commit()
        return novel

    def update_novel_metadata(self, novel_id: int, payload: dict | None):
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        payload = dict(payload or {})
        if "title" in payload:
            title = str(payload.get("title") or "").strip()
            if not title:
                raise ValueError("title is required")
            novel.title = title[:255]
        if "subtitle" in payload:
            subtitle = str(payload.get("subtitle") or "").strip()
            novel.subtitle = subtitle[:255] or None
        if "description" in payload:
            description = str(payload.get("description") or "").strip()
            novel.description = description or None
        db.session.add(novel)
        db.session.commit()
        return novel

    def list_bgm_assets(self, project_id: int):
        assets = self._asset_service.list_assets(project_id, asset_type="cinema_novel_bgm")
        return [self._serialize_asset(asset.id) for asset in assets]

    def upload_bgm_asset(self, project_id: int, user_id: int, upload_file):
        if not upload_file:
            raise ValueError("file is required")
        original_file_name = str(getattr(upload_file, "filename", "") or "").strip()
        asset = self._asset_service.create_asset(
            project_id,
            {
                "asset_type": "cinema_novel_bgm",
                "upload_file": upload_file,
                "metadata_json": json_util.dumps(
                    {
                        "source": "cinema_novel_bgm_upload",
                        "uploaded_by_user_id": user_id,
                        "original_file_name": original_file_name,
                    }
                ),
            },
        )
        mime_type = str(asset.mime_type or "").split(";", 1)[0].strip().lower()
        if not (mime_type.startswith("audio/") or mime_type == "application/ogg"):
            self._asset_service.delete_asset(asset.id)
            raise ValueError("audio file is required")
        return self._serialize_asset(asset.id)

    def export_powerpoint(self, novel_id: int):
        novel = self.get_novel(novel_id)
        if not novel:
            return None

        try:
            from pptx import Presentation
            from pptx.dml.color import RGBColor
            from pptx.enum.shapes import MSO_SHAPE
            from pptx.enum.text import PP_ALIGN
            from pptx.util import Inches, Pt
        except ImportError as exc:
            raise RuntimeError("python-pptx is required to export PowerPoint files") from exc

        from PIL import Image, ImageEnhance, ImageFilter, ImageOps

        chapters = self.list_chapters(novel.id)
        storage_root = Path(current_app.config["STORAGE_ROOT"]).resolve()
        export_dir = storage_root / "projects" / str(novel.project_id) / "exports" / "cinema_novels"
        work_dir = export_dir / "_work" / f"novel_{novel.id}"
        export_dir.mkdir(parents=True, exist_ok=True)
        work_dir.mkdir(parents=True, exist_ok=True)

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6]

        def add_background(slide, image_path: str | None, key: str):
            canvas_path = self._prepare_powerpoint_canvas(
                image_path,
                work_dir / f"{key}.jpg",
                Image,
                ImageOps,
                ImageFilter,
                ImageEnhance,
            )
            if canvas_path:
                slide.shapes.add_picture(str(canvas_path), 0, 0, width=prs.slide_width, height=prs.slide_height)
                return
            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
            bg.fill.solid()
            bg.fill.fore_color.rgb = RGBColor(18, 22, 30)
            bg.line.fill.background()

        def add_bottom_dialog(slide, text: str, speaker: str = "", *, footer: str = ""):
            overlay = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.55),
                Inches(5.28),
                Inches(12.25),
                Inches(1.72),
            )
            overlay.fill.solid()
            overlay.fill.fore_color.rgb = RGBColor(8, 10, 16)
            overlay.fill.transparency = 45
            overlay.line.fill.background()

            if speaker:
                name_box = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(0.8),
                    Inches(5.02),
                    Inches(2.45),
                    Inches(0.44),
                )
                name_box.fill.solid()
                name_box.fill.fore_color.rgb = RGBColor(38, 56, 92)
                name_box.fill.transparency = 20
                name_box.line.fill.background()
                name_frame = name_box.text_frame
                name_frame.clear()
                name_frame.margin_left = Inches(0.12)
                name_frame.margin_right = Inches(0.12)
                paragraph = name_frame.paragraphs[0]
                paragraph.alignment = PP_ALIGN.CENTER
                run = paragraph.add_run()
                run.text = speaker[:18]
                run.font.name = "Yu Gothic"
                run.font.size = Pt(13)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

            text_box = slide.shapes.add_textbox(Inches(0.86), Inches(5.52), Inches(11.72), Inches(1.08))
            frame = text_box.text_frame
            frame.clear()
            frame.word_wrap = True
            frame.margin_left = Inches(0.04)
            frame.margin_right = Inches(0.04)
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT
            run = paragraph.add_run()
            run.text = self._powerpoint_scene_text(text)
            run.font.name = "Yu Gothic"
            run.font.size = Pt(self._powerpoint_scene_font_size(text))
            run.font.color.rgb = RGBColor(255, 255, 255)

            if footer:
                footer_box = slide.shapes.add_textbox(Inches(10.85), Inches(6.62), Inches(1.75), Inches(0.24))
                footer_frame = footer_box.text_frame
                footer_frame.clear()
                footer_paragraph = footer_frame.paragraphs[0]
                footer_paragraph.alignment = PP_ALIGN.RIGHT
                footer_run = footer_paragraph.add_run()
                footer_run.text = footer
                footer_run.font.name = "Yu Gothic"
                footer_run.font.size = Pt(8)
                footer_run.font.color.rgb = RGBColor(188, 198, 215)

        title_image = self._asset_file_path(novel.poster_asset_id or novel.cover_asset_id)
        title_slide = prs.slides.add_slide(blank_layout)
        add_background(title_slide, title_image, "title")
        add_bottom_dialog(
            title_slide,
            "\n".join([part for part in [novel.title or "Untitled", novel.subtitle or novel.description or ""] if part]),
            "",
            footer="title",
        )

        for chapter in chapters:
            scenes = self._load_json(chapter.scene_json, default=[])
            if not isinstance(scenes, list):
                scenes = []
            chapter_fallback = self._asset_file_path(chapter.cover_asset_id) or title_image
            for scene_index, scene in enumerate(scenes):
                if not isinstance(scene, dict):
                    continue
                scene_text = str(scene.get("text") or "").strip()
                if not scene_text:
                    continue
                slide = prs.slides.add_slide(blank_layout)
                image_path = (
                    self._asset_file_path(scene.get("still_asset_id"))
                    or self._asset_file_path(scene.get("background_asset_id"))
                    or chapter_fallback
                )
                add_background(slide, image_path, f"chapter_{chapter.id}_scene_{scene_index + 1}")
                speaker = str(scene.get("speaker") or "").strip()
                footer = f"{int(chapter.chapter_no or 0):02d}-{scene_index + 1:02d}"
                add_bottom_dialog(slide, scene_text, speaker, footer=footer)

        filename_base = self._safe_powerpoint_filename(novel.title or f"cinema_novel_{novel.id}")
        filename = f"{filename_base}_{novel.id}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.pptx"
        output_path = export_dir / filename
        prs.save(output_path)
        self._set_powerpoint_shape_alpha(output_path, "080A10", 20000)
        return str(output_path), filename

    def export_epub(self, novel_id: int, *, writing_mode: str = "horizontal"):
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        chapters = self.list_chapters(novel.id)
        storage_root = Path(current_app.config["STORAGE_ROOT"]).resolve()
        export_dir = storage_root / "projects" / str(novel.project_id) / "exports" / "cinema_novels"
        export_dir.mkdir(parents=True, exist_ok=True)

        filename_base = self._safe_powerpoint_filename(novel.title or f"cinema_novel_{novel.id}")
        is_short_comic = getattr(novel, "mode", "") == "short_comic_video"
        vertical = writing_mode == "vertical" and not is_short_comic
        mode_suffix = "_vertical" if vertical else ""
        filename = f"{filename_base}_{novel.id}{mode_suffix}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.epub"
        output_path = export_dir / filename

        image_items = {}
        manifest_images = []

        def add_image(asset_id, *, name_hint="image"):
            try:
                asset_id = int(asset_id or 0)
            except (TypeError, ValueError):
                asset_id = 0
            if not asset_id:
                return None
            if asset_id in image_items:
                return image_items[asset_id]["href"]
            asset = self._asset_service.get_asset(asset_id)
            if not asset or not asset.file_path:
                return None
            source_path = Path(asset.file_path)
            if not source_path.exists():
                return None
            mime_type = str(asset.mime_type or mimetypes.guess_type(str(source_path))[0] or "image/jpeg").split(";", 1)[0]
            if not mime_type.startswith("image/"):
                return None
            extension = source_path.suffix.lower() or mimetypes.guess_extension(mime_type) or ".jpg"
            href = f"images/{name_hint}_{asset_id}{extension}"
            item_id = f"img_{asset_id}"
            image_items[asset_id] = {"id": item_id, "href": href, "path": source_path, "media_type": mime_type}
            manifest_images.append(image_items[asset_id])
            return href

        cover_href = add_image(novel.poster_asset_id or novel.cover_asset_id, name_hint="cover")
        chapter_docs = []
        if is_short_comic:
            chapter_docs.extend(self._epub_short_comic_documents(novel, chapters, add_image))
        else:
            chapter_docs.extend(self._epub_novel_documents(novel, chapters, add_image))

        if not chapter_docs:
            chapter_docs.append(
                {
                    "id": "chapter_1",
                    "href": "chapters/chapter_1.xhtml",
                    "title": novel.title or "Untitled",
                    "body": f"<h1>{html_escape(novel.title or 'Untitled')}</h1><p>{html_escape(novel.description or novel.subtitle or '')}</p>",
                }
            )

        nav_doc = self._epub_nav_document(novel, chapter_docs)
        ncx_doc = self._epub_ncx_document(novel, chapter_docs)
        opf_doc = self._epub_package_document(novel, chapter_docs, manifest_images, cover_href=cover_href)
        css_doc = self._epub_stylesheet(vertical=vertical)

        with zipfile.ZipFile(output_path, "w") as epub:
            epub.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
            epub.writestr(
                "META-INF/container.xml",
                """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
  <rootfiles>
    <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
  </rootfiles>
</container>""",
                compress_type=zipfile.ZIP_DEFLATED,
            )
            epub.writestr("OEBPS/styles/book.css", css_doc, compress_type=zipfile.ZIP_DEFLATED)
            epub.writestr("OEBPS/nav.xhtml", nav_doc, compress_type=zipfile.ZIP_DEFLATED)
            epub.writestr("OEBPS/toc.ncx", ncx_doc, compress_type=zipfile.ZIP_DEFLATED)
            epub.writestr("OEBPS/content.opf", opf_doc, compress_type=zipfile.ZIP_DEFLATED)
            for doc in chapter_docs:
                epub.writestr(
                    f"OEBPS/{doc['href']}",
                    self._epub_xhtml_document(doc["title"], doc["body"], body_class=doc.get("body_class", "")),
                    compress_type=zipfile.ZIP_DEFLATED,
                )
            for item in manifest_images:
                epub.write(item["path"], f"OEBPS/{item['href']}", compress_type=zipfile.ZIP_DEFLATED)
        return str(output_path), filename

    def _epub_novel_documents(self, novel: CinemaNovel, chapters: list[CinemaNovelChapter], add_image):
        docs = []
        for index, chapter in enumerate(chapters, start=1):
            title = chapter.title or f"Chapter {index}"
            body_parts = [f"<h1>{html_escape(title)}</h1>"]

            markdown_html = self._epub_markdown_to_html(chapter.body_markdown or "", skip_heading=title)
            if markdown_html:
                body_parts.append(markdown_html)

            scenes = self._load_json(chapter.scene_json, default=[])
            if not isinstance(scenes, list):
                scenes = []
            scene_docs = []
            used_scene_images = set()
            for scene_index, scene in enumerate(scenes, start=1):
                if not isinstance(scene, dict):
                    continue
                text = str(scene.get("text") or "").strip()
                speaker = str(scene.get("speaker") or "").strip()
                image_href = add_image(scene.get("still_asset_id") or scene.get("background_asset_id"), name_hint=f"chapter_{index}_scene_{scene_index}")
                if image_href and image_href not in used_scene_images:
                    used_scene_images.add(image_href)
                    image_title = f"{title} image {len(scene_docs) + 1}"
                    scene_docs.append(
                        {
                            "id": f"chapter_{index}_image_{len(scene_docs) + 1}",
                            "href": f"chapters/chapter_{index}_image_{len(scene_docs) + 1}.xhtml",
                            "title": image_title,
                            "body_class": "image-page-document",
                            "body": (
                                '<section class="image-page">'
                                f'<figure class="scene-image-full"><img src="../{html_escape(image_href, quote=True)}" alt="{html_escape(image_title, quote=True)}"/>'
                                "</figure>"
                                "</section>"
                            ),
                        }
                    )
                elif text and not markdown_html:
                    body_parts.append(self._epub_scene_paragraph(text, speaker))

            docs.append(
                {
                    "id": f"chapter_{index}",
                    "href": f"chapters/chapter_{index}.xhtml",
                    "title": title,
                    "body_class": "text-page",
                    "body": "\n".join(part for part in body_parts if part),
                }
            )
            docs.extend(scene_docs)
        if not docs:
            docs.append(
                {
                    "id": "chapter_1",
                    "href": "chapters/chapter_1.xhtml",
                    "title": novel.title or "Untitled",
                    "body_class": "text-page",
                    "body": self._epub_markdown_to_html(novel.description or novel.subtitle or "", skip_heading=novel.title or "") or f"<h1>{html_escape(novel.title or 'Untitled')}</h1>",
                }
            )
        return docs

    def _epub_short_comic_documents(self, novel: CinemaNovel, chapters: list[CinemaNovelChapter], add_image):
        docs = []
        panel_no = 1
        cover_href = add_image(novel.poster_asset_id or novel.cover_asset_id, name_hint="cover")
        if cover_href:
            docs.append(
                {
                    "id": "cover",
                    "href": "chapters/cover.xhtml",
                    "title": novel.title or "Cover",
                    "body": "\n".join(
                        [
                            f"<h1>{html_escape(novel.title or 'Untitled')}</h1>",
                            f'<figure class="panel"><img src="../{html_escape(cover_href, quote=True)}" alt="{html_escape(novel.title or "Cover", quote=True)}"/></figure>',
                            f"<p>{html_escape(novel.description or novel.subtitle or '')}</p>" if (novel.description or novel.subtitle) else "",
                        ]
                    ),
                }
            )
        for chapter in chapters:
            scenes = self._load_json(chapter.scene_json, default=[])
            if not isinstance(scenes, list):
                scenes = []
            for scene in scenes:
                if not isinstance(scene, dict):
                    continue
                text = str(scene.get("text") or scene.get("title") or "").strip()
                title = str(scene.get("title") or text or f"Panel {panel_no}").strip()
                speaker = str(scene.get("speaker") or "").strip()
                image_href = add_image(scene.get("still_asset_id") or scene.get("background_asset_id"), name_hint=f"panel_{panel_no}")
                if not image_href and not text:
                    continue
                body_parts = [f"<h1>{html_escape(title[:80])}</h1>"]
                if image_href:
                    body_parts.append(
                        f'<figure class="panel"><img src="../{html_escape(image_href, quote=True)}" alt="{html_escape(title, quote=True)}"/></figure>'
                    )
                if text:
                    body_parts.append(self._epub_scene_paragraph(text, speaker))
                docs.append(
                    {
                        "id": f"panel_{panel_no}",
                        "href": f"chapters/panel_{panel_no}.xhtml",
                        "title": title[:80],
                        "body": "\n".join(body_parts),
                    }
                )
                panel_no += 1
        return docs

    def _epub_markdown_to_html(self, markdown: str, *, skip_heading: str = "") -> str:
        lines = str(markdown or "").replace("\r\n", "\n").split("\n")
        blocks = []
        paragraph = []
        skip_key = self._epub_heading_key(skip_heading)
        skipped_first_heading = False

        def flush_paragraph():
            if not paragraph:
                return
            blocks.append("<p>" + "<br/>".join(html_escape(line) for line in paragraph) + "</p>")
            paragraph.clear()

        for raw_line in lines:
            line = raw_line.strip()
            if not line:
                flush_paragraph()
                continue
            heading = re.match(r"^(#{1,3})\s+(.+)$", line)
            if heading:
                flush_paragraph()
                level = min(3, len(heading.group(1)))
                heading_text = heading.group(2).strip()
                if not skipped_first_heading and skip_key and self._epub_heading_key(heading_text) == skip_key:
                    skipped_first_heading = True
                    continue
                blocks.append(f"<h{level}>{html_escape(heading_text)}</h{level}>")
                continue
            if not skipped_first_heading and skip_key and not blocks and not paragraph and self._epub_heading_key(line) == skip_key:
                skipped_first_heading = True
                continue
            bullet = re.match(r"^[-*]\s+(.+)$", line)
            if bullet:
                flush_paragraph()
                blocks.append(f"<p class=\"bullet\">・{html_escape(bullet.group(1).strip())}</p>")
                continue
            paragraph.append(line)
        flush_paragraph()
        return "\n".join(blocks)

    def _epub_heading_key(self, value: str) -> str:
        return re.sub(r"\s+", "", str(value or "").strip().lower())

    def _epub_scene_caption(self, text: str, speaker: str = "") -> str:
        text = str(text or "").strip()
        speaker = str(speaker or "").strip()
        if not text:
            return speaker
        caption = f"{speaker}: {text}" if speaker else text
        return caption[:140]

    def _epub_scene_paragraph(self, text: str, speaker: str = "") -> str:
        text = str(text or "").strip()
        speaker = str(speaker or "").strip()
        if speaker:
            return f'<p><strong>{html_escape(speaker)}</strong><br/>{html_escape(text)}</p>'
        return f"<p>{html_escape(text)}</p>"

    def _epub_xhtml_document(self, title: str, body: str, *, body_class: str = "") -> str:
        class_attr = f' class="{html_escape(body_class, quote=True)}"' if body_class else ""
        return f"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" lang="ja" xml:lang="ja">
<head>
  <meta charset="UTF-8"/>
  <title>{html_escape(title or "Untitled")}</title>
  <link rel="stylesheet" type="text/css" href="../styles/book.css"/>
</head>
<body{class_attr}>
{body}
</body>
</html>"""

    def _epub_nav_document(self, novel: CinemaNovel, chapter_docs: list[dict]) -> str:
        items = "\n".join(
            f'      <li><a href="{html_escape(doc["href"], quote=True)}">{html_escape(doc["title"] or "Untitled")}</a></li>'
            for doc in chapter_docs
        )
        return f"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE html>
<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops" lang="ja" xml:lang="ja">
<head>
  <meta charset="UTF-8"/>
  <title>{html_escape(novel.title or "Untitled")} 目次</title>
  <link rel="stylesheet" type="text/css" href="styles/book.css"/>
</head>
<body>
  <nav epub:type="toc" id="toc">
    <h1>目次</h1>
    <ol>
{items}
    </ol>
  </nav>
</body>
</html>"""

    def _epub_ncx_document(self, novel: CinemaNovel, chapter_docs: list[dict]) -> str:
        nav_points = []
        for index, doc in enumerate(chapter_docs, start=1):
            nav_points.append(
                f"""  <navPoint id="navPoint-{index}" playOrder="{index}">
    <navLabel><text>{html_escape(doc["title"] or "Untitled")}</text></navLabel>
    <content src="{html_escape(doc["href"], quote=True)}"/>
  </navPoint>"""
            )
        return f"""<?xml version="1.0" encoding="UTF-8"?>
<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1">
  <head>
    <meta name="dtb:uid" content="novelcreator-cinema-novel-{int(novel.id)}"/>
    <meta name="dtb:depth" content="1"/>
    <meta name="dtb:totalPageCount" content="0"/>
    <meta name="dtb:maxPageNumber" content="0"/>
  </head>
  <docTitle><text>{html_escape(novel.title or "Untitled")}</text></docTitle>
  <navMap>
{chr(10).join(nav_points)}
  </navMap>
</ncx>"""

    def _epub_package_document(self, novel: CinemaNovel, chapter_docs: list[dict], image_items: list[dict], *, cover_href: str | None = None) -> str:
        modified = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")
        manifest_parts = [
            '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>',
            '<item id="ncx" href="toc.ncx" media-type="application/x-dtbncx+xml"/>',
            '<item id="css" href="styles/book.css" media-type="text/css"/>',
        ]
        for doc in chapter_docs:
            manifest_parts.append(
                f'<item id="{html_escape(doc["id"], quote=True)}" href="{html_escape(doc["href"], quote=True)}" media-type="application/xhtml+xml"/>'
            )
        for item in image_items:
            properties = ' properties="cover-image"' if cover_href and item["href"] == cover_href else ""
            manifest_parts.append(
                f'<item id="{html_escape(item["id"], quote=True)}" href="{html_escape(item["href"], quote=True)}" media-type="{html_escape(item["media_type"], quote=True)}"{properties}/>'
            )
        spine_items = "\n".join(f'    <itemref idref="{html_escape(doc["id"], quote=True)}"/>' for doc in chapter_docs)
        cover_meta = ""
        if cover_href:
            cover_item = next((item for item in image_items if item["href"] == cover_href), None)
            if cover_item:
                cover_meta = f'\n    <meta name="cover" content="{html_escape(cover_item["id"], quote=True)}"/>'
        return f"""<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book-id">
  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">
    <dc:identifier id="book-id">novelcreator-cinema-novel-{int(novel.id)}</dc:identifier>
    <dc:title>{html_escape(novel.title or "Untitled")}</dc:title>
    <dc:language>ja</dc:language>
    <meta property="dcterms:modified">{modified}</meta>{cover_meta}
  </metadata>
  <manifest>
    {chr(10).join(manifest_parts)}
  </manifest>
  <spine toc="ncx">
{spine_items}
  </spine>
</package>"""

    def _epub_stylesheet(self, *, vertical: bool = False) -> str:
        vertical_css = """
.text-page {
  -epub-writing-mode: vertical-rl;
  writing-mode: vertical-rl;
  text-orientation: mixed;
  max-height: 42em;
}
.text-page p {
  text-align: justify;
}
.text-page figure {
  max-height: 90%;
}
.text-page figcaption {
  text-align: start;
}
""" if vertical else ""
        return """body {
  color: #1f2937;
  font-family: serif;
  line-height: 1.8;
  margin: 0;
  padding: 1.25em;
}
h1, h2, h3 {
  color: #111827;
  font-family: sans-serif;
  line-height: 1.35;
}
p {
  margin: 0 0 1em;
}
img {
  display: block;
  height: auto;
  max-width: 100%;
}
figure {
  margin: 1.25em 0;
  page-break-inside: avoid;
}
.image-page {
  page-break-before: always;
  break-before: page;
  min-height: 92vh;
}
.scene-image-full {
  align-items: center;
  display: flex;
  flex-direction: column;
  justify-content: center;
  min-height: 92vh;
  text-align: center;
  -epub-writing-mode: horizontal-tb;
  writing-mode: horizontal-tb;
}
.scene-image-full img {
  max-height: 88vh;
  max-width: 100%;
  object-fit: contain;
}
.image-page figcaption {
  display: none;
}
figcaption {
  color: #4b5563;
  font-size: 0.9em;
  margin-top: 0.6em;
}
.panel img, .scene-image img, .chapter-image img {
  border-radius: 0.25em;
}
.bullet {
  margin-left: 1em;
}""" + vertical_css

    def export_short_video(self, novel_id: int, payload: dict | None = None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None

        from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont, ImageOps

        chapters = self.list_chapters(novel.id)
        storage_root = Path(current_app.config["STORAGE_ROOT"]).resolve()
        export_dir = storage_root / "projects" / str(novel.project_id) / "exports" / "cinema_novels"
        work_dir = export_dir / "_video_work" / f"novel_{novel.id}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}"
        frames_dir = work_dir / "frames"
        export_dir.mkdir(parents=True, exist_ok=True)
        frames_dir.mkdir(parents=True, exist_ok=True)

        font_regular = self._load_video_font(ImageFont, size=42)
        font_small = self._load_video_font(ImageFont, size=28)
        font_title = self._load_video_font(ImageFont, size=54)
        font_footer = self._load_video_font(ImageFont, size=22)
        title_image = self._asset_file_path(novel.poster_asset_id or novel.cover_asset_id)
        orientation = str(payload.get("orientation") or "portrait").strip().lower()
        landscape = orientation in {"landscape", "horizontal", "wide", "pc"}
        frame_paths = []
        durations = []

        def add_frame(image_path, text, *, speaker="", footer="", title=False, duration=5.0):
            index = len(frame_paths) + 1
            frame_path = frames_dir / f"frame_{index:04d}.png"
            self._render_short_video_frame(
                image_path,
                frame_path,
                text,
                speaker=speaker,
                footer=footer,
                title=title,
                Image=Image,
                ImageDraw=ImageDraw,
                ImageEnhance=ImageEnhance,
                ImageFilter=ImageFilter,
                ImageOps=ImageOps,
                font_regular=font_regular,
                font_small=font_small,
                font_title=font_title,
                font_footer=font_footer,
                landscape=landscape,
            )
            frame_paths.append(frame_path)
            durations.append(duration)

        add_frame(
            title_image,
            "\n".join([part for part in [novel.title or "Untitled", novel.subtitle or novel.description or ""] if part]),
            footer="title",
            title=True,
            duration=5.2,
        )
        previous_scene_image_path = None

        for chapter in chapters:
            scenes = self._load_json(chapter.scene_json, default=[])
            if not isinstance(scenes, list):
                scenes = []
            chapter_fallback = self._asset_file_path(chapter.cover_asset_id)
            for scene_index, scene in enumerate(scenes):
                if not isinstance(scene, dict):
                    continue
                scene_text = self._short_video_scene_text(str(scene.get("text") or ""))
                if not scene_text:
                    continue
                direct_image_path = (
                    self._asset_file_path(scene.get("still_asset_id"))
                    or self._asset_file_path(scene.get("background_asset_id"))
                )
                image_path = direct_image_path or previous_scene_image_path or chapter_fallback
                if direct_image_path:
                    previous_scene_image_path = direct_image_path
                duration = 5.0 if len(scene_text) < 50 else 5.8
                add_frame(
                    image_path,
                    scene_text,
                    speaker=str(scene.get("speaker") or "").strip(),
                    footer=f"{int(chapter.chapter_no or 0):02d}-{scene_index + 1:02d}",
                    duration=duration,
                )

        if not frame_paths:
            raise ValueError("no scenes were found for video export")

        filename_base = self._safe_powerpoint_filename(novel.title or f"cinema_novel_{novel.id}")
        suffix = "landscape" if landscape else "portrait"
        filename = f"{filename_base}_{novel.id}_{suffix}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.mp4"
        output_path = export_dir / filename
        concat_path = work_dir / "concat.txt"
        with concat_path.open("w", encoding="utf-8") as handle:
            for frame_path, duration in zip(frame_paths, durations):
                handle.write(f"file '{frame_path.as_posix()}'\n")
                handle.write(f"duration {duration:.2f}\n")
            handle.write(f"file '{frame_paths[-1].as_posix()}'\n")

        command = [
            "ffmpeg",
            "-y",
            "-f",
            "concat",
            "-safe",
            "0",
            "-i",
            str(concat_path),
            "-vf",
            "fps=30,format=yuv420p",
            "-c:v",
            "libx264",
            "-preset",
            "medium",
            "-crf",
            "22",
            "-movflags",
            "+faststart",
            str(output_path),
        ]
        result = subprocess.run(command, capture_output=True, text=True, cwd=str(Path(current_app.root_path).parent))
        if result.returncode != 0:
            raise RuntimeError((result.stderr or result.stdout or "ffmpeg failed").strip()[-1000:])
        self._apply_bgm_to_video(output_path, novel.project_id, payload)
        return str(output_path), filename

    def export_comic_short_video(self, novel_id: int, payload: dict | None = None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None

        from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont, ImageOps

        production = self._load_json(novel.production_json, default={})
        manual_short_video = production.get("source_type") == "manual_short_video"

        if not manual_short_video and not (novel.poster_asset_id or novel.cover_asset_id):
            try:
                self.generate_short_comic_thumbnail(novel.id, {})
                novel = self.get_novel(novel_id)
            except Exception:
                current_app.logger.exception("failed to generate short comic thumbnail before video export")
        production = self._load_json(novel.production_json, default={})
        if not manual_short_video and not production.get("end_card_asset_id"):
            try:
                self.generate_short_comic_end_card(novel.id, {})
                novel = self.get_novel(novel_id)
                production = self._load_json(novel.production_json, default={})
            except Exception:
                current_app.logger.exception("failed to generate short comic end card before video export")

        chapters = self.list_chapters(novel.id)
        storage_root = Path(current_app.config["STORAGE_ROOT"]).resolve()
        export_dir = storage_root / "projects" / str(novel.project_id) / "exports" / "cinema_novels"
        work_dir = export_dir / "_comic_video_work" / f"novel_{novel.id}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}"
        frames_dir = work_dir / "frames"
        export_dir.mkdir(parents=True, exist_ok=True)
        frames_dir.mkdir(parents=True, exist_ok=True)

        fonts = {
            "headline": self._load_video_font(ImageFont, size=76),
            "headline_small": self._load_video_font(ImageFont, size=58),
            "label": self._load_video_font(ImageFont, size=32),
            "footer": self._load_video_font(ImageFont, size=24),
        }
        first_scene_image_path = None
        last_scene_image_path = None
        if manual_short_video:
            for chapter in chapters:
                scenes = self._load_json(chapter.scene_json, default=[])
                if not isinstance(scenes, list):
                    continue
                chapter_fallback = self._asset_file_path(chapter.cover_asset_id)
                for scene in scenes:
                    if not isinstance(scene, dict):
                        continue
                    image_path = (
                        self._asset_file_path(scene.get("still_asset_id"))
                        or self._asset_file_path(scene.get("background_asset_id"))
                        or chapter_fallback
                    )
                    if image_path:
                        first_scene_image_path = first_scene_image_path or image_path
                        last_scene_image_path = image_path

        title_image = first_scene_image_path if manual_short_video else self._asset_file_path(novel.poster_asset_id or novel.cover_asset_id)
        frame_paths = []
        durations = []

        def add_panel(image_path, caption, *, label="", footer="", title=False, fade_in=False, duration=2.55):
            frame_path = frames_dir / f"frame_{len(frame_paths) + 1:04d}.png"
            self._render_comic_short_frame(
                image_path,
                frame_path,
                caption,
                label=label,
                footer=footer,
                title=title,
                Image=Image,
                ImageDraw=ImageDraw,
                ImageEnhance=ImageEnhance,
                ImageFilter=ImageFilter,
                ImageOps=ImageOps,
                fonts=fonts,
                mobile_visible=bool(getattr(novel, "mobile_visible", True)),
            )
            if fade_in:
                fade_paths = self._short_video_frame_fade_in_paths(frame_path, Image, frame_count=9)
                for fade_path in fade_paths:
                    frame_paths.append(fade_path)
                    durations.append(1 / 30)
                frame_paths.append(frame_path)
                durations.append(max(0.5, duration - (len(fade_paths) / 30)))
            else:
                frame_paths.append(frame_path)
                durations.append(duration)

        if title_image:
            add_panel(
                title_image,
                novel.title or "Untitled",
                label=novel.subtitle or novel.description or "short comic",
                footer="title",
                title=True,
                fade_in=False,
                duration=1.0 if manual_short_video else 2.8,
            )
        previous_image_path = None
        scene_entries = []

        for chapter in chapters:
            scenes = self._load_json(chapter.scene_json, default=[])
            if not isinstance(scenes, list):
                scenes = []
            chapter_fallback = self._asset_file_path(chapter.cover_asset_id)
            for scene_index, scene in enumerate(scenes):
                if not isinstance(scene, dict):
                    continue
                scene_text = str(scene.get("text") or "").strip()
                caption = "" if scene.get("comic_page") else self._comic_panel_caption(scene_text)
                if not caption and not scene.get("comic_page"):
                    continue
                image_path = (
                    self._asset_file_path(scene.get("still_asset_id"))
                    or self._asset_file_path(scene.get("background_asset_id"))
                    or previous_image_path
                    or chapter_fallback
                )
                if not image_path:
                    continue
                previous_image_path = image_path
                scene_entries.append(
                    {
                        "image_path": image_path,
                        "caption": caption,
                        "label": "" if scene.get("comic_page") else str(scene.get("speaker") or f"SCENE {int(chapter.chapter_no or 0):02d}-{scene_index + 1:02d}").strip(),
                        "footer": f"{int(chapter.chapter_no or 0):02d}-{scene_index + 1:02d}",
                    }
                )

        for entry_index, entry in enumerate(scene_entries):
            add_panel(
                entry["image_path"],
                entry["caption"],
                label=entry["label"],
                footer=entry["footer"],
                duration=1.0 if manual_short_video and entry_index == len(scene_entries) - 1 else 3.45,
            )

        end_card_image = None if manual_short_video else self._asset_file_path(production.get("end_card_asset_id"))
        if end_card_image and end_card_image != previous_image_path:
            add_panel(
                end_card_image,
                "END" if manual_short_video else "LAPLACE CITY",
                label="END",
                footer="end",
                title=True,
                fade_in=True,
                duration=1.8,
            )

        if not frame_paths:
            raise ValueError("no scenes were found for comic video export")

        filename_base = self._safe_powerpoint_filename(novel.title or f"cinema_novel_{novel.id}")
        filename = f"{filename_base}_comic.mp4"
        output_path = export_dir / filename
        concat_path = work_dir / "concat.txt"
        with concat_path.open("w", encoding="utf-8") as handle:
            for frame_path, duration in zip(frame_paths, durations):
                handle.write(f"file '{frame_path.as_posix()}'\n")
                handle.write(f"duration {duration:.2f}\n")
            handle.write(f"file '{frame_paths[-1].as_posix()}'\n")

        command = [
            "ffmpeg",
            "-y",
            "-f",
            "concat",
            "-safe",
            "0",
            "-i",
            str(concat_path),
            "-vf",
            "fps=30,format=yuv420p",
            "-c:v",
            "libx264",
            "-preset",
            "medium",
            "-crf",
            "21",
            "-movflags",
            "+faststart",
            str(output_path),
        ]
        result = subprocess.run(command, capture_output=True, text=True, cwd=str(Path(current_app.root_path).parent))
        if result.returncode != 0:
            raise RuntimeError((result.stderr or result.stdout or "ffmpeg failed").strip()[-1000:])
        self._apply_bgm_to_video(output_path, novel.project_id, payload)
        return str(output_path), filename

    def create_short_comic_novel(self, project_id: int, user_id: int, payload: dict | None):
        payload = dict(payload or {})
        comic_layout = self._short_comic_layout(payload)
        source_novel_id = self._source_novel_id_from_payload(project_id, payload)
        inherited_character_ids = self._source_novel_character_ids_from_payload(project_id, payload)
        inherited_character_names = self._character_names_for_ids(project_id, inherited_character_ids)
        storyboard = (
            self.generate_comic_page_storyboard(project_id, payload)
            if comic_layout == "page"
            else self.generate_short_comic_storyboard(project_id, payload)
        )
        title = str(storyboard.get("title") or payload.get("title") or "ショート漫画").strip() or "ショート漫画"
        subtitle = str(storyboard.get("logline") or ("漫画ページ" if comic_layout == "page" else "縦ショート漫画")).strip()
        panels = storyboard.get("pages") if comic_layout == "page" and isinstance(storyboard.get("pages"), list) else storyboard.get("panels")
        panels = panels if isinstance(panels, list) else []
        if not panels:
            raise ValueError("short comic storyboard did not include panels")

        novel = CinemaNovel(
            project_id=project_id,
            created_by_user_id=user_id,
            title=title,
            subtitle=subtitle[:255],
            description=subtitle,
            status=str(payload.get("status") or "draft").strip() or "draft",
            mobile_visible=self._normalize_bool(payload.get("mobile_visible", True)),
            mode="short_comic_video",
            production_json=json_util.dumps(
                {
                    "source_input": payload,
                    "storyboard": storyboard,
                    "comic_layout": comic_layout,
                    "target_panel_count": storyboard.get("target_panel_count"),
                    "actual_panel_count": len(panels),
                    "source_novel_id": source_novel_id,
                    "inherited_character_ids": inherited_character_ids,
                }
            ),
        )
        db.session.add(novel)
        db.session.flush()

        scenes = []
        for index, panel in enumerate(panels):
            if not isinstance(panel, dict):
                continue
            caption = (
                self._normalize_comic_page_caption(panel.get("page_title") or panel.get("caption") or panel.get("summary") or panel.get("text") or "")
                if comic_layout == "page"
                else self._normalize_short_comic_caption(panel.get("caption") or panel.get("text") or "")
            )
            if not caption:
                continue
            panel_character_ids = self._normalize_character_ids(panel.get("character_ids"))
            scene_character_ids = panel_character_ids or inherited_character_ids
            panel_characters = panel.get("characters") if isinstance(panel.get("characters"), list) else []
            characters = [str(item).strip() for item in panel_characters if str(item).strip()]
            for name in inherited_character_names:
                if name and name not in characters:
                    characters.append(name)
            page_panels = panel.get("panels") if isinstance(panel.get("panels"), list) else []
            dialogue_lines = []
            for item in page_panels:
                if not isinstance(item, dict):
                    continue
                speaker = str(item.get("speaker") or "").strip()
                speech = str(item.get("speech") or item.get("dialogue") or "").strip()
                if speech:
                    dialogue_lines.append(f"{speaker}: {speech}" if speaker else speech)
            scenes.append(
                {
                    "speaker": str(panel.get("speaker") or "").strip(),
                    "text": caption,
                    "caption": caption,
                    "comic_layout": comic_layout,
                    "comic_page": comic_layout == "page",
                    "page_title": str(panel.get("page_title") or caption).strip(),
                    "page_panels": page_panels,
                    "dialogue": dialogue_lines,
                    "tone": str(panel.get("tone") or "").strip(),
                    "emotion": str(panel.get("emotion") or "").strip(),
                    "shot": str(panel.get("shot") or "").strip(),
                    "viewer_context": str(panel.get("viewer_context") or "").strip(),
                    "visible_event": str(panel.get("visible_event") or "").strip(),
                    "joke": str(panel.get("joke") or "").strip(),
                    "story_function": str(panel.get("story_function") or "").strip(),
                    "visual_focus": str(panel.get("visual_focus") or "").strip(),
                    "image_prompt": str(panel.get("image_prompt") or "").strip(),
                    "characters": characters,
                    "character_ids": scene_character_ids,
                    "background_asset_id": None,
                    "still_asset_id": None,
                    "choice_list": [],
                    "panel_index": index,
                }
            )
        if not scenes:
            raise ValueError("short comic storyboard did not include usable panels")

        chapter = CinemaNovelChapter(
            novel_id=novel.id,
            chapter_no=1,
            title="漫画ページ" if comic_layout == "page" else "ショート漫画",
            body_markdown="\n".join(f"{index + 1}. {scene['caption']}" for index, scene in enumerate(scenes)),
            scene_json=json_util.dumps(scenes),
            sort_order=0,
        )
        db.session.add(chapter)
        db.session.commit()

        if payload.get("generate_images", True):
            self.generate_short_comic_thumbnail(novel.id, payload)
            panel_result = self.generate_short_comic_panel_images(novel.id, {"parallel": True, "max_workers": 5, "overwrite": False})
            created_count = len((panel_result or {}).get("assets") or [])
            failed_count = len((panel_result or {}).get("failed_assets") or [])
            production = self._load_json(novel.production_json, default={})
            production["panel_image_generation"] = {
                "expected": len(scenes),
                "created": created_count,
                "failed": failed_count,
                "finished_at": datetime.utcnow().isoformat(),
            }
            novel.production_json = json_util.dumps(production)
            db.session.add(novel)
            db.session.commit()
            if scenes and created_count == 0:
                raise RuntimeError(f"panel image generation created 0/{len(scenes)} images")
            self.generate_short_comic_end_card(novel.id, payload)
        return novel

    def create_manual_short_video(self, project_id: int, user_id: int, payload: dict | None):
        payload = dict(payload or {})
        title = str(payload.get("title") or "ショート動画").strip() or "ショート動画"
        raw_scenes = payload.get("scenes") if isinstance(payload.get("scenes"), list) else []
        if not raw_scenes:
            raw_scenes = [
                {
                    "caption": str(payload.get("caption") or "").strip(),
                    "image_prompt": str(payload.get("image_prompt") or "").strip(),
                    "character_ids": payload.get("character_ids") or [],
                }
            ]
        scenes = []
        for index, raw_scene in enumerate(raw_scenes):
            if not isinstance(raw_scene, dict):
                continue
            scene = self._manual_short_video_scene(project_id, raw_scene, index=index)
            if scene["caption"] or scene["image_prompt"]:
                scenes.append(scene)
        if not scenes:
            scenes.append(self._manual_short_video_scene(project_id, {"caption": "", "image_prompt": ""}, index=0))
        novel = CinemaNovel(
            project_id=project_id,
            created_by_user_id=user_id,
            title=title[:255],
            subtitle="手動作成ショート動画",
            description=str(payload.get("description") or "").strip() or "Feedのような短い入力から作る手動ショート動画",
            status=str(payload.get("status") or "draft").strip() or "draft",
            mobile_visible=self._normalize_bool(payload.get("mobile_visible", False)),
            mode="short_comic_video",
            production_json=json_util.dumps(
                {
                    "source_type": "manual_short_video",
                    "source_input": payload,
                    "comic_layout": "short",
                    "actual_panel_count": len(scenes),
                }
            ),
        )
        db.session.add(novel)
        db.session.flush()
        chapter = CinemaNovelChapter(
            novel_id=novel.id,
            chapter_no=1,
            title="ショート動画",
            body_markdown="\n".join(f"{index + 1}. {scene.get('caption') or scene.get('image_prompt')}" for index, scene in enumerate(scenes)),
            scene_json=json_util.dumps(scenes),
            sort_order=0,
        )
        db.session.add(chapter)
        db.session.commit()
        return novel

    def create_chat_session_novel(self, session_id: int, user_id: int, payload: dict | None = None):
        payload = dict(payload or {})
        chat_session = ChatSession.query.filter(
            ChatSession.id == session_id,
            ChatSession.deleted_at.is_(None),
        ).first()
        if not chat_session:
            return None
        messages = ChatMessage.query.filter_by(session_id=session_id).order_by(ChatMessage.order_no.asc(), ChatMessage.id.asc()).all()
        exportable_messages = [
            message for message in messages
            if str(message.message_text or "").strip()
            and str(message.sender_type or "").strip().lower() not in {"system", "debug"}
        ]
        if not exportable_messages:
            raise ValueError("ノベル化できるチャット本文がありません。")
        session_images = SessionImage.query.filter_by(session_id=session_id).order_by(SessionImage.created_at.asc(), SessionImage.id.asc()).all()
        scene_images = [
            image for image in session_images
            if image.asset_id and str(image.image_type or "").lower() not in {"costume", "costume_initial", "costume_reference", "closet_costume"}
        ]
        settings = self._user_setting_service.apply_cinema_novel_text_generation_settings(payload.get("text_options") or {})
        student_speaker_name = self._chat_session_learning_student_name(chat_session)
        learning_role_character_ids = self._chat_session_learning_role_character_ids(chat_session)
        source_payload = self._chat_session_novel_source_payload(
            chat_session,
            exportable_messages,
            scene_images,
            student_speaker_name=student_speaker_name,
        )
        draft = self._generate_chat_session_novel_draft(source_payload, settings)
        title = str(draft.get("title") or payload.get("title") or chat_session.title or "チャットノベル").strip()[:255] or "チャットノベル"
        subtitle = str(draft.get("subtitle") or "チャットセッションから作成").strip()[:255] or None
        scenes = self._normalize_chat_session_novel_scenes(
            draft.get("scenes"),
            exportable_messages,
            scene_images,
            project_id=chat_session.project_id,
            player_speaker_name=chat_session.player_name,
            student_speaker_name=student_speaker_name,
            learning_character_ids=learning_role_character_ids,
        )
        if not scenes:
            raise ValueError("ノベルシーンを作成できませんでした。")
        first_asset_id = self._first_scene_asset_id(scenes)
        novel = CinemaNovel(
            project_id=chat_session.project_id,
            created_by_user_id=user_id,
            title=title,
            subtitle=subtitle,
            description=str(draft.get("description") or f"チャットセッション「{chat_session.title or chat_session.id}」をノベル化しました。").strip(),
            status="draft",
            mobile_visible=True,
            mode="cinema_novel",
            cover_asset_id=first_asset_id,
            poster_asset_id=first_asset_id,
            production_json=json_util.dumps(
                {
                    "source_type": "live_chat_session_novel",
                    "source_session_id": chat_session.id,
                    "source_room_id": chat_session.room_id,
                    "source_message_ids": [message.id for message in exportable_messages],
                    "source_image_ids": [image.id for image in scene_images],
                    "usage": draft.get("usage"),
                    "model": draft.get("model"),
                }
            ),
        )
        db.session.add(novel)
        db.session.flush()
        chapter = CinemaNovelChapter(
            novel_id=novel.id,
            chapter_no=1,
            title=str(draft.get("chapter_title") or "セッション回想").strip()[:255] or "セッション回想",
            body_markdown=self._chat_session_novel_body_markdown(title, scenes),
            scene_json=json_util.dumps(scenes),
            cover_asset_id=first_asset_id,
            sort_order=1,
        )
        db.session.add(chapter)
        db.session.commit()
        return novel

    def _chat_session_learning_student_name(self, chat_session) -> str | None:
        room_snapshot = self._load_json(getattr(chat_session, "room_snapshot_json", None), default={})
        settings = self._load_json(getattr(chat_session, "settings_json", None), default={})
        room_snapshot = room_snapshot if isinstance(room_snapshot, dict) else {}
        settings = settings if isinstance(settings, dict) else {}
        genre = (
            room_snapshot.get("genre")
            or room_snapshot.get("live_chat_genre")
            or settings.get("live_chat_genre")
            or settings.get("genre")
        )
        if str(genre or "").strip().lower() != "learning":
            return None
        name = str(room_snapshot.get("student_character_name") or "").strip()
        if name:
            return name
        try:
            character_id = int(room_snapshot.get("student_character_id") or settings.get("learning_student_character_id") or 0)
        except (TypeError, ValueError):
            character_id = 0
        if character_id:
            character = Character.query.filter_by(id=character_id, project_id=chat_session.project_id).first()
            if character:
                return str(character.name or "").strip() or None
        return None

    def _chat_session_learning_role_character_ids(self, chat_session) -> list[int]:
        room_snapshot = self._load_json(getattr(chat_session, "room_snapshot_json", None), default={})
        settings = self._load_json(getattr(chat_session, "settings_json", None), default={})
        room_snapshot = room_snapshot if isinstance(room_snapshot, dict) else {}
        settings = settings if isinstance(settings, dict) else {}
        genre = (
            room_snapshot.get("genre")
            or room_snapshot.get("live_chat_genre")
            or settings.get("live_chat_genre")
            or settings.get("genre")
        )
        if str(genre or "").strip().lower() != "learning":
            return []
        return self._normalize_character_ids(
            [
                room_snapshot.get("teacher_character_id") or room_snapshot.get("character_id"),
                room_snapshot.get("student_character_id") or settings.get("learning_student_character_id"),
            ]
        )

    def _chat_session_novel_source_payload(self, chat_session, messages, images, *, student_speaker_name: str | None = None):
        room_snapshot = self._load_json(getattr(chat_session, "room_snapshot_json", None), default={})
        settings = self._load_json(getattr(chat_session, "settings_json", None), default={})
        room_snapshot = room_snapshot if isinstance(room_snapshot, dict) else {}
        settings = settings if isinstance(settings, dict) else {}
        message_items = []
        for message in messages[-80:]:
            sender_type = str(message.sender_type or "").strip().lower()
            speaker_name = message.speaker_name or ("プレイヤー" if sender_type == "user" else "")
            if student_speaker_name and sender_type == "user":
                speaker_name = student_speaker_name
            message_items.append(
                {
                    "id": message.id,
                    "order_no": message.order_no,
                    "sender_type": message.sender_type,
                    "speaker_name": speaker_name,
                    "text": str(message.message_text or "").strip()[:900],
                    "created_at": message.created_at.isoformat() if getattr(message, "created_at", None) else None,
                }
            )
        image_items = []
        for index, image in enumerate(images[-40:], start=1):
            state_json = self._load_json(getattr(image, "state_json", None), default={})
            state_json = state_json if isinstance(state_json, dict) else {}
            displayed = state_json.get("displayed_image_observation") if isinstance(state_json.get("displayed_image_observation"), dict) else {}
            summary = displayed.get("short_summary") or state_json.get("focus_summary") or state_json.get("conversation_image_prompt") or ""
            image_items.append(
                {
                    "index": index,
                    "session_image_id": image.id,
                    "asset_id": image.asset_id,
                    "image_type": image.image_type,
                    "prompt_text": str(image.prompt_text or "").strip()[:700],
                    "summary": str(summary or "").strip()[:700],
                    "created_at": image.created_at.isoformat() if getattr(image, "created_at", None) else None,
                }
            )
        return {
            "session": {
                "id": chat_session.id,
                "title": chat_session.title,
                "player_name": chat_session.player_name,
                "student_speaker_name": student_speaker_name,
                "room_id": chat_session.room_id,
                "room_title": room_snapshot.get("room_title") or room_snapshot.get("title"),
                "genre": room_snapshot.get("genre") or settings.get("live_chat_genre") or settings.get("genre"),
            },
            "messages": message_items,
            "images": image_items,
        }

    def _generate_chat_session_novel_draft(self, source_payload: dict, settings: dict):
        prompt = "\n".join(
            [
                "次のライブチャット履歴を、ビジュアルノベルとして読みやすい1章に編集してください。",
                "目的は動画ではなく、既存のノベル画面で読むためのシーン化です。",
                "会話ログをそのまま全件コピーせず、流れがわかるように整理し、テロップ/地の文とセリフを混ぜてください。",
                "学習ルームの場合は説明の要点を自然な講義ノベルにし、恋愛ルームの場合は空気感と掛け合いを残してください。",
                "画像がある場合は、合いそうな scene に image_index を指定してください。画像を無理に全シーンへ割り当てないでください。",
                "必ずJSONだけで返してください。",
                "形式: {\"title\":\"...\",\"subtitle\":\"...\",\"description\":\"...\",\"chapter_title\":\"...\",\"scenes\":[{\"type\":\"narration|dialogue\",\"speaker\":\"\",\"text\":\"...\",\"source_message_ids\":[1,2],\"image_index\":1}]}",
                "scenes は 8〜24件。text は1シーンあたり日本語で40〜180字程度。speaker は地の文なら空文字。",
                "",
                json_util.dumps(source_payload),
            ]
        )
        result = self._text_ai_client.extract_state_json(prompt, model=settings.get("model"))
        parsed = result.get("parsed_json")
        if not isinstance(parsed, dict):
            raise ValueError("ノベル化AIのJSONを解析できませんでした。")
        parsed["model"] = result.get("model")
        parsed["usage"] = result.get("usage")
        return parsed

    def _normalize_chat_session_novel_scenes(
        self,
        raw_scenes,
        messages,
        images,
        *,
        project_id: int | None = None,
        player_speaker_name: str | None = None,
        student_speaker_name: str | None = None,
        learning_character_ids=None,
    ):
        raw_scenes = raw_scenes if isinstance(raw_scenes, list) else []
        message_by_id = {int(message.id): message for message in messages}
        image_by_index = {index: image for index, image in enumerate(images[-40:], start=1)}
        player_aliases = {str(player_speaker_name or "").strip(), "プレイヤー", "Player", "player"}
        player_aliases = {item for item in player_aliases if item}
        learning_character_ids = self._normalize_character_ids(learning_character_ids)
        learning_character_names = self._character_names_for_ids(project_id, learning_character_ids) if project_id and learning_character_ids else []
        scenes = []
        for raw in raw_scenes[:40]:
            if not isinstance(raw, dict):
                continue
            text = str(raw.get("text") or "").strip()
            if not text:
                continue
            speaker = str(raw.get("speaker") or "").strip()[:80]
            source_ids = []
            for value in raw.get("source_message_ids") or []:
                try:
                    message_id = int(value)
                except (TypeError, ValueError):
                    continue
                if message_id in message_by_id:
                    source_ids.append(message_id)
            scene_type = str(raw.get("type") or ("dialogue" if speaker else "narration")).strip().lower()
            if scene_type not in {"dialogue", "narration"}:
                scene_type = "dialogue" if speaker else "narration"
            if student_speaker_name and scene_type == "dialogue":
                source_messages = [message_by_id.get(message_id) for message_id in source_ids]
                has_user_source = any(str(getattr(message, "sender_type", "") or "").strip().lower() == "user" for message in source_messages if message)
                if has_user_source and (not speaker or speaker in player_aliases):
                    speaker = student_speaker_name
            try:
                image = image_by_index.get(int(raw.get("image_index") or 0))
            except (TypeError, ValueError):
                image = None
            if not image:
                image = self._image_for_source_messages(source_ids, message_by_id, images)
            scene_payload = {
                    "id": f"01-{len(scenes) + 1:03d}",
                    "type": scene_type,
                    "speaker": speaker if scene_type == "dialogue" else "",
                    "text": text,
                    "background_asset_id": None,
                    "still_asset_id": image.asset_id if image else None,
                    "choice_list": [],
                    "source_message_ids": source_ids,
                    "source_session_image_id": image.id if image else None,
                }
            if learning_character_ids:
                scene_payload["character_ids"] = learning_character_ids
                scene_payload["characters"] = learning_character_names
            scenes.append(scene_payload)
        if scenes:
            return scenes
        for message in messages[-24:]:
            text = str(message.message_text or "").strip()
            if not text:
                continue
            image = self._image_for_source_messages([message.id], message_by_id, images)
            sender = str(message.sender_type or "").lower()
            speaker = message.speaker_name or ("プレイヤー" if sender == "user" else "")
            if student_speaker_name and sender == "user":
                speaker = student_speaker_name
            scene_payload = {
                    "id": f"01-{len(scenes) + 1:03d}",
                    "type": "dialogue" if sender != "system" else "narration",
                    "speaker": speaker,
                    "text": text,
                    "background_asset_id": None,
                    "still_asset_id": image.asset_id if image else None,
                    "choice_list": [],
                    "source_message_ids": [message.id],
                    "source_session_image_id": image.id if image else None,
                }
            if learning_character_ids:
                scene_payload["character_ids"] = learning_character_ids
                scene_payload["characters"] = learning_character_names
            scenes.append(scene_payload)
        return scenes

    def _image_for_source_messages(self, source_ids, message_by_id, images):
        if not images or not source_ids:
            return None
        source_messages = [message_by_id.get(int(message_id)) for message_id in source_ids if message_by_id.get(int(message_id))]
        source_messages = [message for message in source_messages if getattr(message, "created_at", None)]
        if not source_messages:
            return None
        anchor = min(message.created_at for message in source_messages)
        before = [image for image in images if getattr(image, "created_at", None) and image.created_at <= anchor]
        if before:
            return before[-1]
        return images[0]

    def _chat_session_novel_body_markdown(self, title: str, scenes: list[dict]):
        lines = [f"# {title}", ""]
        for scene in scenes:
            speaker = str(scene.get("speaker") or "").strip()
            text = str(scene.get("text") or "").strip()
            if not text:
                continue
            if speaker:
                lines.append(f"**{speaker}**")
            lines.append(text)
            lines.append("")
        return "\n".join(lines).strip()

    def _first_scene_asset_id(self, scenes: list[dict]):
        for scene in scenes or []:
            asset_id = scene.get("still_asset_id") or scene.get("background_asset_id")
            if asset_id:
                return asset_id
        return None

    def _manual_short_video_scene(self, project_id: int, raw_scene: dict, *, index: int = 0) -> dict:
        character_ids = self._normalize_character_ids(raw_scene.get("character_ids"))
        characters = self._character_names_for_ids(project_id, character_ids)
        outfit_ids = self._normalize_scene_outfit_ids(project_id, character_ids, raw_scene.get("outfit_ids"))
        caption = self._normalize_short_comic_caption(raw_scene.get("caption") or raw_scene.get("text") or "")
        image_prompt = str(raw_scene.get("image_prompt") or raw_scene.get("scene_text") or "").strip()
        return {
            "id": str(raw_scene.get("id") or f"manual-{datetime.utcnow().strftime('%Y%m%d%H%M%S%f')}-{index}"),
            "type": "manual_short_video",
            "speaker": "",
            "text": caption,
            "caption": caption,
            "comic_layout": "short",
            "comic_page": False,
            "page_title": caption,
            "page_panels": [],
            "dialogue": [],
            "tone": str(raw_scene.get("tone") or "manual").strip(),
            "emotion": str(raw_scene.get("emotion") or "").strip(),
            "shot": str(raw_scene.get("shot") or "social video still").strip(),
            "visual_focus": image_prompt,
            "image_prompt": image_prompt,
            "character_ids": character_ids,
            "outfit_ids": outfit_ids,
            "characters": characters,
            "background_asset_id": None,
            "still_asset_id": raw_scene.get("still_asset_id"),
            "choice_list": [],
            "panel_index": index,
        }

    def generate_short_comic_thumbnail(self, novel_id: int, payload: dict | None = None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        production = self._load_json(novel.production_json, default={})
        manual_short_video = production.get("source_type") == "manual_short_video"
        storyboard = production.get("storyboard") if isinstance(production.get("storyboard"), dict) else {}
        panels = storyboard.get("pages") if isinstance(storyboard.get("pages"), list) else storyboard.get("panels")
        panels = panels if isinstance(panels, list) else []
        first_panel = panels[0] if panels and isinstance(panels[0], dict) else {}
        title = str(storyboard.get("title") or novel.title or "").strip()
        logline = str(storyboard.get("logline") or novel.subtitle or novel.description or "").strip()
        if manual_short_video:
            chapters = self.list_chapters(novel.id)
            first_chapter = chapters[0] if chapters else None
            scenes = self._load_json(first_chapter.scene_json, default=[]) if first_chapter else []
            scenes = scenes if isinstance(scenes, list) else []
            first_scene = next((scene for scene in scenes if isinstance(scene, dict)), {})
            if first_scene:
                first_panel = first_scene
                scene_line = str(first_scene.get("image_prompt") or first_scene.get("visual_focus") or "").strip()
                scene_caption = str(first_scene.get("caption") or first_scene.get("text") or "").strip()
                logline = scene_line or scene_caption or logline
        headline = self._normalize_short_comic_caption(payload.get("caption") or title or logline or "ショート漫画")
        source_input = production.get("source_input") if isinstance(production.get("source_input"), dict) else payload
        main_character = self._main_character_name_from_payload(novel.project_id, source_input)
        if manual_short_video and not main_character:
            main_character = "、".join(str(item) for item in first_panel.get("characters") or [] if str(item).strip())
        searchable = "\n".join(
            [
                title,
                logline,
                main_character,
                " ".join(str(item) for item in first_panel.get("characters") or []),
                str(first_panel.get("visual_focus") or ""),
                str(first_panel.get("image_prompt") or ""),
            ]
        )
        reference_asset_ids = self._scene_character_reference_asset_ids(novel.project_id, first_panel) if manual_short_video else []
        if manual_short_video:
            reference_asset_ids = self._apply_novel_session_outfit_references(
                novel,
                reference_asset_ids,
                character_ids=self._normalize_character_ids(first_panel.get("character_ids")),
            )
        if reference_asset_ids:
            reference_paths, reference_asset_ids = self._resolve_reference_image_paths(reference_asset_ids)
        else:
            references = self._matching_character_references(novel.project_id, searchable, limit=3)
            reference_asset_ids = self._apply_novel_session_outfit_references(
                novel,
                [item.get("base_asset_id") for item in references if item.get("base_asset_id")],
                character_ids=[item.get("id") for item in references if item.get("id")],
            )
            reference_paths, reference_asset_ids = self._resolve_reference_image_paths(reference_asset_ids)
        options = self._user_setting_service.apply_cinema_novel_image_generation_settings(payload.get("image_options") or payload)
        options["size"] = "1024x1536" if bool(getattr(novel, "mobile_visible", True)) else "1536x1024"
        options["quality"] = options.get("quality") or "medium"
        options["model"] = options.get("model") or options.get("image_ai_model") or "gpt-image-2"
        art_style = self._short_comic_art_style(source_input)
        prompt = self._short_comic_thumbnail_prompt(novel, headline, first_panel, main_character=main_character, logline=logline, art_style=art_style)
        if manual_short_video:
            prompt = "\n".join(
                [
                    prompt,
                    "",
                    "Manual short video opening rule:",
                    "Base this opening thumbnail on the first manual scene's scene text, selected reference characters, image headline, visible action, props, and emotion.",
                    "Do not invent an unrelated manga premise. The opening must feel like a strong title image for the same short video the viewer will watch next.",
                    "Use Feed-like impact: event-first composition, readable environment, dramatic social-media incident energy, and expressive character reaction.",
                ]
            )
        revision_prompt = "\n".join(
            str(payload.get(key) or "").strip()
            for key in ("image_prompt", "revision_prompt")
            if str(payload.get(key) or "").strip()
        )
        if revision_prompt:
            prompt = f"{prompt}\n\nUser revision request:\n{revision_prompt}"
        asset = self._generate_cinema_asset(
            project_id=novel.project_id,
            asset_type="cinema_novel_comic_thumbnail",
            file_prefix=f"short_comic_{novel.id}_thumbnail",
            prompt=prompt,
            image_options=options,
            metadata={
                "source": "short_comic_thumbnail",
                "novel_id": novel.id,
                "headline": headline,
                "reference_asset_ids": reference_asset_ids,
            },
            reference_paths=reference_paths,
        )
        self._push_asset_history(production, "cover_asset_history_ids", novel.poster_asset_id or novel.cover_asset_id)
        novel.cover_asset_id = asset.id
        novel.poster_asset_id = asset.id
        novel.production_json = json_util.dumps(production)
        db.session.add(novel)
        db.session.commit()
        return {
            "novel": self.serialize_novel(novel, include_chapters=True),
            "asset": self._serialize_asset(asset.id),
        }

    def generate_short_comic_end_card(self, novel_id: int, payload: dict | None = None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        production = self._load_json(novel.production_json, default={})
        storyboard = production.get("storyboard") if isinstance(production.get("storyboard"), dict) else {}
        source_input = production.get("source_input") if isinstance(production.get("source_input"), dict) else payload
        panels = storyboard.get("pages") if isinstance(storyboard.get("pages"), list) else storyboard.get("panels")
        panels = panels if isinstance(panels, list) else []
        last_panel = panels[-1] if panels and isinstance(panels[-1], dict) else {}
        main_character = self._main_character_name_from_payload(novel.project_id, source_input)
        searchable = "\n".join(
            [
                str(storyboard.get("title") or novel.title or ""),
                main_character,
                " ".join(str(item) for item in last_panel.get("characters") or []),
                str(last_panel.get("visual_focus") or ""),
                str(last_panel.get("image_prompt") or ""),
            ]
        )
        references = self._matching_character_references(novel.project_id, searchable, limit=3)
        reference_paths, reference_asset_ids = self._resolve_reference_image_paths(
            [item.get("base_asset_id") for item in references if item.get("base_asset_id")]
        )
        options = self._user_setting_service.apply_cinema_novel_image_generation_settings(payload.get("image_options") or payload)
        options["size"] = "1024x1536" if bool(getattr(novel, "mobile_visible", True)) else "1536x1024"
        options["quality"] = options.get("quality") or "medium"
        options["model"] = options.get("model") or options.get("image_ai_model") or "gpt-image-2"
        art_style = self._short_comic_art_style(source_input)
        prompt = self._short_comic_end_card_prompt(novel, last_panel, main_character=main_character, art_style=art_style)
        revision_prompt = "\n".join(
            str(payload.get(key) or "").strip()
            for key in ("caption", "image_prompt", "revision_prompt")
            if str(payload.get(key) or "").strip()
        )
        if revision_prompt:
            prompt = f"{prompt}\n\nUser revision request:\n{revision_prompt}"
        asset = self._generate_cinema_asset(
            project_id=novel.project_id,
            asset_type="cinema_novel_comic_end_card",
            file_prefix=f"short_comic_{novel.id}_end_card",
            prompt=prompt,
            image_options=options,
            metadata={
                "source": "short_comic_end_card",
                "novel_id": novel.id,
                "reference_asset_ids": reference_asset_ids,
            },
            reference_paths=reference_paths,
        )
        self._push_asset_history(production, "end_card_asset_history_ids", production.get("end_card_asset_id"))
        production["end_card_asset_id"] = asset.id
        novel.production_json = json_util.dumps(production)
        db.session.add(novel)
        db.session.commit()
        return {
            "novel": self.serialize_novel(novel, include_chapters=True),
            "asset": self._serialize_asset(asset.id),
        }

    def generate_short_comic_storyboard(self, project_id: int, payload: dict | None):
        payload = dict(payload or {})
        project = Project.query.get(project_id)
        world = World.query.filter_by(project_id=project_id).first()
        main_character = self._main_character_name_from_payload(project_id, payload)
        genre = str(payload.get("genre") or "面白い").strip() or "面白い"
        theme = str(payload.get("theme") or "").strip()
        title = str(payload.get("title") or "").strip()
        art_style = self._short_comic_art_style(payload)
        prompt_visual_label = "manga thumbnail panel" if art_style == "anime" else "reference-style visual still"
        is_photo_book = "写真集" in genre
        genre_guidance = (
            "ジャンルが写真集の場合: 起承転結の物語より、主人公の魅力を見せる短いビジュアルカット集にしてください。"
            "各コマはダンス、歌唱、朝の挨拶、街歩き、ステージ照明、カフェ、振り向き、クローズアップ、全身ポーズ、"
            "衣装違い、笑顔、少し照れた表情など、同じ構図の繰り返しを避けた多彩なシーンにする。"
            "caption は『おはよう』『歌声が光る』『踊る夜』のように短く、サムネ文字として映える言葉にする。"
            "image_prompt は人物写真集・MVスチル・アイドル/モデル撮影のような構図、ポーズ、光、背景を具体的に書く。"
            #"ただし露骨な性的表現や過度な露出にはしない。"
        ) if is_photo_book else (
            "ジャンルに合わせて、短い物語として続きが気になる展開、変化、オチ、余韻を作ってください。"
        )
        try:
            target_panel_count = int(payload.get("target_panel_count") or 20)
        except (TypeError, ValueError):
            target_panel_count = 20
        target_panel_count = max(4, min(40, target_panel_count))
        reference_sources = self._reference_sources(payload)
        character_context = (
            self._registered_character_context(project_id, main_character=main_character)
            if self._reference_source_enabled(reference_sources, "characters")
            else ""
        )
        reference_context = self._production_reference_context(
            project_id,
            reference_sources=reference_sources,
            main_character=main_character,
            project=project,
            world=world,
            include_characters=False,
        )
        source_novel_context = self._source_novel_context_for_prompt(project_id, payload)
        source_character_names = self._character_names_for_ids(
            project_id,
            self._source_novel_character_ids_from_payload(project_id, payload),
        )
        prompt = "\n".join(
            [
                "Return only JSON.",
                "縦ショート漫画動画用の絵コンテを日本語で作ってください。",
                "これは小説ではなく、1コマ1画像で流すショート動画です。",
                (
                    "原作ノベルが提供されています。その章順、主要場面、感情の流れ、結末を圧縮してショート漫画化してください。別の新作にしないこと。"
                    if source_novel_context
                    else ""
                ),
                (
                    "Source novel fixed cast: "
                    + ", ".join(source_character_names)
                    + ". Keep these characters as the active cast, include them in characters/image_prompt when the scene is based on their conversation, and do not replace them with unrelated characters."
                    if source_character_names
                    else ""
                ),
                f"Hard constraint: 指定主人公は「{main_character or '未指定'}」です。未指定でない場合、title/logline/panels は必ずこの主人公を中心に作ること。",
                f"Hard constraint: main_character が「{main_character or ''}」なら、caption または image_prompt または characters に「{main_character or ''}」を頻繁に入れること。",
                "Hard constraint: 登録キャラクター文脈に別キャラがいても、指定主人公を別キャラへ置き換えないこと。",
                f"目安コマ数: {target_panel_count}。物語として自然なら±2コマまで許可。",
                genre_guidance,
                "初見視聴者向けに再脚本化すること。原作を知らない人が、誰が・何をしようとして・何がズレて・なぜ面白いのかを字幕だけで追える構成にする。",
                "冒頭3コマは必ず前提を明示する。例: 主人公の役割、今回やろうとしている実験/事件/勝負、相手や障害。",
                "中盤は因果をつなぐ。単なる名詞見出しではなく『保存したくて動けない』『祈りすぎて殴れない』のように、原因とズレが分かる短文にする。",
                "終盤はオチの意味を明示する。何に失敗し、何が意外に成功したのかが初見でも分かるようにする。",
                "各コマの caption は日本語22文字以内。画像内にそのまま大きく入れるので短く、強く、誤字なく。名詞だけのラベルは禁止。",
                "caption は『赤グローブ館長』のような既読者向け見出しではなく、『拳闘文化を保存します』『保存したくて動けない』のような初見向けの状況説明にする。",
                "caption にはできるだけ動詞を入れる。『誰が何をする/できない/勘違いする/止まる/守る』が分かる文にする。",
                "各コマに viewer_context, visible_event, joke, story_function を入れる。viewer_context は初見者に必要な前提、visible_event は画面で見せる行動、joke は笑いのズレ、story_function は setup/escalation/payoff のいずれか。",
                (
                    f"各コマの image_prompt は gpt-image-2 用。スマホ版なら9:16 vertical {prompt_visual_label}、"
                    f"スマホ版でないなら16:9 horizontal {prompt_visual_label}。no speech bubbles except the exact headline text."
                ),
                "image_prompt には caption を画像内テキストとして正確に入れる指示を含める。",
                "シリアス、コメディ、アクション、感動などの tone をコマごとに明示し、構図と表情に反映する。",
                "複数キャラ可。ただし各コマは基本1〜2人、必要時のみ3人。",
                "登録キャラクターを使う場合、characters に名前を入れる。",
                f"Art style: {art_style}.",
                "Art style が reference_image の場合、漫画化・アニメ化せず、登録キャラクターの基準画像の雰囲気・絵柄・衣装感・レンダリングを最優先する。",
                "Art style が anime の場合、鮮やかで読みやすい現代アニメ調・漫画サムネ調にする。",
                "Main character が指定されている場合、その人物を必ず主人公にしてください。別の主人公へ置き換えないでください。",
                "主人公は冒頭、転換点、ラストに必ず登場させ、可能なら多くのコマの characters に含めてください。",
                "Use only the enabled reference sections below as story material. Do not invent facts that contradict them.",
                "Required JSON shape:",
                '{"title": "...", "logline": "...", "first_time_viewer_summary": "...", "target_panel_count": 20, "panels": [{"caption": "...", "viewer_context": "...", "visible_event": "...", "joke": "...", "story_function": "setup", "tone": "comedy", "emotion": "...", "shot": "close-up", "visual_focus": "...", "characters": ["..."], "image_prompt": "..."}]}',
                "",
                f"Requested title: {title}",
                f"Genre: {genre}",
                f"Main character: {main_character or 'AIが選定'}",
                f"Theme/request: {theme}",
                f"Output aspect: {'9:16 vertical smartphone panel' if self._normalize_bool(payload.get('mobile_visible', True)) else '16:9 horizontal desktop panel'}",
                "",
                "Enabled reference sources:",
                ", ".join(reference_sources),
                "",
                "Reference material:",
                reference_context or "追加参照なし",
                "",
                "Source novel to adapt:",
                source_novel_context or "none",
                "",
                "Registered characters:",
                character_context or "なし",
            ]
        )
        settings = self._user_setting_service.apply_cinema_novel_text_generation_settings(payload.get("text_options") or {})
        result = self._text_ai_client.extract_state_json(
            prompt,
            model=settings.get("model"),
        )
        parsed = result.get("parsed_json")
        if not isinstance(parsed, dict):
            raise RuntimeError("short comic storyboard response was not valid JSON")
        panels = parsed.get("panels")
        if not isinstance(panels, list) or not panels:
            raise RuntimeError("short comic storyboard did not include panels")
        if main_character:
            for key in ("title", "logline"):
                if isinstance(parsed.get(key), str):
                    parsed[key] = parsed[key].replace("??", main_character).replace("？？", main_character)
            for panel in panels:
                if not isinstance(panel, dict):
                    continue
                for key in ("caption", "visual_focus", "image_prompt"):
                    if isinstance(panel.get(key), str):
                        panel[key] = panel[key].replace("??", main_character).replace("？？", main_character)
                characters = panel.get("characters") if isinstance(panel.get("characters"), list) else []
                characters = [main_character if str(item).strip() in {"??", "？？", ""} else str(item) for item in characters]
                if main_character not in characters:
                    characters = [main_character] + characters[:2]
                panel["characters"] = characters
                image_prompt = str(panel.get("image_prompt") or "")
                if main_character not in image_prompt:
                    panel["image_prompt"] = f"Main protagonist {main_character}. {image_prompt}".strip()
        parsed["target_panel_count"] = target_panel_count
        parsed["model"] = result.get("model")
        parsed["usage"] = result.get("usage")
        return parsed

    def generate_comic_page_storyboard(self, project_id: int, payload: dict | None):
        payload = dict(payload or {})
        project = Project.query.get(project_id)
        world = World.query.filter_by(project_id=project_id).first()
        main_character = self._main_character_name_from_payload(project_id, payload)
        genre = str(payload.get("genre") or "面白い").strip() or "面白い"
        theme = str(payload.get("theme") or "").strip()
        title = str(payload.get("title") or "").strip()
        art_style = self._short_comic_art_style(payload)
        try:
            target_page_count = int(payload.get("target_page_count") or payload.get("target_panel_count") or 4)
        except (TypeError, ValueError):
            target_page_count = 4
        target_page_count = max(1, min(12, target_page_count))
        reference_sources = self._reference_sources(payload)
        character_context = (
            self._registered_character_context(project_id, main_character=main_character)
            if self._reference_source_enabled(reference_sources, "characters")
            else ""
        )
        reference_context = self._production_reference_context(
            project_id,
            reference_sources=reference_sources,
            main_character=main_character,
            project=project,
            world=world,
            include_characters=False,
        )
        source_novel_context = self._source_novel_context_for_prompt(project_id, payload)
        source_character_names = self._character_names_for_ids(
            project_id,
            self._source_novel_character_ids_from_payload(project_id, payload),
        )
        prompt = "\n".join(
            [
                "Return only JSON.",
                "Create a Japanese manga-page storyboard, not a prose novel and not single-panel social video frames.",
                "The output will be used by gpt-image-2 to draw one complete manga page per page item.",
                (
                    "Source novel is provided. Adapt that novel's plot, chapter order, important scenes, character emotions, and ending into manga pages. Do not make a different story."
                    if source_novel_context
                    else ""
                ),
                "Each manga page should contain 2 to 6 panels, clear gutters, speech balloons, short readable dialogue, and optional small narration boxes.",
                "Keep each speech balloon short. Use Japanese dialogue of about 4 to 16 characters per balloon whenever possible.",
                "Avoid tiny text, dense paragraphs, excessive captions, or too many balloons.",
                f"Target page count: {target_page_count}. A little fewer is acceptable if the story works.",
                f"Hard constraint: main protagonist is {main_character or 'not specified'}. If specified, use this protagonist consistently.",
                "Use registered characters if relevant. Do not replace the specified protagonist with another character.",
                (
                    "Source novel fixed cast: "
                    + ", ".join(source_character_names)
                    + ". Keep these characters as the active cast, include them in characters/image_prompt when the page is based on their conversation, and do not replace them with unrelated characters."
                    if source_character_names
                    else ""
                ),
                "For each page, provide page_title, summary, tone, characters, and panels.",
                "For each panel, provide panel_no, composition, action, speaker, speech, sfx, and background.",
                "Also provide image_prompt for each page. The image_prompt must describe the whole page layout and include the exact dialogue/sfx text to draw.",
                "If art_style is reference_image, preserve reference character identity and rendering taste, while composing the page as manga panels.",
                "If art_style is anime, use polished modern manga/anime page art.",
                "Required JSON shape:",
                '{"title":"...","logline":"...","target_page_count":4,"pages":[{"page_title":"...","summary":"...","tone":"comedy","characters":["..."],"panels":[{"panel_no":1,"composition":"...","action":"...","speaker":"...","speech":"...","sfx":"...","background":"..."}],"image_prompt":"..."}]}',
                "",
                f"Requested title: {title}",
                f"Genre: {genre}",
                f"Main character: {main_character or 'AI chooses'}",
                f"Theme/request: {theme}",
                f"Page aspect: {'9:16 vertical smartphone manga page' if self._normalize_bool(payload.get('mobile_visible', True)) else 'portrait manga page, readable on desktop'}",
                f"Art style: {art_style}",
                "",
                "Enabled reference sources:",
                ", ".join(reference_sources),
                "",
                "Reference material:",
                reference_context or "none",
                "",
                "Source novel to adapt:",
                source_novel_context or "none",
                "",
                "Registered characters:",
                character_context or "none",
            ]
        )
        settings = self._user_setting_service.apply_cinema_novel_text_generation_settings(payload.get("text_options") or {})
        result = self._text_ai_client.extract_state_json(
            prompt,
            model=settings.get("model"),
        )
        parsed = result.get("parsed_json")
        if not isinstance(parsed, dict):
            raise RuntimeError("comic page storyboard response was not valid JSON")
        pages = parsed.get("pages")
        if not isinstance(pages, list) or not pages:
            raise RuntimeError("comic page storyboard did not include pages")
        if main_character:
            for key in ("title", "logline"):
                if isinstance(parsed.get(key), str):
                    parsed[key] = parsed[key].replace("??", main_character).replace("？？", main_character)
            for page in pages:
                if not isinstance(page, dict):
                    continue
                for key in ("page_title", "summary", "image_prompt"):
                    if isinstance(page.get(key), str):
                        page[key] = page[key].replace("??", main_character).replace("？？", main_character)
                characters = page.get("characters") if isinstance(page.get("characters"), list) else []
                characters = [main_character if str(item).strip() in {"??", "？？", ""} else str(item) for item in characters]
                if main_character not in characters:
                    characters = [main_character] + characters[:3]
                page["characters"] = characters
                image_prompt = str(page.get("image_prompt") or "")
                if main_character not in image_prompt:
                    page["image_prompt"] = f"Main protagonist {main_character}. {image_prompt}".strip()
        parsed["target_page_count"] = target_page_count
        parsed["target_panel_count"] = target_page_count
        parsed["comic_layout"] = "page"
        parsed["model"] = result.get("model")
        parsed["usage"] = result.get("usage")
        return parsed

    def generate_short_comic_panel_images(self, novel_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        chapters = self.list_chapters(novel.id)
        chapter = chapters[0] if chapters else None
        if not chapter:
            return None
        scenes = self._load_json(chapter.scene_json, default=[])
        if not isinstance(scenes, list):
            return None
        overwrite = bool(payload.get("overwrite"))
        options = self._user_setting_service.apply_cinema_novel_image_generation_settings(payload.get("image_options") or payload)
        options["size"] = "1024x1536" if bool(getattr(novel, "mobile_visible", True)) else "1536x1024"
        options["quality"] = options.get("quality") or "medium"
        options["model"] = options.get("model") or options.get("image_ai_model") or "gpt-image-2"
        production = self._load_json(novel.production_json, default={})
        source_input = production.get("source_input") if isinstance(production.get("source_input"), dict) else {}
        art_style = self._short_comic_art_style(source_input)
        target_scene_index = payload.get("scene_index")
        if target_scene_index is not None:
            try:
                target_scene_index = int(target_scene_index)
            except (TypeError, ValueError):
                raise ValueError("scene_index is invalid")
            if target_scene_index < 0 or target_scene_index >= len(scenes):
                raise ValueError("scene was not found")
        image_jobs = []
        for scene_index, scene in enumerate(scenes):
            if target_scene_index is not None and scene_index != target_scene_index:
                continue
            if not isinstance(scene, dict):
                continue
            if scene.get("still_asset_id") and not overwrite:
                continue
            caption = self._normalize_short_comic_caption(scene.get("caption") or scene.get("text") or "")
            if not caption and not str(scene.get("image_prompt") or scene.get("visual_focus") or "").strip():
                continue
            searchable = "\n".join(
                [
                    caption,
                    str(scene.get("visual_focus") or ""),
                    str(scene.get("image_prompt") or ""),
                    " ".join(str(item) for item in scene.get("characters") or []),
                ]
            )
            prompt_searchable = "\n".join(
                [
                    caption,
                    str(scene.get("visual_focus") or ""),
                    str(scene.get("image_prompt") or ""),
                ]
            )
            prompt_references = self._matching_character_references(novel.project_id, prompt_searchable, limit=3)
            effective_scene = dict(scene)
            if prompt_references:
                prompt_character_ids = [item.get("id") for item in prompt_references if item.get("id")]
                effective_scene["character_ids"] = prompt_character_ids
                effective_scene["characters"] = [
                    str(item.get("name") or item.get("nickname") or "").strip()
                    for item in prompt_references
                    if str(item.get("name") or item.get("nickname") or "").strip()
                ]
            use_current_image = bool(payload.get("use_current_image")) and bool(scene.get("still_asset_id"))
            current_image_paths, current_image_asset_ids = (
                self._resolve_reference_image_paths([scene.get("still_asset_id")])
                if use_current_image
                else ([], [])
            )
            explicit_reference_ids = []
            if not prompt_references:
                explicit_reference_ids = self._scene_character_reference_asset_ids(novel.project_id, scene)
                explicit_reference_ids = self._apply_novel_session_outfit_references(
                    novel,
                    explicit_reference_ids,
                    character_ids=self._normalize_character_ids(scene.get("character_ids")),
                )
            if explicit_reference_ids:
                reference_paths, reference_asset_ids = self._resolve_reference_image_paths(explicit_reference_ids)
            else:
                # Prefer the text the user just edited. Existing storyboard character labels can be stale
                # after manual prompt edits, and should not override an explicit name in the new prompt.
                references = prompt_references
                if not references:
                    references = self._matching_character_references(novel.project_id, searchable, limit=3)
                reference_asset_ids = self._apply_novel_session_outfit_references(
                    novel,
                    [item.get("base_asset_id") for item in references if item.get("base_asset_id")],
                    character_ids=[item.get("id") for item in references if item.get("id")],
                )
                reference_paths, reference_asset_ids = self._resolve_reference_image_paths(reference_asset_ids)
            prompt = (
                self._comic_page_image_prompt(novel, effective_scene, art_style=art_style)
                if scene.get("comic_page")
                else self._short_comic_image_prompt(novel, effective_scene, caption, art_style=art_style)
            )
            if current_image_paths:
                prompt = "\n".join(
                    [
                        prompt,
                        "",
                        "Current displayed image revision mode:",
                        "Use the first attached image as the current displayed panel to revise, but the revision note must visibly change the result.",
                        "Preserve useful continuity from the current image only where it does not weaken the requested change.",
                        "If the current image is too stiff, plain, expressionless, or too close to a character sheet, make the new result more dynamic and cinematic instead of preserving that stiffness.",
                        "Use the remaining attached reference images only for character identity, facial features, hairstyle, outfit design, and rendering consistency.",
                    ]
                )
                reference_paths = [*current_image_paths, *reference_paths]
                reference_asset_ids = [*current_image_asset_ids, *reference_asset_ids]
            image_jobs.append(
                {
                    "scene_index": scene_index,
                    "project_id": novel.project_id,
                    "asset_type": "cinema_novel_comic_page" if scene.get("comic_page") else "cinema_novel_comic_panel",
                    "file_prefix": f"comic_page_{novel.id}_{scene_index + 1}" if scene.get("comic_page") else f"short_comic_{novel.id}_panel_{scene_index + 1}",
                    "prompt": prompt,
                    "image_options": options,
                    "reference_paths": reference_paths,
                    "reference_asset_ids": reference_asset_ids,
                    "metadata": {
                        "source": "short_comic_panel",
                        "novel_id": novel.id,
                        "chapter_id": chapter.id,
                        "scene_index": scene_index,
                        "caption": caption,
                        "reference_asset_ids": reference_asset_ids,
                        "current_image_revision_asset_ids": current_image_asset_ids,
                    },
                }
            )

        created_assets = []
        failed_assets = []
        for job, result, error in self._generate_cinema_asset_jobs(
            image_jobs,
            parallel=bool(payload.get("parallel", True)),
            max_workers=payload.get("max_workers"),
        ):
            if error or not result:
                failed_assets.append({"scene_index": job.get("scene_index"), "error": error or "image generation failed"})
                continue
            asset = self._create_cinema_asset_from_result(
                project_id=job["project_id"],
                asset_type=job["asset_type"],
                file_prefix=job["file_prefix"],
                image_options=job["image_options"],
                metadata=job["metadata"],
                result=result,
            )
            scene = scenes[job["scene_index"]]
            if not isinstance(scene, dict):
                scene = {}
                scenes[job["scene_index"]] = scene
            self._push_scene_still_history(scene, scene.get("still_asset_id"))
            scene["still_asset_id"] = asset.id
            created_assets.append(self._serialize_asset(asset.id))
            chapter.scene_json = json_util.dumps(scenes)
            db.session.add(chapter)
            db.session.commit()
        return {
            "novel": self.serialize_novel(novel, include_chapters=True),
            "chapter": self.serialize_chapter(chapter),
            "assets": created_assets,
            "failed_assets": failed_assets,
        }

    def update_comic_scene(self, novel_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        chapter_id = payload.get("chapter_id")
        try:
            chapter_id = int(chapter_id)
            scene_index = int(payload.get("scene_index"))
        except (TypeError, ValueError):
            raise ValueError("chapter_id and scene_index are required")
        chapter = self.get_chapter(chapter_id)
        if not chapter or int(chapter.novel_id) != int(novel.id):
            raise ValueError("chapter was not found")
        scenes = self._load_json(chapter.scene_json, default=[])
        if not isinstance(scenes, list) or scene_index < 0 or scene_index >= len(scenes):
            raise ValueError("scene was not found")
        scene = scenes[scene_index]
        if not isinstance(scene, dict):
            scene = {}
            scenes[scene_index] = scene

        if "caption" in payload or "text" in payload:
            raw_text = payload.get("caption") if "caption" in payload else payload.get("text")
            text = str(raw_text or "").strip()
            scene["caption"] = text
            scene["text"] = text
        for key in ("page_title", "image_prompt", "visual_focus"):
            if key in payload:
                value = str(payload.get(key) or "").strip()
                if value:
                    scene[key] = value
                else:
                    scene.pop(key, None)
        if "character_ids" in payload:
            character_ids = self._normalize_character_ids(payload.get("character_ids"))
            scene["character_ids"] = character_ids
            scene["characters"] = self._character_names_for_ids(novel.project_id, character_ids)
        elif "characters" in payload and isinstance(payload.get("characters"), list):
            scene["characters"] = [str(item).strip() for item in payload.get("characters") if str(item).strip()]
        elif any(key in payload for key in ("caption", "text", "image_prompt", "visual_focus")):
            prompt_searchable = "\n".join(
                [
                    str(scene.get("caption") or scene.get("text") or ""),
                    str(scene.get("visual_focus") or ""),
                    str(scene.get("image_prompt") or ""),
                ]
            )
            prompt_references = self._matching_character_references(novel.project_id, prompt_searchable, limit=3)
            if prompt_references:
                character_ids = self._normalize_character_ids([item.get("id") for item in prompt_references if item.get("id")])
                scene["character_ids"] = character_ids
                scene["characters"] = self._character_names_for_ids(novel.project_id, character_ids)
        if "outfit_ids" in payload:
            character_ids = self._normalize_character_ids(scene.get("character_ids"))
            scene["outfit_ids"] = self._normalize_scene_outfit_ids(novel.project_id, character_ids, payload.get("outfit_ids"))

        if "selected_asset_id" in payload:
            try:
                selected_asset_id = int(payload.get("selected_asset_id") or 0)
            except (TypeError, ValueError):
                raise ValueError("selected_asset_id is invalid")
            asset = Asset.query.get(selected_asset_id)
            if not asset or getattr(asset, "deleted_at", None) or int(asset.project_id) != int(novel.project_id):
                raise ValueError("selected image was not found")
            current_asset_id = scene.get("still_asset_id")
            if int(current_asset_id or 0) != selected_asset_id:
                self._push_scene_still_history(scene, current_asset_id)
                scene["still_asset_id"] = selected_asset_id
            scene["still_asset_history_ids"] = [
                asset_id for asset_id in self._scene_still_history_ids(scene)
                if int(asset_id) != selected_asset_id
            ]

        chapter.scene_json = json_util.dumps(scenes)
        db.session.add(chapter)
        db.session.commit()
        return {
            "novel": self.serialize_novel(novel, include_chapters=True),
            "chapter": self.serialize_chapter(chapter),
            "scene_index": scene_index,
        }

    def mutate_comic_scene(self, novel_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        chapter_id = payload.get("chapter_id")
        try:
            chapter_id = int(chapter_id)
            scene_index = int(payload.get("scene_index") or 0)
        except (TypeError, ValueError):
            raise ValueError("chapter_id and scene_index are required")
        chapter = self.get_chapter(chapter_id)
        if not chapter or int(chapter.novel_id) != int(novel.id):
            raise ValueError("chapter was not found")
        scenes = self._load_json(chapter.scene_json, default=[])
        if not isinstance(scenes, list):
            scenes = []
        operation = str(payload.get("operation") or "").strip()
        if operation not in {"add_after", "copy_after", "delete", "reorder"}:
            raise ValueError("operation is required")
        if operation in {"copy_after", "delete"} and (scene_index < 0 or scene_index >= len(scenes)):
            raise ValueError("scene was not found")
        if operation == "reorder":
            try:
                from_scene_index = int(payload.get("from_scene_index"))
                to_scene_index = int(payload.get("to_scene_index"))
            except (TypeError, ValueError):
                raise ValueError("from_scene_index and to_scene_index are required")
            if from_scene_index < 0 or from_scene_index >= len(scenes) or to_scene_index < 0 or to_scene_index >= len(scenes):
                raise ValueError("scene was not found")
            scene = scenes.pop(from_scene_index)
            scenes.insert(to_scene_index, scene)
            chapter.scene_json = json_util.dumps(scenes)
            db.session.add(chapter)
            db.session.commit()
            return {
                "novel": self.serialize_novel(novel, include_chapters=True),
                "chapter": self.serialize_chapter(chapter),
                "scene_index": to_scene_index,
                "operation": operation,
            }

        production = self._load_json(novel.production_json, default={})
        comic_layout = production.get("comic_layout") if isinstance(production, dict) else None
        current_scene = scenes[scene_index] if 0 <= scene_index < len(scenes) and isinstance(scenes[scene_index], dict) else {}
        new_scene_index = max(0, min(scene_index + 1, len(scenes)))

        if operation == "delete":
            scenes.pop(scene_index)
            new_scene_index = min(scene_index, max(0, len(scenes) - 1))
        elif operation == "copy_after":
            scene = copy.deepcopy(current_scene)
            scene["id"] = f"manual-{datetime.utcnow().strftime('%Y%m%d%H%M%S%f')}"
            scene.pop("still_asset_history_ids", None)
            scene["still_asset_id"] = None
            scenes.insert(new_scene_index, scene)
        else:
            caption = str(payload.get("caption") or "新しいコマ").strip() or "新しいコマ"
            image_prompt = str(payload.get("image_prompt") or current_scene.get("image_prompt") or current_scene.get("visual_focus") or "").strip()
            character_ids = self._normalize_character_ids(payload.get("character_ids") or current_scene.get("character_ids") or [])
            characters = self._character_names_for_ids(novel.project_id, character_ids)
            outfit_ids = self._normalize_scene_outfit_ids(novel.project_id, character_ids, payload.get("outfit_ids") or current_scene.get("outfit_ids"))
            if not characters:
                characters = current_scene.get("characters") if isinstance(current_scene.get("characters"), list) else []
            scene = {
                "id": f"manual-{datetime.utcnow().strftime('%Y%m%d%H%M%S%f')}",
                "type": "narration",
                "speaker": "",
                "text": caption,
                "caption": caption,
                "page_title": caption,
                "visual_focus": image_prompt or caption,
                "image_prompt": image_prompt or caption,
                "tone": str(current_scene.get("tone") or "").strip(),
                "emotion": str(current_scene.get("emotion") or "").strip(),
                "shot": str(current_scene.get("shot") or "thumbnail composition").strip(),
                "character_ids": character_ids,
                "outfit_ids": outfit_ids,
                "characters": characters,
                "comic_page": bool(current_scene.get("comic_page")) or comic_layout == "page",
                "page_panels": copy.deepcopy(current_scene.get("page_panels") or []),
                "dialogue": copy.deepcopy(current_scene.get("dialogue") or []),
                "background_asset_id": None,
                "still_asset_id": None,
                "choice_list": [],
            }
            scenes.insert(new_scene_index, scene)

        chapter.scene_json = json_util.dumps(scenes)
        db.session.add(chapter)
        db.session.commit()
        return {
            "novel": self.serialize_novel(novel, include_chapters=True),
            "chapter": self.serialize_chapter(chapter),
            "scene_index": new_scene_index,
            "operation": operation,
        }

    def update_comic_special_image(self, novel_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        kind = str(payload.get("kind") or "").strip()
        if kind not in {"cover", "end_card"}:
            raise ValueError("kind is required")
        try:
            selected_asset_id = int(payload.get("selected_asset_id") or 0)
        except (TypeError, ValueError):
            raise ValueError("selected_asset_id is invalid")
        asset = Asset.query.get(selected_asset_id)
        if not asset or getattr(asset, "deleted_at", None) or int(asset.project_id) != int(novel.project_id):
            raise ValueError("selected image was not found")

        production = self._load_json(novel.production_json, default={})
        if kind == "cover":
            current_asset_id = novel.poster_asset_id or novel.cover_asset_id
            if int(current_asset_id or 0) != selected_asset_id:
                self._push_asset_history(production, "cover_asset_history_ids", current_asset_id)
                novel.cover_asset_id = selected_asset_id
                novel.poster_asset_id = selected_asset_id
            production["cover_asset_history_ids"] = [
                asset_id for asset_id in self._asset_history_ids(production, "cover_asset_history_ids")
                if int(asset_id) != selected_asset_id
            ]
        else:
            current_asset_id = production.get("end_card_asset_id")
            if int(current_asset_id or 0) != selected_asset_id:
                self._push_asset_history(production, "end_card_asset_history_ids", current_asset_id)
                production["end_card_asset_id"] = selected_asset_id
            production["end_card_asset_history_ids"] = [
                asset_id for asset_id in self._asset_history_ids(production, "end_card_asset_history_ids")
                if int(asset_id) != selected_asset_id
            ]
        novel.production_json = json_util.dumps(production)
        db.session.add(novel)
        db.session.commit()
        return {"novel": self.serialize_novel(novel, include_chapters=True), "kind": kind}

    def regenerate_comic_special_image(self, novel_id: int, payload: dict | None):
        payload = dict(payload or {})
        kind = str(payload.get("kind") or "").strip()
        if kind == "cover":
            return self.generate_short_comic_thumbnail(novel_id, payload)
        if kind == "end_card":
            return self.generate_short_comic_end_card(novel_id, payload)
        raise ValueError("kind is required")

    def _scene_still_history_ids(self, scene: dict) -> list[int]:
        raw_history = scene.get("still_asset_history_ids") if isinstance(scene, dict) else []
        if not isinstance(raw_history, list):
            raw_history = []
        history = []
        for asset_id in raw_history:
            try:
                asset_id = int(asset_id)
            except (TypeError, ValueError):
                continue
            if asset_id and asset_id not in history:
                history.append(asset_id)
        return history[:12]

    def _push_scene_still_history(self, scene: dict, asset_id):
        if not isinstance(scene, dict):
            return
        try:
            asset_id = int(asset_id or 0)
        except (TypeError, ValueError):
            asset_id = 0
        if not asset_id:
            return
        history = [item for item in self._scene_still_history_ids(scene) if int(item) != asset_id]
        scene["still_asset_history_ids"] = [asset_id, *history][:12]

    def _asset_history_ids(self, source: dict, key: str) -> list[int]:
        raw_history = source.get(key) if isinstance(source, dict) else []
        if not isinstance(raw_history, list):
            raw_history = []
        history = []
        for asset_id in raw_history:
            try:
                asset_id = int(asset_id)
            except (TypeError, ValueError):
                continue
            if asset_id and asset_id not in history:
                history.append(asset_id)
        return history[:12]

    def _push_asset_history(self, source: dict, key: str, asset_id):
        if not isinstance(source, dict):
            return
        try:
            asset_id = int(asset_id or 0)
        except (TypeError, ValueError):
            asset_id = 0
        if not asset_id:
            return
        history = [item for item in self._asset_history_ids(source, key) if int(item) != asset_id]
        source[key] = [asset_id, *history][:12]

    def _asset_file_path(self, asset_id: int | None):
        if not asset_id:
            return None
        asset = Asset.query.get(asset_id)
        if not asset or getattr(asset, "deleted_at", None) or not asset.file_path:
            return None
        path = Path(asset.file_path)
        return str(path) if path.exists() else None

    def _prepare_powerpoint_canvas(self, image_path, output_path, Image, ImageOps, ImageFilter, ImageEnhance):
        if not image_path:
            return None
        try:
            source = ImageOps.exif_transpose(Image.open(image_path)).convert("RGB")
        except Exception:
            return None

        target_size = (1920, 1080)
        ratio = target_size[0] / target_size[1]
        source_ratio = source.width / source.height if source.height else ratio
        if source_ratio > ratio:
            crop_width = int(source.height * ratio)
            left = max(0, (source.width - crop_width) // 2)
            crop_box = (left, 0, left + crop_width, source.height)
        else:
            crop_height = int(source.width / ratio) if source.width else source.height
            top = max(0, (source.height - crop_height) // 2)
            crop_box = (0, top, source.width, top + crop_height)

        background = source.crop(crop_box).resize(target_size, Image.Resampling.LANCZOS)
        background = background.filter(ImageFilter.GaussianBlur(radius=22))
        background = ImageEnhance.Brightness(background).enhance(0.55)
        foreground = ImageOps.contain(source, target_size, Image.Resampling.LANCZOS)
        canvas = background.copy()
        canvas.paste(foreground, ((target_size[0] - foreground.width) // 2, (target_size[1] - foreground.height) // 2))
        canvas.save(output_path, "JPEG", quality=92, optimize=True)
        return output_path

    def _powerpoint_scene_text(self, text: str) -> str:
        text = re.sub(r"\s+", " ", str(text or "")).strip()
        if len(text) <= 260:
            return text
        return text[:257].rstrip() + "..."

    def _powerpoint_scene_font_size(self, text: str) -> int:
        length = len(str(text or ""))
        if length > 220:
            return 15
        if length > 160:
            return 17
        return 19

    def _safe_powerpoint_filename(self, value: str) -> str:
        safe = re.sub(r"[\\/:*?\"<>|\s]+", "_", str(value or "").strip(), flags=re.UNICODE).strip("._")
        return safe[:80] or "cinema_novel"

    def _load_video_font(self, ImageFont, *, size: int):
        candidates = [
            Path("C:/Windows/Fonts/YuGothM.ttc"),
            Path("C:/Windows/Fonts/YuGothR.ttc"),
            Path("C:/Windows/Fonts/meiryo.ttc"),
            Path("C:/Windows/Fonts/msgothic.ttc"),
        ]
        for path in candidates:
            if path.exists():
                try:
                    return ImageFont.truetype(str(path), size=size)
                except Exception:
                    continue
        return ImageFont.load_default()

    def _render_short_video_frame(
        self,
        image_path,
        output_path,
        text,
        *,
        speaker="",
        footer="",
        title=False,
        Image,
        ImageDraw,
        ImageEnhance,
        ImageFilter,
        ImageOps,
        font_regular,
        font_small,
        font_title,
        font_footer,
        landscape=False,
    ):
        target_size = (1920, 1080) if landscape else (1080, 1920)
        try:
            source = ImageOps.exif_transpose(Image.open(image_path)).convert("RGB") if image_path else None
        except Exception:
            source = None

        if source:
            ratio = target_size[0] / target_size[1]
            source_ratio = source.width / source.height if source.height else ratio
            if source_ratio > ratio:
                crop_width = int(source.height * ratio)
                left = max(0, (source.width - crop_width) // 2)
                crop_box = (left, 0, left + crop_width, source.height)
            else:
                crop_height = int(source.width / ratio) if source.width else source.height
                top = max(0, (source.height - crop_height) // 2)
                crop_box = (0, top, source.width, top + crop_height)
            background = source.crop(crop_box).resize(target_size, Image.Resampling.LANCZOS)
            background = background.filter(ImageFilter.GaussianBlur(radius=22 if landscape else 26))
            background = ImageEnhance.Brightness(background).enhance(0.62)
            foreground_box = (1420, 900) if landscape else (1080, 1440)
            foreground = ImageOps.contain(source, foreground_box, Image.Resampling.LANCZOS)
            canvas = background.copy()
            canvas.paste(foreground, ((target_size[0] - foreground.width) // 2, 42 if landscape else 70))
        else:
            canvas = Image.new("RGB", target_size, (18, 22, 30))

        overlay = Image.new("RGBA", target_size, (0, 0, 0, 0))
        draw = ImageDraw.Draw(overlay)
        box = (140, 742 if not title else 704, 1780, 1016) if landscape else (64, 1310 if not title else 1260, 1016, 1782)
        draw.rounded_rectangle(box, radius=34, fill=(8, 10, 16, 178))
        if speaker:
            speaker_box = (box[0] + 28, box[1] - 50, box[0] + 316, box[1] + 20)
            draw.rounded_rectangle(speaker_box, radius=20, fill=(38, 56, 92, 210))
            draw.text((speaker_box[0] + 144, box[1] - 36), speaker[:16], font=font_small, fill=(255, 255, 255, 255), anchor="ma")

        text_font = font_title if title else font_regular
        lines = self._wrap_video_text(text, draw, text_font, max_width=1500 if landscape else 850)
        line_height = 66 if title else 58
        total_height = len(lines) * line_height
        y = box[1] + max(36, ((box[3] - box[1]) - total_height) // 2)
        for line in lines:
            draw.text((target_size[0] // 2, y), line, font=text_font, fill=(255, 255, 255, 255), anchor="ma")
            y += line_height
        if footer:
            footer_pos = (target_size[0] - 70, target_size[1] - 58) if landscape else (970, 1810)
            draw.text(footer_pos, footer, font=font_footer, fill=(215, 222, 235, 230), anchor="ra")

        canvas = Image.alpha_composite(canvas.convert("RGBA"), overlay).convert("RGB")
        canvas.save(output_path, "PNG")

    def _render_comic_short_frame(
        self,
        image_path,
        output_path,
        caption,
        *,
        label="",
        footer="",
        title=False,
        Image,
        ImageDraw,
        ImageEnhance,
        ImageFilter,
        ImageOps,
        fonts,
        mobile_visible=True,
    ):
        target_size = (1080, 1920) if mobile_visible else (1920, 1080)
        try:
            source = ImageOps.exif_transpose(Image.open(image_path)).convert("RGB") if image_path else None
        except Exception:
            source = None

        if source:
            ratio = target_size[0] / target_size[1]
            source_ratio = source.width / source.height if source.height else ratio
            if source_ratio > ratio:
                crop_width = int(source.height * ratio)
                left = max(0, (source.width - crop_width) // 2)
                crop_box = (left, 0, left + crop_width, source.height)
            else:
                crop_height = int(source.width / ratio) if source.width else source.height
                top = max(0, (source.height - crop_height) // 2)
                crop_box = (0, top, source.width, top + crop_height)
            background = source.crop(crop_box).resize(target_size, Image.Resampling.LANCZOS)
            background = background.filter(ImageFilter.GaussianBlur(radius=24))
            background = ImageEnhance.Brightness(background).enhance(0.45)
            foreground = ImageOps.contain(source, target_size, Image.Resampling.LANCZOS)
            canvas = background.copy()
            canvas.paste(foreground, ((target_size[0] - foreground.width) // 2, (target_size[1] - foreground.height) // 2))
            canvas = ImageEnhance.Contrast(canvas).enhance(1.05)
            canvas = ImageEnhance.Color(canvas).enhance(1.06)
        else:
            canvas = Image.new("RGB", target_size, (18, 22, 30))
        canvas.save(output_path, "PNG")

    def _short_video_frame_fade_in_paths(self, frame_path: Path, Image, *, frame_count: int = 9):
        source = Image.open(frame_path).convert("RGB")
        fade_dir = frame_path.parent / f"{frame_path.stem}_fade"
        fade_dir.mkdir(parents=True, exist_ok=True)
        fade_paths = []
        black = Image.new("RGB", source.size, (0, 0, 0))
        count = max(2, int(frame_count))
        for index in range(count - 1):
            alpha = index / (count - 1)
            frame = Image.blend(black, source, alpha)
            path = fade_dir / f"{frame_path.stem}_fade_{index + 1:02d}.png"
            frame.save(path, "PNG")
            fade_paths.append(path)
        return fade_paths

    def _wrap_video_text(self, text: str, draw, font, *, max_width: int):
        text = re.sub(r"\s+", " ", str(text or "")).strip()
        if not text:
            return []
        lines = []
        current = ""
        for char in text:
            candidate = current + char
            if draw.textlength(candidate, font=font) <= max_width or not current:
                current = candidate
            else:
                lines.append(current)
                current = char
        if current:
            lines.append(current)
        return lines[:5]

    def _short_video_scene_text(self, text: str) -> str:
        text = re.sub(r"\s+", " ", str(text or "")).strip()
        limit = 74
        if len(text) <= limit:
            return text
        end = max(text.rfind("。", 0, limit), text.rfind("、", 0, limit), text.rfind("！", 0, limit), text.rfind("？", 0, limit))
        if end >= 36:
            value = text[: end + 1].rstrip()
            return value if value.endswith(("。", "！", "？")) else value + "..."
        return text[: limit - 3].rstrip() + "..."

    def _comic_panel_caption(self, text: str) -> str:
        text = re.sub(r"\s+", " ", str(text or "")).strip(" 「」")
        if not text:
            return ""
        subject_match = re.match(r"^([ぁ-んァ-ン一-龥A-Za-z0-9ー]{1,8})が", text)
        subject = subject_match.group(1) if subject_match else ""
        if subject and "受付" in text and ("着いた" in text or "到着" in text):
            return f"{subject}、受付で異変。"
        if subject and "笑" in text and ("浮" in text or "重力" in text):
            return f"{subject}、笑って浮く。"
        if "会議" in text and ("椅子" in text or "浮" in text):
            return "会議室、全員ふわふわ。"
        if "ベル" in text and ("鳴" in text or "浮" in text):
            return "ベルが不穏に鳴る。"
        if "禁止" in text and "笑" in text:
            return "笑ったらアウト。"
        for separator in ("。", "！", "？"):
            index = text.find(separator)
            if 8 <= index <= 24:
                text = text[: index + 1]
                break
        replacements = [
            ("という", "って"),
            ("していた", "した"),
            ("している", "してる"),
            ("だった", "だ"),
        ]
        for before, after in replacements:
            text = text.replace(before, after)
        limit = 22
        if len(text) <= limit:
            return text
        end = max(text.rfind("、", 0, limit), text.rfind("。", 0, limit), text.rfind("！", 0, limit), text.rfind("？", 0, limit))
        if end >= 10:
            return text[:end].rstrip("、。！？") + "..."
        return text[: limit - 3].rstrip() + "..."

    def _normalize_short_comic_caption(self, text: str) -> str:
        text = re.sub(r"\s+", " ", str(text or "")).strip(" 「」")
        if len(text) <= 24:
            return text
        end = max(text.rfind("、", 0, 24), text.rfind("。", 0, 24), text.rfind("！", 0, 24), text.rfind("？", 0, 24))
        if end >= 10:
            return text[: end + 1].rstrip()
        return text[:21].rstrip() + "..."

    def _main_character_name_from_payload(self, project_id: int, payload: dict) -> str:
        raw_name = str(payload.get("main_character") or "").strip()
        if raw_name:
            return raw_name
        try:
            character_id = int(payload.get("main_character_id") or 0)
        except (TypeError, ValueError):
            character_id = 0
        if character_id:
            character = Character.query.filter(
                Character.id == character_id,
                Character.project_id == project_id,
                Character.deleted_at.is_(None),
            ).first()
            if character:
                return str(character.name or character.nickname or "").strip()
        return ""

    def _normalize_comic_page_caption(self, text: str) -> str:
        text = re.sub(r"\s+", " ", str(text or "")).strip(" 「」")
        if len(text) <= 36:
            return text
        end = max(text.rfind("、", 0, 36), text.rfind("。", 0, 36), text.rfind("！", 0, 36), text.rfind("？", 0, 36))
        if end >= 12:
            return text[: end + 1].rstrip()
        return text[:33].rstrip() + "..."

    def _short_comic_image_prompt(self, novel, scene: dict, caption: str, *, art_style: str = "reference_image") -> str:
        if str(scene.get("type") or "").strip() == "manual_short_video":
            return self._manual_short_video_image_prompt(novel, scene, caption, art_style=art_style)
        tone = str(scene.get("tone") or "").strip() or "dramatic"
        emotion = str(scene.get("emotion") or "").strip()
        shot = str(scene.get("shot") or "").strip() or "thumbnail composition"
        visual_focus = str(scene.get("visual_focus") or "").strip()
        base_prompt = str(scene.get("image_prompt") or "").strip()
        viewer_context = str(scene.get("viewer_context") or "").strip()
        visible_event = str(scene.get("visible_event") or "").strip()
        joke = str(scene.get("joke") or "").strip()
        story_function = str(scene.get("story_function") or "").strip()
        characters = "、".join(str(item) for item in scene.get("characters") or [] if str(item).strip())
        character_contexts = self._scene_character_contexts(novel.project_id, scene)
        is_mobile = bool(getattr(novel, "mobile_visible", True))
        if art_style == "anime":
            aspect_instruction = (
                "9:16 vertical smartphone manga panel. Keep the full composition inside a tall phone frame; do not crop important faces, bodies, props, or headline text."
                if is_mobile
                else "16:9 horizontal desktop manga panel. Use a wide cinematic composition; keep the full composition inside the frame and do not crop important faces, props, or headline text."
            )
            composition_instruction = "Use an eye-catching manga thumbnail composition, high readability, cinematic lighting, and clear character acting."
            headline_instruction = "The headline should be bold, clean Japanese lettering with strong contrast, like a viral short manga thumbnail."
        else:
            aspect_instruction = (
                "9:16 vertical smartphone visual still in the exact reference-image style. Keep the full composition inside a tall phone frame; do not crop important faces, bodies, props, or headline text."
                if is_mobile
                else "16:9 horizontal desktop visual still in the exact reference-image style. Use a wide cinematic composition; keep the full composition inside the frame and do not crop important faces, props, or headline text."
            )
            composition_instruction = "Use a dramatic social-video still composition while preserving the reference image rendering style. Do not convert the character into anime, manga, cel shading, comic line art, chibi, or a different illustration style."
            headline_instruction = "The headline should be bold, clean Japanese lettering with strong contrast, but the character/background art must remain in the reference-image style."
        style_instruction = self._short_comic_art_style_instruction(art_style)
        text_lines = (
            [
                "Text rendering requirement:",
                f"Add this exact Japanese headline text inside the image, large and readable: 「{caption}」",
                headline_instruction,
                "Do not add any other text, gibberish, watermark, logo, speech bubble text, UI text, or extra captions.",
                "Keep enough contrast behind the headline using a simple dark translucent band or outlined letters if needed.",
            ]
            if caption
            else [
                "Text rendering requirement:",
                "Do not add headline text, captions, subtitles, UI text, gibberish, watermark, logo, or speech bubble text.",
            ]
        )
        return "\n".join(
            [
                f"Create one finished {aspect_instruction}",
                style_instruction,
                composition_instruction,
                f"Novel title: {novel.title or ''}",
                f"Characters in this panel: {characters or 'use the scene context'}",
                "Registered character identity constraints:",
                "Every named character in this panel is a registered character represented by the attached reference images and the settings below.",
                "Do not replace a registered character with a generic boxer, a random man, a stunt double, or a newly invented person.",
                "Preserve each registered character's gender presentation, face, hair, body silhouette, outfit logic, color mood, accessories, and fixed design traits from the reference images/settings.",
                "If the scene requires boxing gloves or action poses, add only those props/poses while keeping the registered character's identity intact.",
                f"Selected character settings: {json_util.dumps(character_contexts) if character_contexts else 'none'}",
                f"Tone: {tone}",
                f"Emotion: {emotion}",
                f"Shot/composition: {shot}",
                f"Story function: {story_function}",
                f"First-time viewer context: {viewer_context}",
                f"Visible event to make obvious: {visible_event}",
                f"Comedic/dramatic gap to show: {joke}",
                f"Visual focus: {visual_focus}",
                "",
                *text_lines,
                "",
                "Scene prompt:",
                base_prompt or visual_focus or caption,
            ]
        )

    def _manual_short_video_image_prompt(self, novel, scene: dict, caption: str, *, art_style: str = "reference_image") -> str:
        is_mobile = bool(getattr(novel, "mobile_visible", True))
        aspect_instruction = (
            "9:16 vertical smartphone visual novel event CG / social-media incident photo"
            if is_mobile
            else "16:9 horizontal cinematic visual novel event CG / social-media incident photo"
        )
        scene_text = str(scene.get("image_prompt") or scene.get("visual_focus") or "").strip()
        characters = ", ".join(str(item) for item in scene.get("characters") or [] if str(item).strip())
        project = Project.query.get(novel.project_id)
        world = World.query.filter_by(project_id=novel.project_id).first()
        character_contexts = self._scene_character_contexts(novel.project_id, scene)
        visual_direction = self._manual_short_video_visual_direction(scene_text, caption)
        scene_brief = self._manual_short_video_scene_brief(scene_text, caption)
        style_instruction = self._short_comic_art_style_instruction(art_style)
        text_lines = (
            [
                "Text overlay requirement:",
                f"Add this exact Japanese headline text inside the image, large and readable: 「{caption}」",
                "Place the headline like a polished short-video caption in a reserved top or bottom band. The headline must be secondary to the event CG composition and must not force a thumbnail-like centered character pose.",
                "Do not cover character faces, hands, food, props, motion, or the key action with text.",
                "Do not add any other text, gibberish, watermark, logo, speech bubble text, UI text, or extra captions.",
            ]
            if caption
            else [
                "Text overlay requirement:",
                "Do not add captions, subtitles, UI text, gibberish, watermark, logo, or speech bubble text.",
            ]
        )
        return "\n".join(
            [
                "Manual short video dynamic mode:",
                "Create a high-impact Feed-like image for a manual short video scene.",
                "Use the attached reference images as the primary source of character identity and art style.",
                "When a selected outfit reference is attached for a character, treat that outfit as the highest-priority clothing reference for that character.",
                "Preserve selected outfit design, silhouette, fabric impression, color palette, accessories, fixed parts, and NG rules unless the scene text explicitly requests a compatible change.",
                "If a character has both an outfit reference and a base identity reference, those references describe the same character; do not create duplicate people from them.",
                "If multiple reference images are provided, include all selected referenced characters in the scene unless the scene text clearly says otherwise.",
                "Keep the same face, hair, outfit design logic, coloring, rendering quality, material detail, and mood for every referenced character.",
                "All visible characters must share one unified high-end 3D, semi-realistic, cinematic game-CG style. Do not render one character as lower-detail anime, chibi, manga, sketch, or flat illustration.",
                style_instruction,
                f"Target frame: {aspect_instruction}.",
                "No speech bubbles, no UI, no logo, no watermark.",
                "Do not make a simple standing portrait, idle pose, catalog pose, or generic promotional still.",
                "The image must depict the user's scene text as one clear visible moment: concrete action, props, facial expression, body language, and readable environment.",
                "Make it feel like a dramatic social-media incident photo or visual novel event CG: caught-in-the-act composition, expressive reaction, visible cause of the situation, and a clear environment.",
                "The characters should be doing something specific: reacting, pointing, stumbling, holding an object, eating, reaching, laughing, inspecting a device, shielding themselves from chaos, or being caught mid-action.",
                "Use dynamic camera language where appropriate: dutch angle, close foreground object, over-the-shoulder framing, shallow depth of field, motion blur, steam, dramatic lighting, cluttered evidence, or a comedic reveal in the background.",
                "Compose like the Feed image generator: event-first, prop/action-first, expressive reaction, visible environment, not a poster pose.",
                f"Primary visual scene brief: {scene_brief}",
                "Treat the primary visual scene brief as the main source of truth for composition, action, props, environment, and emotion.",
                f"Visual direction: {visual_direction}",
                "The viewer should understand the situation from the image alone.",
                f"Short video title: {novel.title or ''}",
                f"Selected reference characters: {characters or 'none explicitly selected'}",
                f"Scene text to visualize: {scene_text or caption}",
                f"World: {getattr(project, 'title', '') or ''}. {getattr(project, 'summary', '') or ''}",
                f"World setting: {(getattr(world, 'overview', '') or '')} {(getattr(world, 'tone', '') or '')}" if world else "World setting: ",
                f"Selected character settings: {json_util.dumps(character_contexts)}",
                "",
                *text_lines,
                "",
                "Safety/style guardrail: keep teasing or romantic-comedy moments tasteful and non-explicit; no nudity, no exposed intimate body parts, no sexual framing.",
            ]
        )

    def _manual_short_video_scene_brief(self, scene_text: str, caption: str = "") -> str:
        text = str(scene_text or "").strip()
        caption_text = str(caption or "").strip()
        if text and caption_text and caption_text not in text:
            return f"{text} / visible headline emotion: {caption_text}"
        if text:
            return text
        if caption_text:
            return caption_text
        return "A clear character incident with visible action, props, readable environment, and expressive reaction."

    def _manual_short_video_visual_direction(self, scene_text: str, caption: str = "") -> str:
        text = f"{scene_text or ''}\n{caption or ''}".lower()
        food_words = ("ラーメン", "料理", "食べ", "飯", "ごはん", "麺", "辛", "甘", "スイーツ", "カフェ", "肉", "ソース", "湯気")
        tangle_words = ("ロープ", "ケーブル", "コード", "リボン", "絡ま", "引っか", "ひっか", "ほどけ")
        fall_words = ("転", "こけ", "つまず", "尻もち", "倒れ", "落ち", "よろけ")
        romance_words = ("照れ", "赤面", "ドキ", "告白", "距離", "手を", "近い", "ラブ", "恋")
        report_words = ("発見", "目撃", "現地", "速報", "事件", "騒ぎ", "トラブル")
        photo_words = ("写真", "撮影", "ポーズ", "自撮り", "映り", "写り")
        if any(word in text for word in food_words):
            return (
                "Food reaction event CG. Put the food or ramen bowl large in the foreground, with steam/sauce/chopsticks as dynamic foreground elements. "
                "Show the selected characters leaning into the food, eating or reacting mid-action, with expressive faces and a readable restaurant/cafe environment. "
                "Avoid a plain seated portrait; make the food, hands, steam, and reaction drive the composition."
            )
        if any(word in text for word in tangle_words):
            return (
                "Tangled prop mishap event CG. Show the rope/cable/ribbon visibly caught on clothing, hands, chair, mic, or nearby fixture. "
                "Use diagonal lines and awkward body language to show motion and confusion; keep it tasteful and non-explicit."
            )
        if any(word in text for word in fall_words):
            return (
                "Near-fall or recovery moment. Capture the instant of stumbling or regaining balance: tilted camera, reaching hand, flying prop, surprised expression, and environment clues."
            )
        if any(word in text for word in romance_words):
            return (
                "Small romantic-comedy accident as visual evidence. Use flustered expression, awkward distance, hands/eye contact, warm dramatic lighting, and one concrete prop that explains the moment."
            )
        if any(word in text for word in report_words):
            return (
                "Live field-report incident composition. Put one character in the foreground reacting, with the concrete incident unfolding behind or beside them. Make the cause visible."
            )
        if any(word in text for word in photo_words):
            return (
                "Failed photo / behind-the-scenes shot. The character is trying for a nice pose, but a funny background detail, prop, timing mistake, or expression ruins it."
            )
        return (
            "Feed-style event CG. Choose the most imageable beat from the scene text and show one clear action, one prop or environmental clue, and one strong emotion. "
            "Use an off-center composition, dynamic camera angle, foreground object, and expressive body language so it does not look like a character sheet."
        )

    def _short_comic_thumbnail_prompt(self, novel, headline: str, panel: dict, *, main_character: str = "", logline: str = "", art_style: str = "reference_image") -> str:
        is_mobile = bool(getattr(novel, "mobile_visible", True))
        aspect_instruction = (
            "9:16 vertical smartphone thumbnail"
            if is_mobile
            else "16:9 horizontal desktop thumbnail"
        )
        visual_focus = str(panel.get("visual_focus") or panel.get("image_prompt") or logline or "").strip()
        characters = "、".join(str(item) for item in panel.get("characters") or [] if str(item).strip())
        if main_character and main_character not in characters:
            characters = "、".join([main_character, characters]).strip("、")
        style_instruction = self._short_comic_art_style_instruction(art_style)
        if art_style == "anime":
            series_label = "short manga video series"
            design_instruction = "Use bold manga thumbnail design, dramatic composition, clear focal character, strong contrast, and polished lighting."
        else:
            series_label = "short visual video series"
            design_instruction = "Make it clickable with dramatic composition, clear focal character, strong contrast, and polished lighting, while preserving the exact reference-image rendering style. Do not convert the art into anime, manga, cel shading, or comic line art."
        return "\n".join(
            [
                f"Create a finished {aspect_instruction} for a {series_label}.",
                "This is the standalone opening thumbnail, not panel 1.",
                style_instruction,
                "Make it more iconic, poster-like, and clickable than an ordinary story panel.",
                design_instruction,
                f"Novel title: {novel.title or ''}",
                f"Main/visible characters: {characters or main_character or 'use the story context'}",
                f"Story hook: {logline}",
                f"Visual focus: {visual_focus}",
                "",
                "Text rendering requirement:",
                f"Add this exact Japanese headline text inside the image, large and readable: 「{headline}」",
                "Use clean bold Japanese lettering. Do not add any other text, gibberish, watermark, logo, speech bubbles, UI text, or extra captions.",
                "Keep all important faces, props, and headline text fully inside the frame.",
            ]
        )

    def _comic_page_image_prompt(self, novel, scene: dict, *, art_style: str = "reference_image") -> str:
        page_title = str(scene.get("page_title") or scene.get("caption") or "Manga page").strip()
        summary = str(scene.get("visual_focus") or scene.get("text") or "").strip()
        tone = str(scene.get("tone") or "").strip() or "comedy"
        characters = ", ".join(str(item) for item in scene.get("characters") or [] if str(item).strip())
        panels = scene.get("page_panels") if isinstance(scene.get("page_panels"), list) else []
        panel_lines = []
        exact_texts = []
        for index, panel in enumerate(panels[:6], start=1):
            if not isinstance(panel, dict):
                continue
            speaker = str(panel.get("speaker") or "").strip()
            speech = str(panel.get("speech") or panel.get("dialogue") or "").strip()
            sfx = str(panel.get("sfx") or "").strip()
            if speech:
                exact_texts.append(speech)
            if sfx:
                exact_texts.append(sfx)
            panel_lines.append(
                f"Panel {index}: composition={panel.get('composition') or ''}; action={panel.get('action') or ''}; "
                f"speaker={speaker}; speech={speech}; sfx={sfx}; background={panel.get('background') or ''}"
            )
        if not panel_lines:
            panel_lines.append(f"Panel 1: introduce {characters or 'the protagonist'}; speech={page_title}")
        style_instruction = self._short_comic_art_style_instruction(art_style)
        identity_instruction = (
            "Preserve the reference character identity, costume impression, face, hair, color mood, and rendering taste. Use manga panel layout and speech balloons, but do not redesign the character into a different person."
            if art_style == "reference_image"
            else "Use polished modern manga/anime page art with expressive faces, clean linework, readable balloons, and dynamic panel rhythm."
        )
        return "\n".join(
            [
                "Create one complete finished Japanese manga page image.",
                "This is a normal manga page, not a social-video thumbnail and not a single illustration.",
                style_instruction,
                identity_instruction,
                "Page layout: 2 to 6 panels with clear white gutters, readable panel order, varied close-up / medium / wide shots.",
                "Add speech balloons and optional small narration boxes directly inside the page.",
                "Use only the exact Japanese text listed below. Do not add gibberish, watermark, UI text, random signs, or extra captions.",
                "Keep every speech balloon large and readable on a smartphone. Avoid tiny dense text.",
                "The page should feel like a real manga page: panel borders, balloons, expressive reactions, timing, and visual comedy/drama.",
                f"Novel title: {novel.title or ''}",
                f"Page title: {page_title}",
                f"Tone: {tone}",
                f"Characters: {characters or 'use the story context'}",
                f"Page summary: {summary}",
                "",
                "Panel plan:",
                "\n".join(panel_lines),
                "",
                "Exact visible text to render:",
                "\n".join(f"- {text}" for text in exact_texts[:12]) or f"- {page_title}",
                "",
                f"Base image prompt: {scene.get('image_prompt') or ''}",
            ]
        )

    def _short_comic_end_card_prompt(self, novel, panel: dict, *, main_character: str = "", art_style: str = "reference_image") -> str:
        is_mobile = bool(getattr(novel, "mobile_visible", True))
        aspect_instruction = (
            "9:16 vertical smartphone end card"
            if is_mobile
            else "16:9 horizontal desktop end card"
        )
        visual_focus = str(panel.get("visual_focus") or panel.get("image_prompt") or "").strip()
        characters = "、".join(str(item) for item in panel.get("characters") or [] if str(item).strip())
        if main_character and main_character not in characters:
            characters = "、".join([main_character, characters]).strip("、")
        style_instruction = self._short_comic_art_style_instruction(art_style)
        finish_instruction = (
            "Use dramatic lighting, confident pose, strong silhouette, and a premium cinematic anime finish."
            if art_style == "anime"
            else "Use dramatic lighting, confident pose, strong silhouette, and a premium cinematic finish while preserving the exact reference-image rendering style. Do not convert the art into anime, manga, cel shading, or comic line art."
        )
        return "\n".join(
            [
                f"Create a finished {aspect_instruction} for the final frame of a short visual video.",
                "This is an end card, not a story panel.",
                style_instruction,
                "Make the main protagonist look extremely cool, heroic, iconic, and poster-like.",
                finish_instruction,
                f"Novel title: {novel.title or ''}",
                f"Main/visible characters: {characters or main_character or 'use the story context'}",
                f"Story context: {visual_focus}",
                "",
                "Text rendering requirement:",
                'Add exact logo text: "LAPLACE CITY"',
                'Add small text below or nearby: "END"',
                "Do not add any other text, gibberish, watermark, logo, speech bubbles, UI text, or extra captions.",
                "Keep all important faces, props, LAPLACE CITY, and END fully inside the frame.",
            ]
        )

    def _short_comic_art_style(self, payload: dict | None) -> str:
        value = str((payload or {}).get("short_comic_art_style") or "reference_image").strip().lower()
        return "anime" if value in {"anime", "animation", "アニメ調"} else "reference_image"

    def _short_comic_layout(self, payload: dict | None) -> str:
        value = str((payload or {}).get("comic_layout") or "short").strip().lower()
        return "page" if value in {"page", "pages", "manga_page", "comic_page", "漫画ページ"} else "short"

    def _short_comic_art_style_instruction(self, art_style: str) -> str:
        if art_style == "anime":
            return "Art style: vivid modern anime and manga thumbnail style, clean cel shading, expressive faces, readable shapes, polished commercial animation look."
        return "Art style: follow the attached/reference character images as the primary visual standard. Preserve their exact rendering style, medium, texture, line/paint treatment, costume impression, facial identity, color mood, and rendering taste. Do not anime-ify, manga-ify, cel-shade, chibi-fy, or redesign the character."

    def _set_powerpoint_shape_alpha(self, pptx_path: Path, color_hex: str, alpha_value: int):
        from zipfile import ZIP_DEFLATED, ZipFile

        source = Path(pptx_path)
        temp_path = source.with_suffix(".tmp.pptx")
        color_hex = str(color_hex or "").upper()
        alpha_value = max(0, min(100000, int(alpha_value)))
        target = f'<a:srgbClr val="{color_hex}"/>'
        replacement = f'<a:srgbClr val="{color_hex}"><a:alpha val="{alpha_value}"/></a:srgbClr>'

        with ZipFile(source, "r") as zin, ZipFile(temp_path, "w", ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename.startswith("ppt/slides/slide") and item.filename.endswith(".xml"):
                    text = data.decode("utf-8")
                    text = text.replace(target, replacement)
                    data = text.encode("utf-8")
                zout.writestr(item, data)
        os.replace(temp_path, source)

    def _apply_bgm_to_video(self, video_path: Path, project_id: int, payload: dict):
        bgm_asset_id = int(payload.get("bgm_asset_id") or 0)
        if not bgm_asset_id:
            return
        asset = self._asset_service.get_asset(bgm_asset_id)
        if not asset or int(asset.project_id or 0) != int(project_id):
            raise ValueError("selected BGM was not found")
        if asset.asset_type != "cinema_novel_bgm":
            raise ValueError("selected asset is not a BGM file")
        audio_path = Path(asset.file_path)
        if not audio_path.exists():
            raise ValueError("selected BGM file was not found")
        try:
            volume = float(payload.get("bgm_volume") or 0.45)
        except (TypeError, ValueError):
            volume = 0.45
        volume = max(0.0, min(1.0, volume))
        video_duration = self._probe_media_duration(video_path)
        audio_filter = f"volume={volume:.3f}"
        if video_duration and video_duration > 2.5:
            fade_start = max(0.0, video_duration - 2.0)
            audio_filter = f"{audio_filter},afade=t=out:st={fade_start:.2f}:d=2"
        mixed_path = video_path.with_name(f"{video_path.stem}_bgm{video_path.suffix}")
        command = [
            "ffmpeg",
            "-y",
            "-i",
            str(video_path),
            "-stream_loop",
            "-1",
            "-i",
            str(audio_path),
            "-map",
            "0:v:0",
            "-map",
            "1:a:0",
            "-c:v",
            "copy",
            "-c:a",
            "aac",
            "-b:a",
            "160k",
            "-filter:a",
            audio_filter,
            "-shortest",
            "-movflags",
            "+faststart",
            str(mixed_path),
        ]
        result = subprocess.run(command, capture_output=True, text=True, cwd=str(Path(current_app.root_path).parent))
        if result.returncode != 0:
            raise RuntimeError((result.stderr or result.stdout or "ffmpeg audio mix failed").strip()[-1000:])
        os.replace(mixed_path, video_path)

    def _probe_media_duration(self, media_path: Path) -> float | None:
        command = [
            "ffprobe",
            "-v",
            "error",
            "-show_entries",
            "format=duration",
            "-of",
            "default=noprint_wrappers=1:nokey=1",
            str(media_path),
        ]
        try:
            result = subprocess.run(command, capture_output=True, text=True, cwd=str(Path(current_app.root_path).parent))
            if result.returncode != 0:
                return None
            return float(str(result.stdout or "").strip())
        except Exception:
            return None

    def _apply_mobile_novel_image_options(self, novel, options: dict) -> dict:
        options = dict(options or {})
        if bool(getattr(novel, "mobile_visible", False)):
            settings = self._user_setting_service.get_global_settings()
            options["size"] = settings.get("mobile_default_size") or "1024x1536"
            options["client_viewport"] = "portrait_mobile"
        return options

    def _novel_image_layout_instruction(self, novel) -> str:
        if bool(getattr(novel, "mobile_visible", False)):
            return "Portrait mobile composition, 9:16 key visual. Keep main characters and the background naturally framed for a tall phone screen."
        return "Landscape cinematic composition, 16:9 key visual."

    def delete_novel(self, novel_id: int) -> bool:
        novel = self.get_novel(novel_id)
        if not novel:
            return False
        now = datetime.utcnow()
        novel.deleted_at = now
        novel.status = "archived"
        for chapter in CinemaNovelChapter.query.filter(CinemaNovelChapter.novel_id == novel.id).all():
            chapter.deleted_at = now
            db.session.add(chapter)
        for entry in CinemaNovelLoreEntry.query.filter(CinemaNovelLoreEntry.novel_id == novel.id).all():
            entry.deleted_at = now
            db.session.add(entry)
        reviews = CinemaNovelReview.query.filter(CinemaNovelReview.novel_id == novel.id).all()
        feed_post_ids = []
        memory_note_ids = []
        for review in reviews:
            review.deleted_at = now
            if review.feed_post_id:
                feed_post_ids.append(review.feed_post_id)
            if review.memory_note_id:
                memory_note_ids.append(review.memory_note_id)
            db.session.add(review)
        for impression in CinemaNovelCharacterImpression.query.filter(CinemaNovelCharacterImpression.novel_id == novel.id).all():
            impression.deleted_at = now
            if impression.memory_note_id:
                memory_note_ids.append(impression.memory_note_id)
            db.session.add(impression)
        if feed_post_ids:
            for post in FeedPost.query.filter(FeedPost.id.in_(feed_post_ids)).all():
                post.deleted_at = now
                post.status = "archived"
                db.session.add(post)
        source_refs = [f"cinema_novel:{novel.id}", f"cinema_novel:{novel.id}:impressions"]
        note_query = CharacterMemoryNote.query.filter(
            (CharacterMemoryNote.source_ref.in_(source_refs))
            | (CharacterMemoryNote.id.in_(memory_note_ids or [-1]))
        )
        for note in note_query.all():
            note.enabled = False
            db.session.add(note)
        CinemaNovelProgress.query.filter(CinemaNovelProgress.novel_id == novel.id).delete(synchronize_session=False)
        db.session.add(novel)
        db.session.commit()
        return True

    def list_chapters(self, novel_id: int):
        return CinemaNovelChapter.query.filter(
            CinemaNovelChapter.novel_id == novel_id,
            CinemaNovelChapter.deleted_at.is_(None),
        ).order_by(CinemaNovelChapter.chapter_no.asc(), CinemaNovelChapter.sort_order.asc(), CinemaNovelChapter.id.asc()).all()

    def get_chapter(self, chapter_id: int):
        return CinemaNovelChapter.query.filter(
            CinemaNovelChapter.id == chapter_id,
            CinemaNovelChapter.deleted_at.is_(None),
        ).first()

    def serialize_novel(self, novel, *, include_chapters: bool = False, user_id: int | None = None):
        if not novel:
            return None
        production = self._load_json(novel.production_json)
        if not isinstance(production, dict):
            production = {}
        end_card_asset_id = production.get("end_card_asset_id")
        cover_history_ids = [
            asset_id for asset_id in self._asset_history_ids(production, "cover_asset_history_ids")
            if int(asset_id) != int((novel.poster_asset_id or novel.cover_asset_id) or 0)
        ]
        end_card_history_ids = [
            asset_id for asset_id in self._asset_history_ids(production, "end_card_asset_history_ids")
            if int(asset_id) != int(end_card_asset_id or 0)
        ]
        payload = {
            "id": novel.id,
            "project_id": novel.project_id,
            "created_by_user_id": novel.created_by_user_id,
            "title": novel.title,
            "subtitle": novel.subtitle,
            "description": novel.description,
            "status": novel.status,
            "mobile_visible": bool(getattr(novel, "mobile_visible", True)),
            "mode": novel.mode,
            "cover_asset_id": novel.cover_asset_id,
            "poster_asset_id": novel.poster_asset_id,
            "cover_asset": self._serialize_asset(novel.cover_asset_id),
            "poster_asset": self._serialize_asset(novel.poster_asset_id),
            "source_path": novel.source_path,
            "production_json": production,
            "comic_cover_asset_history": [
                asset for asset in (self._serialize_asset(asset_id) for asset_id in cover_history_ids)
                if asset
            ],
            "comic_end_card_asset_id": end_card_asset_id,
            "comic_end_card_asset": self._serialize_asset(end_card_asset_id),
            "comic_end_card_asset_history": [
                asset for asset in (self._serialize_asset(asset_id) for asset_id in end_card_history_ids)
                if asset
            ],
            "sort_order": novel.sort_order,
            "created_at": novel.created_at.isoformat() if novel.created_at else None,
            "updated_at": novel.updated_at.isoformat() if novel.updated_at else None,
        }
        chapters = self.list_chapters(novel.id)
        payload["chapter_count"] = len(chapters)
        if include_chapters:
            payload["chapters"] = [self.serialize_chapter(chapter) for chapter in chapters]
        if user_id:
            payload["progress"] = self.serialize_progress(self.get_progress(user_id, novel.id))
            payload["reviews"] = [self.serialize_review(review) for review in self.list_reviews(novel.id, user_id=user_id)]
        else:
            payload["reviews"] = [self.serialize_review(review) for review in self.list_reviews(novel.id)]
        payload["lore_entries"] = [self.serialize_lore_entry(entry) for entry in self.list_lore_entries(novel.id)]
        return payload

    def list_reviews(self, novel_id: int, *, user_id: int | None = None):
        query = CinemaNovelReview.query.filter(
            CinemaNovelReview.novel_id == novel_id,
            CinemaNovelReview.deleted_at.is_(None),
        )
        if user_id:
            query = query.filter(CinemaNovelReview.user_id == user_id)
        return query.order_by(CinemaNovelReview.updated_at.desc(), CinemaNovelReview.id.desc()).all()

    def serialize_review(self, review) -> dict | None:
        if not review:
            return None
        character = Character.query.get(review.character_id)
        thumbnail_id = getattr(character, "thumbnail_asset_id", None) or getattr(character, "base_asset_id", None)
        return {
            "id": review.id,
            "novel_id": review.novel_id,
            "character_id": review.character_id,
            "user_id": review.user_id,
            "feed_post_id": review.feed_post_id,
            "memory_note_id": review.memory_note_id,
            "review_text": review.review_text,
            "memory_note": review.memory_note,
            "rating_label": review.rating_label,
            "status": review.status,
            "metadata": self._load_json(review.metadata_json, default={}),
            "impressions": [
                self.serialize_character_impression(impression)
                for impression in self.list_character_impressions(
                    review.novel_id,
                    reviewer_character_id=review.character_id,
                    user_id=review.user_id,
                )
            ],
            "character": {
                "id": character.id,
                "name": character.name,
                "nickname": character.nickname,
                "thumbnail_asset": self._serialize_asset(thumbnail_id),
            } if character else None,
            "created_at": review.created_at.isoformat() if review.created_at else None,
            "updated_at": review.updated_at.isoformat() if review.updated_at else None,
        }

    def create_character_review(self, novel_id: int, user_id: int, payload: dict | None = None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        try:
            character_id = int(payload.get("character_id") or 0)
        except (TypeError, ValueError):
            raise ValueError("character_id is required")
        character = Character.query.filter(
            Character.id == character_id,
            Character.project_id == novel.project_id,
            Character.deleted_at.is_(None),
        ).first()
        if not character:
            raise ValueError("character was not found")
        existing = CinemaNovelReview.query.filter(
            CinemaNovelReview.novel_id == novel.id,
            CinemaNovelReview.character_id == character.id,
            CinemaNovelReview.user_id == user_id,
            CinemaNovelReview.deleted_at.is_(None),
        ).first()
        lore_entries = self.ensure_novel_lore(novel.id)
        result = self._generate_character_review(novel, character)
        review_text = str(result.get("feed_review") or "").strip()
        memory_note_text = str(result.get("memory_note") or "").strip()
        rating_label = str(result.get("rating_label") or "").strip()[:80] or None
        if not review_text:
            raise RuntimeError("review response did not include review text")
        if not memory_note_text:
            memory_note_text = self._fallback_review_memory_note(novel, character, review_text)
        feed_post = self._upsert_review_feed_post(
            novel=novel,
            character=character,
            user_id=user_id,
            review_text=review_text,
            existing_feed_post_id=existing.feed_post_id if existing else None,
        )
        memory_note = self._upsert_review_memory_note(
            novel=novel,
            character=character,
            user_id=user_id,
            note_text=memory_note_text,
            existing_memory_note_id=existing.memory_note_id if existing else None,
        )
        impressions = self._generate_and_upsert_character_impressions(
            novel=novel,
            character=character,
            user_id=user_id,
            lore_entries=lore_entries,
        )
        impression_memory = self._upsert_impression_memory_note(
            novel=novel,
            character=character,
            user_id=user_id,
            impressions=impressions,
        )
        if impression_memory:
            for impression in impressions:
                impression.memory_note_id = impression_memory.id
            db.session.commit()
        metadata = {
            "source": "cinema_novel_review",
            "model": result.get("model"),
            "usage": result.get("usage"),
            "review_summary": result.get("review_summary"),
            "lore_entry_count": len(lore_entries or []),
            "impression_count": len(impressions or []),
        }
        if existing:
            existing.feed_post_id = feed_post.id if feed_post else None
            existing.memory_note_id = memory_note.id if memory_note else None
            existing.review_text = review_text
            existing.memory_note = memory_note_text
            existing.rating_label = rating_label
            existing.status = "published"
            existing.metadata_json = json_util.dumps(metadata)
            db.session.commit()
            return existing
        review = CinemaNovelReview(
            novel_id=novel.id,
            character_id=character.id,
            user_id=user_id,
            feed_post_id=feed_post.id if feed_post else None,
            memory_note_id=memory_note.id if memory_note else None,
            review_text=review_text,
            memory_note=memory_note_text,
            rating_label=rating_label,
            status="published",
            metadata_json=json_util.dumps(metadata),
        )
        db.session.add(review)
        db.session.commit()
        return review

    def list_lore_entries(self, novel_id: int):
        return CinemaNovelLoreEntry.query.filter(
            CinemaNovelLoreEntry.novel_id == novel_id,
            CinemaNovelLoreEntry.deleted_at.is_(None),
        ).order_by(CinemaNovelLoreEntry.sort_order.asc(), CinemaNovelLoreEntry.id.asc()).all()

    def serialize_lore_entry(self, entry) -> dict | None:
        if not entry:
            return None
        return {
            "id": entry.id,
            "novel_id": entry.novel_id,
            "lore_type": entry.lore_type,
            "name": entry.name,
            "summary": entry.summary,
            "role_note": entry.role_note,
            "source_note": entry.source_note,
            "sort_order": entry.sort_order,
            "metadata": self._load_json(entry.metadata_json, default={}),
            "created_at": entry.created_at.isoformat() if entry.created_at else None,
            "updated_at": entry.updated_at.isoformat() if entry.updated_at else None,
        }

    def list_character_impressions(
        self,
        novel_id: int,
        *,
        reviewer_character_id: int | None = None,
        user_id: int | None = None,
    ):
        query = CinemaNovelCharacterImpression.query.filter(
            CinemaNovelCharacterImpression.novel_id == novel_id,
            CinemaNovelCharacterImpression.deleted_at.is_(None),
        )
        if reviewer_character_id:
            query = query.filter(CinemaNovelCharacterImpression.reviewer_character_id == reviewer_character_id)
        if user_id:
            query = query.filter(CinemaNovelCharacterImpression.user_id == user_id)
        return query.order_by(CinemaNovelCharacterImpression.id.asc()).all()

    def serialize_character_impression(self, impression) -> dict | None:
        if not impression:
            return None
        target = Character.query.get(impression.target_character_id) if impression.target_character_id else None
        return {
            "id": impression.id,
            "novel_id": impression.novel_id,
            "reviewer_character_id": impression.reviewer_character_id,
            "user_id": impression.user_id,
            "target_name": impression.target_name,
            "target_character_id": impression.target_character_id,
            "target_character": {
                "id": target.id,
                "name": target.name,
                "nickname": target.nickname,
            } if target else None,
            "impression_text": impression.impression_text,
            "talk_hint": impression.talk_hint,
            "memory_note_id": impression.memory_note_id,
            "metadata": self._load_json(impression.metadata_json, default={}),
            "created_at": impression.created_at.isoformat() if impression.created_at else None,
            "updated_at": impression.updated_at.isoformat() if impression.updated_at else None,
        }

    def ensure_novel_lore(self, novel_id: int, *, force: bool = False):
        novel = self.get_novel(novel_id)
        if not novel:
            return []
        existing = self.list_lore_entries(novel.id)
        if existing and not force:
            return existing
        result = self._generate_novel_lore(novel)
        entries = result.get("entries") if isinstance(result, dict) else []
        if not isinstance(entries, list):
            entries = []
        saved = []
        for index, item in enumerate(entries[:40], start=1):
            if not isinstance(item, dict):
                continue
            name = str(item.get("name") or "").strip()[:255]
            summary = str(item.get("summary") or "").strip()
            if not name or not summary:
                continue
            lore_type = str(item.get("lore_type") or "other").strip().lower()[:50] or "other"
            entry = CinemaNovelLoreEntry.query.filter(
                CinemaNovelLoreEntry.novel_id == novel.id,
                CinemaNovelLoreEntry.lore_type == lore_type,
                CinemaNovelLoreEntry.name == name,
            ).first()
            if not entry:
                entry = CinemaNovelLoreEntry(novel_id=novel.id, lore_type=lore_type, name=name)
            entry.summary = summary[:1600]
            entry.role_note = str(item.get("role_note") or "").strip()[:1200] or None
            entry.source_note = str(item.get("source_note") or "").strip()[:1200] or None
            entry.sort_order = index
            entry.metadata_json = json_util.dumps(
                {
                    "source": "cinema_novel_lore_generation",
                    "model": result.get("model"),
                    "usage": result.get("usage"),
                }
            )
            entry.deleted_at = None
            db.session.add(entry)
            saved.append(entry)
        db.session.commit()
        return self.list_lore_entries(novel.id)

    def serialize_chapter(self, chapter):
        if not chapter:
            return None
        scenes = self._load_json(chapter.scene_json, default=[])
        if isinstance(scenes, list):
            scenes = [self._serialize_scene(scene) for scene in scenes]
        return {
            "id": chapter.id,
            "novel_id": chapter.novel_id,
            "chapter_no": chapter.chapter_no,
            "title": chapter.title,
            "body_markdown": chapter.body_markdown,
            "scene_json": scenes,
            "cover_asset_id": chapter.cover_asset_id,
            "cover_asset": self._serialize_asset(chapter.cover_asset_id),
            "generated_assets": self._chapter_generated_assets(chapter),
            "sort_order": chapter.sort_order,
            "created_at": chapter.created_at.isoformat() if chapter.created_at else None,
            "updated_at": chapter.updated_at.isoformat() if chapter.updated_at else None,
        }

    def _serialize_scene(self, scene):
        if not isinstance(scene, dict):
            return scene
        payload = dict(scene)
        payload["background_asset"] = self._serialize_asset(payload.get("background_asset_id"))
        payload["still_asset"] = self._serialize_asset(payload.get("still_asset_id"))
        history_ids = [
            asset_id for asset_id in self._scene_still_history_ids(payload)
            if int(asset_id) != int(payload.get("still_asset_id") or 0)
        ]
        payload["still_asset_history_ids"] = history_ids
        payload["still_asset_history"] = [
            asset for asset in (self._serialize_asset(asset_id) for asset_id in history_ids)
            if asset
        ]
        return payload

    def _chapter_generated_assets(self, chapter):
        novel = self.get_novel(chapter.novel_id)
        if not novel:
            return []
        assets = Asset.query.filter(
            Asset.project_id == novel.project_id,
            Asset.deleted_at.is_(None),
            Asset.asset_type.in_(["cinema_novel_chapter_cover", "cinema_novel_scene_still"]),
        ).order_by(Asset.created_at.asc(), Asset.id.asc()).all()
        payload = []
        seen_asset_ids = set()
        for asset in assets:
            metadata = self._load_json(asset.metadata_json, default={})
            if not isinstance(metadata, dict) or int(metadata.get("chapter_id") or 0) != int(chapter.id):
                continue
            serialized = self._serialize_asset(asset.id)
            if not serialized or asset.id in seen_asset_ids:
                continue
            seen_asset_ids.add(asset.id)
            source = metadata.get("source") or asset.asset_type
            scene_index = metadata.get("scene_index")
            if source == "cinema_novel_chapter_cover":
                label = "章扉"
            elif isinstance(scene_index, int):
                label = f"scene {scene_index + 1}"
            else:
                label = "劇中スチル"
            serialized["label"] = label
            serialized["source"] = source
            serialized["scene_index"] = scene_index
            payload.append(serialized)
        return payload

    def get_progress(self, user_id: int, novel_id: int):
        return CinemaNovelProgress.query.filter_by(user_id=user_id, novel_id=novel_id).first()

    def serialize_progress(self, progress):
        if not progress:
            return None
        return {
            "id": progress.id,
            "user_id": progress.user_id,
            "novel_id": progress.novel_id,
            "chapter_id": progress.chapter_id,
            "scene_index": progress.scene_index,
            "page_index": progress.page_index,
            "updated_at": progress.updated_at.isoformat() if progress.updated_at else None,
        }

    def save_progress(self, user_id: int, novel_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        try:
            chapter_id = int(payload.get("chapter_id") or 0)
            scene_index = max(0, int(payload.get("scene_index") or 0))
            page_index = max(0, int(payload.get("page_index") or 0))
        except (TypeError, ValueError):
            raise ValueError("invalid progress")
        chapter = self.get_chapter(chapter_id)
        if not chapter or chapter.novel_id != novel.id:
            raise ValueError("invalid chapter")
        scenes = self._load_json(chapter.scene_json, default=[])
        if scenes:
            scene_index = min(scene_index, len(scenes) - 1)
        progress = self.get_progress(user_id, novel.id)
        if not progress:
            progress = CinemaNovelProgress(user_id=user_id, novel_id=novel.id, chapter_id=chapter.id)
            db.session.add(progress)
        progress.chapter_id = chapter.id
        progress.scene_index = scene_index
        progress.page_index = page_index
        db.session.commit()
        return progress

    def import_markdown_folder(self, project_id: int, user_id: int, payload: dict | None):
        payload = dict(payload or {})
        source_path = str(payload.get("source_path") or "").strip()
        if not source_path:
            raise ValueError("source_path is required")
        folder = self._resolve_book_folder(source_path)
        chapter_files = self._chapter_files(folder)
        if not chapter_files:
            raise ValueError("chapter markdown files were not found")
        title = str(payload.get("title") or folder.name).strip() or folder.name
        existing = CinemaNovel.query.filter(
            CinemaNovel.project_id == project_id,
            CinemaNovel.title == title,
            CinemaNovel.source_path == str(folder),
            CinemaNovel.deleted_at.is_(None),
        ).first()
        if existing:
            return existing
        novel = CinemaNovel(
            project_id=project_id,
            created_by_user_id=user_id,
            title=title,
            subtitle=str(payload.get("subtitle") or "ノベル作品").strip() or None,
            description=str(payload.get("description") or "").strip() or None,
            status=str(payload.get("status") or "published").strip() or "published",
            mobile_visible=self._normalize_bool(payload.get("mobile_visible", True)),
            mode="cinema_novel",
            source_path=str(folder),
            production_json=json_util.dumps(
                {
                    "source": "markdown_folder",
                    "reader": "prebuilt",
                    "generation_mode": "build_then_publish",
                    "image_generation": "prebuilt_only",
                }
            ),
        )
        if novel.status not in self.VALID_STATUSES:
            novel.status = "draft"
        db.session.add(novel)
        db.session.flush()
        for index, path in enumerate(chapter_files, start=1):
            body = path.read_text(encoding="utf-8").strip()
            chapter_no, chapter_title = self._extract_chapter_heading(path, body, index)
            db.session.add(
                CinemaNovelChapter(
                    novel_id=novel.id,
                    chapter_no=chapter_no,
                    title=chapter_title,
                    body_markdown=body,
                    scene_json=json_util.dumps(self._markdown_to_scenes(body, chapter_no=chapter_no)),
                    sort_order=index,
                )
            )
        db.session.commit()
        return novel

    def generate_production_outline(self, project_id: int, payload: dict | None):
        payload = dict(payload or {})
        settings = self._user_setting_service.apply_cinema_novel_text_generation_settings(payload.get("text_options") or {})
        title = str(payload.get("title") or "無題のノベル作品").strip()
        main_character = str(payload.get("main_character") or "").strip()
        theme = str(payload.get("theme") or "").strip()
        genre = str(payload.get("genre") or "映画ノベル").strip()
        chapter_count = int(payload.get("chapter_count") or 5)
        chapter_count = max(3, min(12, chapter_count))
        chapter_target_chars = int(payload.get("chapter_target_chars") or settings.get("chapter_target_chars") or 3500)
        project = Project.query.get(project_id)
        world = World.query.filter_by(project_id=project_id).first()
        reference_sources = self._reference_sources(payload)
        registered_character_context = (
            self._registered_character_context(project_id, main_character=main_character)
            if self._reference_source_enabled(reference_sources, "characters")
            else ""
        )
        reference_context = self._production_reference_context(
            project_id,
            reference_sources=reference_sources,
            main_character=main_character,
            project=project,
            world=world,
            include_characters=False,
        )
        prompt = "\n".join(
            [
                "日本語で、ノベル再生用の事前生成ノベルゲーム作品の制作設計書を作成してください。",
                "これはリアルタイム生成ではなく、章立て、各章本文、各章画像を制作モードで事前生成してから公開する作品です。",
                "読者は鑑賞モードで待ち時間なく読み進めます。",
                "章数は短く濃くしてください。標準は5章で、導入、展開、転機、クライマックス、余韻の読み切り構成を優先してください。",
                "各章は長文説明よりも、短い本文と多めの画像でテンポよく読ませる前提にしてください。",
                "登場人物は、主人公を除き、可能な限り下記のDB登録済みキャラクターだけを使ってください。",
                "主人公は指定名を優先してよく、DB未登録でも構いません。ただし脇役、敵役、組織代表、関係者は登録済みキャラクターから選んでください。",
                "知らない新キャラクターを安易に増やさないでください。新キャラクターが必要な場合は、なぜ既存キャラクターで代替できないかを明記してください。",
                "チェックされた参照データだけを素材として使い、参照内容と矛盾する設定を作らないでください。",
                "",
                f"タイトル: {title}",
                f"ジャンル: {genre}",
                f"主役: {main_character or '未指定'}",
                f"テーマ: {theme or '未指定'}",
                f"章数: {chapter_count}",
                f"各章の目標文字数: {chapter_target_chars}",
                "",
                "有効な参照データ:",
                ", ".join(reference_sources),
                "",
                "参照素材:",
                reference_context or "追加参照なし",
                "",
                "DB登録済みキャラクター:",
                registered_character_context or "登録済みキャラクターなし",
                "",
                "必ず以下の構成で出力してください。",
                "1. ログライン",
                "2. 作品コンセプト",
                "3. 主要キャラクターの演出メモ",
                "4. 全体構成",
                "5. 章立て一覧。各章に、章タイトル、目的、主要シーン、劇中スチル案を含める",
                "6. 深掘り生成方針。各章をどう膨らませるか",
                "7. ノベルゲーム化方針。シーン分割、話者名、栞、画像プリロードの前提",
                "説明だけでなく、このまま制作ジョブに渡せる具体度にしてください。",
            ]
        )
        prompt = "\n".join(
            [
                prompt,
                "",
                "Machine-readable chapter list rules:",
                "Under the heading '## 章立て一覧', include one chapter header per chapter in exactly this format:",
                "- 第1章: 章タイトル",
                "  - 概要: 章の内容を1〜3文で書く",
                "  - 目的: この章の物語上の役割を書く",
                "Do not put the chapter title on a separate line. Do not use '# 第1章' as the only chapter line.",
                f"Create exactly {chapter_count} chapter header lines in this format.",
            ]
        )
        result = self._text_ai_client.generate_text(
            prompt,
            model=settings.get("model"),
            temperature=0.75,
            max_tokens=32000,
        )
        return {
            "model": result.get("model"),
            "chapter_target_chars": chapter_target_chars,
            "outline_markdown": result.get("text") or "",
            "usage": result.get("usage"),
        }

    def generate_production_premise(self, project_id: int, payload: dict | None = None):
        payload = dict(payload or {})
        settings = self._user_setting_service.apply_cinema_novel_text_generation_settings(payload.get("text_options") or {})
        project = Project.query.get(project_id)
        world = World.query.filter_by(project_id=project_id).first()
        current_input = payload.get("current_input") or {}
        requested_main_character = str(current_input.get("main_character") or "").strip()
        requested_genre = str(current_input.get("genre") or "").strip()
        reference_sources = self._reference_sources(current_input)
        character_context = (
            self._registered_character_context(project_id, main_character=requested_main_character)
            if self._reference_source_enabled(reference_sources, "characters")
            else ""
        )
        reference_context = self._production_reference_context(
            project_id,
            reference_sources=reference_sources,
            main_character=requested_main_character,
            project=project,
            world=world,
            include_characters=False,
        )
        prompt = "\n".join(
            [
                "Return only JSON.",
                "ノベル用の長編企画を、日本語で1案だけ提案してください。",
                "ユーザーが手入力しなくても章立て制作設計へ進めるための、入力欄の初期案を作ります。",
                "DB登録済みキャラクターを可能な限り使ってください。主人公も登録済みキャラから選ぶのを優先します。",
                "ただし、物語上どうしても必要なら主人公だけは新規名でも構いません。その場合も理由を protagonist_reason に書いてください。",
                "知らない脇役や敵役を増やさず、既存キャラの関係性・思想・口調・弱点を活かしてください。",
                "Required JSON keys: title, main_character, protagonist_reason, genre, chapter_count, theme, concept_note.",
                "chapter_count は 5 を基本にし、必要な場合だけ 3, 4, 6, 8 のいずれかにしてください。",
                "ユーザー指定ジャンルがある場合は必ず反映してください。",
                "ユーザー指定主役がある場合は必ずそのキャラクターを主人公にしてください。",
                "チェックされた参照データだけを素材として使い、参照内容と矛盾する設定を作らないでください。",
                "",
                "User current input:",
                f"title: {str(current_input.get('title') or '').strip()}",
                f"main_character: {requested_main_character or 'AI選定'}",
                f"genre: {requested_genre or 'AI提案'}",
                f"chapter_count: {str(current_input.get('chapter_count') or '5').strip()}",
                f"theme: {str(current_input.get('theme') or '').strip()}",
                "",
                "Enabled reference sources:",
                ", ".join(reference_sources),
                "",
                "Reference material:",
                reference_context or "追加参照なし",
                "",
                "DB registered characters:",
                character_context or "登録済みキャラクターなし",
            ]
        )
        result = self._text_ai_client.generate_text(
            prompt,
            model=settings.get("model"),
            response_format={"type": "json_object"},
            temperature=0.7,
            max_tokens=6000,
        )
        parsed = self._text_ai_client._try_parse_json(result.get("text"))
        if not isinstance(parsed, dict):
            raise RuntimeError("production premise response is invalid")
        chapter_count = parsed.get("chapter_count")
        try:
            chapter_count = int(chapter_count)
        except (TypeError, ValueError):
            chapter_count = 5
        if chapter_count not in {3, 4, 5, 6, 8}:
            chapter_count = 5
        return {
            "title": str(parsed.get("title") or "").strip(),
            "main_character": str(parsed.get("main_character") or "").strip(),
            "protagonist_reason": str(parsed.get("protagonist_reason") or "").strip(),
            "genre": str(parsed.get("genre") or "映画ノベル").strip() or "映画ノベル",
            "chapter_count": chapter_count,
            "theme": str(parsed.get("theme") or "").strip(),
            "concept_note": str(parsed.get("concept_note") or "").strip(),
            "model": result.get("model"),
            "usage": result.get("usage"),
        }

    def _registered_character_context(self, project_id: int, *, main_character: str = "") -> str:
        characters = Character.query.filter(
            Character.project_id == project_id,
            Character.deleted_at.is_(None),
        ).order_by(Character.id.asc()).all()
        if not characters:
            return ""
        main_name = str(main_character or "").strip().lower()
        lines = []
        for character in characters[:60]:
            role_note = "主人公候補または主要人物"
            if main_name and main_name in {str(character.name or "").strip().lower(), str(character.nickname or "").strip().lower()}:
                role_note = "指定主人公と同一または近い登録キャラクター"
            lines.extend(
                [
                    f"- id={character.id} name={character.name or ''} nickname={character.nickname or ''} role_note={role_note}",
                    f"  overview={self._shorten_for_prompt(getattr(character, 'character_summary', None), 900)}",
                    f"  personality={self._shorten_for_prompt(character.personality, 500)}",
                    f"  first_person={character.first_person or ''}",
                    f"  second_person={character.second_person or ''}",
                    f"  speech_style={self._shorten_for_prompt(character.speech_style, 350)}",
                    f"  speech_sample={self._shorten_for_prompt(character.speech_sample, 500)}",
                    f"  appearance={self._shorten_for_prompt(character.appearance_summary, 350)}",
                    f"  ng_rules={self._shorten_for_prompt(character.ng_rules, 250)}",
                ]
            )
        return "\n".join(lines)

    def _reference_sources(self, payload: dict | None) -> list[str]:
        payload = payload if isinstance(payload, dict) else {}
        raw = payload.get("reference_sources")
        if raw is None:
            raw = payload.get("reference_context_sources")
        if raw is None:
            return list(self.DEFAULT_REFERENCE_SOURCES)
        if isinstance(raw, str):
            values = re.split(r"[,|\s]+", raw)
        elif isinstance(raw, (list, tuple, set)):
            values = list(raw)
        else:
            values = []
        sources = []
        for value in values:
            key = str(value or "").strip()
            if key in self.VALID_REFERENCE_SOURCES and key not in sources:
                sources.append(key)
        return sources

    def _reference_source_enabled(self, reference_sources: list[str], key: str) -> bool:
        return key in set(reference_sources or [])

    def _production_reference_context(
        self,
        project_id: int,
        *,
        reference_sources: list[str],
        main_character: str = "",
        project=None,
        world=None,
        include_characters: bool = True,
    ) -> str:
        sources = reference_sources or []
        lines = []
        if self._reference_source_enabled(sources, "worldbuilding"):
            project = project or Project.query.get(project_id)
            lines.extend(
                [
                    "## 世界観",
                    f"title: {getattr(project, 'title', '') or ''}",
                    f"summary: {self._shorten_for_prompt(getattr(project, 'summary', '') or '', 1200)}",
                ]
            )
        if self._reference_source_enabled(sources, "world"):
            world = world or World.query.filter_by(project_id=project_id).first()
            lines.extend(
                [
                    "## ワールド",
                    f"name: {getattr(world, 'name', '') or ''}",
                    f"tone: {getattr(world, 'tone', '') or ''}",
                    f"era: {getattr(world, 'era_description', '') or ''}",
                    f"overview: {self._shorten_for_prompt(getattr(world, 'overview', '') or '', 1600)}",
                    f"technology: {self._shorten_for_prompt(getattr(world, 'technology_level', '') or '', 700)}",
                    f"social_structure: {self._shorten_for_prompt(getattr(world, 'social_structure', '') or '', 700)}",
                    f"rules: {self._shorten_for_prompt(getattr(world, 'rules_json', '') or '', 700)}",
                    f"forbidden: {self._shorten_for_prompt(getattr(world, 'forbidden_json', '') or '', 700)}",
                ]
            )
        if include_characters and self._reference_source_enabled(sources, "characters"):
            lines.extend(
                [
                    "## キャラクター設定",
                    self._registered_character_context(project_id, main_character=main_character) or "登録済みキャラクターなし",
                ]
            )
        if self._reference_source_enabled(sources, "feed"):
            lines.extend(["## Feed", self._recent_feed_context(project_id)])
        if self._reference_source_enabled(sources, "news"):
            lines.extend(["## ニュース", self._recent_world_news_context(project_id)])
        if self._reference_source_enabled(sources, "short_stories"):
            lines.extend(["## キャラクターのショートストーリー", self._recent_short_story_context(project_id)])
        return "\n".join(part for part in lines if part is not None).strip()

    def _recent_feed_context(self, project_id: int, limit: int = 12) -> str:
        posts = (
            FeedPost.query.filter(
                FeedPost.project_id == project_id,
                FeedPost.status == "published",
                FeedPost.deleted_at.is_(None),
            )
            .order_by(FeedPost.published_at.desc(), FeedPost.created_at.desc(), FeedPost.id.desc())
            .limit(limit)
            .all()
        )
        if not posts:
            return "Feed投稿なし"
        character_ids = {post.character_id for post in posts if post.character_id}
        characters = {
            character.id: character
            for character in Character.query.filter(Character.id.in_(character_ids)).all()
        } if character_ids else {}
        lines = []
        for post in posts:
            character = characters.get(post.character_id)
            speaker = getattr(character, "name", None) or f"character_id={post.character_id}"
            lines.append(f"- {speaker}: {self._shorten_for_prompt(post.body, 260)}")
        return "\n".join(lines)

    def _recent_world_news_context(self, project_id: int, limit: int = 8) -> str:
        items = (
            WorldNewsItem.query.filter(
                WorldNewsItem.project_id == project_id,
                WorldNewsItem.status == "published",
                WorldNewsItem.deleted_at.is_(None),
            )
            .order_by(WorldNewsItem.created_at.desc(), WorldNewsItem.id.desc())
            .limit(limit)
            .all()
        )
        if not items:
            return "ニュースなし"
        lines = []
        for item in items:
            summary = item.summary or item.body
            lines.append(
                f"- [{item.news_type or 'news'}] {item.title or ''}: {self._shorten_for_prompt(summary, 320)}"
            )
        return "\n".join(lines)

    def _recent_short_story_context(self, project_id: int, limit: int = 6) -> str:
        sessions = (
            ChatSession.query.filter(
                ChatSession.project_id == project_id,
                ChatSession.deleted_at.is_(None),
            )
            .order_by(ChatSession.updated_at.desc(), ChatSession.id.desc())
            .limit(40)
            .all()
        )
        stories = []
        for session in sessions:
            settings = self._load_json_object(getattr(session, "settings_json", None))
            saved = settings.get("saved_short_stories")
            if not isinstance(saved, list):
                continue
            for story in reversed(saved[-5:]):
                if isinstance(story, dict):
                    stories.append((session, story))
                if len(stories) >= limit:
                    break
            if len(stories) >= limit:
                break
        if not stories:
            return "保存済みショートストーリーなし"
        lines = []
        for session, story in stories[:limit]:
            title = str(story.get("title") or getattr(session, "title", "") or "ショートストーリー").strip()
            synopsis = str(story.get("synopsis") or "").strip()
            body = str(story.get("body") or "").strip()
            excerpt = synopsis or body
            lines.append(f"- {title}: {self._shorten_for_prompt(excerpt, 420)}")
        return "\n".join(lines)

    def _load_json_object(self, value) -> dict:
        if not value:
            return {}
        try:
            parsed = json_util.loads(value) if isinstance(value, str) else value
        except Exception:
            return {}
        return parsed if isinstance(parsed, dict) else {}

    def _generate_character_review(self, novel, character) -> dict:
        settings = self._user_setting_service.apply_cinema_novel_text_generation_settings({})
        novel_context = self._novel_review_context(novel)
        lore_context = self._lore_prompt_context(self.list_lore_entries(novel.id))
        character_context = "\n".join(
            [
                f"name: {character.name or ''}",
                f"nickname: {character.nickname or ''}",
                f"first_person: {character.first_person or ''}",
                f"second_person: {character.second_person or ''}",
                f"summary: {self._shorten_for_prompt(getattr(character, 'character_summary', None), 1000)}",
                f"personality: {self._shorten_for_prompt(character.personality, 900)}",
                f"speech_style: {self._shorten_for_prompt(character.speech_style, 700)}",
                f"speech_sample: {self._shorten_for_prompt(character.speech_sample, 700)}",
                f"ng_rules: {self._shorten_for_prompt(character.ng_rules, 400)}",
            ]
        )
        prompt = "\n".join(
            [
                "Return only JSON.",
                "A registered character has read/watched the following visual novel as an in-world cinema work.",
                "Create a public Feed review and a private memory note for future chat.",
                "Japanese only.",
                "Required keys: feed_review, memory_note, rating_label, review_summary.",
                "feed_review: 80-220 Japanese characters. Write as the character posting to Feed, in their voice. Mention one specific memorable element from the novel.",
                "memory_note: 120-360 Japanese characters. Third-person memory for AI prompt. It must say this character has read/watched the work, what they remember, and how they tend to talk about it.",
                "rating_label: short Japanese label such as 爆笑, 傑作, 怪作, 刺さった, 困惑.",
                "Do not change the character's permanent personality. This is a viewing experience memory.",
                "Do not invent facts that contradict the novel context.",
                "",
                "Character:",
                character_context,
                "",
                "Novel context:",
                novel_context,
                "",
                "Known novel lore:",
                lore_context or "(none yet)",
            ]
        )
        result = self._text_ai_client.generate_text(
            prompt,
            model=settings.get("model"),
            response_format={"type": "json_object"},
            temperature=0.8,
            max_tokens=2000,
        )
        parsed = self._text_ai_client._try_parse_json(result.get("text")) or {}
        if not isinstance(parsed, dict):
            parsed = {}
        parsed["model"] = result.get("model")
        parsed["usage"] = result.get("usage")
        return parsed

    def _generate_novel_lore(self, novel) -> dict:
        settings = self._user_setting_service.apply_cinema_novel_text_generation_settings({})
        prompt = "\n".join(
            [
                "Return only JSON.",
                "Extract reusable in-world knowledge from this visual novel for future character chats.",
                "Japanese only.",
                "Required shape: {\"entries\":[{\"lore_type\":\"character|term|event|location|scene|theme|other\",\"name\":\"...\",\"summary\":\"...\",\"role_note\":\"...\",\"source_note\":\"...\"}]}",
                "Focus on characters, named concepts, important incidents, iconic scenes, relationships, jokes, and emotional hooks.",
                "Character entries must explain how the character appears in this novel, what they want, what makes them funny or memorable, and how they relate to other entries.",
                "Term/event entries must be understandable later without rereading the novel.",
                "Do not invent facts that are not supported by the novel context.",
                "Create 8-24 compact entries.",
                "",
                "Novel context:",
                self._novel_review_context(novel),
            ]
        )
        result = self._text_ai_client.generate_text(
            prompt,
            model=settings.get("model"),
            response_format={"type": "json_object"},
            temperature=0.35,
            max_tokens=5000,
        )
        parsed = self._text_ai_client._try_parse_json(result.get("text")) or {}
        if not isinstance(parsed, dict):
            parsed = {}
        parsed["model"] = result.get("model")
        parsed["usage"] = result.get("usage")
        return parsed

    def _generate_character_impressions(self, novel, character, lore_entries: list) -> dict:
        settings = self._user_setting_service.apply_cinema_novel_text_generation_settings({})
        character_context = "\n".join(
            [
                f"name: {character.name or ''}",
                f"nickname: {character.nickname or ''}",
                f"first_person: {character.first_person or ''}",
                f"summary: {self._shorten_for_prompt(getattr(character, 'character_summary', None), 1000)}",
                f"personality: {self._shorten_for_prompt(character.personality, 900)}",
                f"speech_style: {self._shorten_for_prompt(character.speech_style, 700)}",
                f"speech_sample: {self._shorten_for_prompt(character.speech_sample, 700)}",
            ]
        )
        prompt = "\n".join(
            [
                "Return only JSON.",
                "A registered character has watched this visual novel. Create that character's private impressions of the novel's characters, terms, and iconic scenes.",
                "Japanese only.",
                "Required shape: {\"impressions\":[{\"target_name\":\"...\",\"impression_text\":\"...\",\"talk_hint\":\"...\"}]}",
                "target_name must match an entry name from Known novel lore when possible.",
                "impression_text: 80-260 Japanese characters. Write in third person. Explain how the reviewing character interprets or reacts to that target.",
                "talk_hint: 40-160 Japanese characters. How this reviewing character should bring it up in future chats.",
                "Prefer 4-10 memorable targets. Include important registered characters and the funniest or most emotionally useful concepts.",
                "Do not change the reviewing character's permanent personality. This is a viewing experience memory.",
                "",
                "Reviewing character:",
                character_context,
                "",
                "Novel:",
                f"title: {novel.title or ''}",
                f"description: {novel.description or ''}",
                "",
                "Known novel lore:",
                self._lore_prompt_context(lore_entries) or "(none)",
            ]
        )
        result = self._text_ai_client.generate_text(
            prompt,
            model=settings.get("model"),
            response_format={"type": "json_object"},
            temperature=0.55,
            max_tokens=4000,
        )
        parsed = self._text_ai_client._try_parse_json(result.get("text")) or {}
        if not isinstance(parsed, dict):
            parsed = {}
        parsed["model"] = result.get("model")
        parsed["usage"] = result.get("usage")
        return parsed

    def _lore_prompt_context(self, entries: list) -> str:
        lines = []
        for entry in entries or []:
            if isinstance(entry, dict):
                lore_type = entry.get("lore_type") or "other"
                name = entry.get("name") or ""
                summary = entry.get("summary") or ""
                role_note = entry.get("role_note") or ""
            else:
                lore_type = entry.lore_type
                name = entry.name
                summary = entry.summary
                role_note = entry.role_note
            if not name or not summary:
                continue
            lines.append(f"- [{lore_type}] {name}: {self._shorten_for_prompt(summary, 600)}")
            if role_note:
                lines.append(f"  role_note={self._shorten_for_prompt(role_note, 300)}")
        return "\n".join(lines)[:9000]

    def _resolve_lore_target_character_id(self, project_id: int, target_name: str) -> int | None:
        normalized = str(target_name or "").strip().lower()
        if not normalized:
            return None
        characters = Character.query.filter(
            Character.project_id == project_id,
            Character.deleted_at.is_(None),
        ).all()
        for character in characters:
            names = {str(character.name or "").strip().lower(), str(character.nickname or "").strip().lower()}
            if normalized in names:
                return character.id
        return None

    def _generate_and_upsert_character_impressions(self, *, novel, character, user_id: int, lore_entries: list):
        result = self._generate_character_impressions(novel, character, lore_entries)
        impressions = result.get("impressions") if isinstance(result, dict) else []
        if not isinstance(impressions, list):
            impressions = []
        saved = []
        for item in impressions[:12]:
            if not isinstance(item, dict):
                continue
            target_name = str(item.get("target_name") or "").strip()[:255]
            impression_text = str(item.get("impression_text") or "").strip()
            if not target_name or not impression_text:
                continue
            row = CinemaNovelCharacterImpression.query.filter(
                CinemaNovelCharacterImpression.novel_id == novel.id,
                CinemaNovelCharacterImpression.reviewer_character_id == character.id,
                CinemaNovelCharacterImpression.user_id == user_id,
                CinemaNovelCharacterImpression.target_name == target_name,
            ).first()
            if not row:
                row = CinemaNovelCharacterImpression(
                    novel_id=novel.id,
                    reviewer_character_id=character.id,
                    user_id=user_id,
                    target_name=target_name,
                )
            row.target_character_id = self._resolve_lore_target_character_id(novel.project_id, target_name)
            row.impression_text = impression_text[:1200]
            row.talk_hint = str(item.get("talk_hint") or "").strip()[:800] or None
            row.metadata_json = json_util.dumps(
                {
                    "source": "cinema_novel_character_impression",
                    "model": result.get("model"),
                    "usage": result.get("usage"),
                }
            )
            row.deleted_at = None
            db.session.add(row)
            saved.append(row)
        db.session.commit()
        return self.list_character_impressions(novel.id, reviewer_character_id=character.id, user_id=user_id)

    def _novel_review_context(self, novel) -> str:
        production = self._load_json(novel.production_json, default={})
        lines = [
            f"title: {novel.title or ''}",
            f"subtitle: {novel.subtitle or ''}",
            f"description: {novel.description or ''}",
        ]
        source_input = production.get("source_input") if isinstance(production, dict) else {}
        if isinstance(source_input, dict):
            premise = source_input.get("premise") if isinstance(source_input.get("premise"), dict) else {}
            if premise:
                lines.extend(
                    [
                        f"genre: {premise.get('genre') or ''}",
                        f"theme: {premise.get('theme') or ''}",
                        f"concept: {premise.get('concept_note') or ''}",
                    ]
                )
        outline = str((production or {}).get("outline_markdown") or "").strip()
        if outline:
            lines.append("production_outline:")
            lines.append(self._shorten_for_prompt(outline, 4500))
        chapters = self.list_chapters(novel.id)
        if chapters:
            lines.append("chapters:")
            for chapter in chapters[:8]:
                body = str(chapter.body_markdown or "").strip()
                lines.append(f"- {chapter.chapter_no}. {chapter.title}: {self._shorten_for_prompt(body, 1200)}")
        return "\n".join(lines)[:12000]

    def _fallback_review_memory_note(self, novel, character, review_text: str) -> str:
        return (
            f"{character.name}はラプ・シネマの上映作品『{novel.title}』を鑑賞済み。"
            f"印象に残った感想として「{review_text[:180]}」という反応を持っている。"
            "今後この作品が話題に出たら、鑑賞済みの体験として自分の口調で反応できる。"
        )

    def _upsert_review_feed_post(self, *, novel, character, user_id: int, review_text: str, existing_feed_post_id: int | None):
        post = FeedPost.query.filter(
            FeedPost.id == existing_feed_post_id,
            FeedPost.deleted_at.is_(None),
        ).first() if existing_feed_post_id else None
        generation_state = json_util.dumps(
            {
                "source": "cinema_novel_review",
                "cinema_novel_id": novel.id,
                "cinema_novel_title": novel.title,
                "character_id": character.id,
            }
        )
        if post:
            post.body = review_text
            post.character_id = character.id
            post.status = "published"
            post.generation_state_json = generation_state
            if not post.published_at:
                post.published_at = datetime.utcnow()
            db.session.commit()
            return post
        post = FeedPost(
            project_id=novel.project_id,
            character_id=character.id,
            created_by_user_id=user_id,
            body=review_text,
            status="published",
            like_count=0,
            generation_state_json=generation_state,
            published_at=datetime.utcnow(),
        )
        db.session.add(post)
        db.session.commit()
        return post

    def _upsert_review_memory_note(self, *, novel, character, user_id: int, note_text: str, existing_memory_note_id: int | None):
        note = CharacterMemoryNote.query.filter(
            CharacterMemoryNote.id == existing_memory_note_id,
            CharacterMemoryNote.user_id == user_id,
            CharacterMemoryNote.character_id == character.id,
        ).first() if existing_memory_note_id else None
        source_ref = f"cinema_novel:{novel.id}"
        if note:
            note.category = "fun_fact"
            note.note = note_text[:1000]
            note.source_type = "cinema_novel_review"
            note.source_ref = source_ref
            note.confidence = 1.0
            note.enabled = True
            db.session.commit()
            return note
        note = CharacterMemoryNote.query.filter(
            CharacterMemoryNote.user_id == user_id,
            CharacterMemoryNote.character_id == character.id,
            CharacterMemoryNote.source_type == "cinema_novel_review",
            CharacterMemoryNote.source_ref == source_ref,
        ).first()
        if note:
            note.note = note_text[:1000]
            note.enabled = True
            db.session.commit()
            return note
        note = CharacterMemoryNote(
            user_id=user_id,
            character_id=character.id,
            category="fun_fact",
            note=note_text[:1000],
            source_type="cinema_novel_review",
            source_ref=source_ref,
            confidence=1.0,
            enabled=True,
            pinned=False,
        )
        db.session.add(note)
        db.session.commit()
        return note

    def _upsert_impression_memory_note(self, *, novel, character, user_id: int, impressions: list):
        if not impressions:
            return None
        lines = []
        for impression in impressions[:8]:
            if not impression.impression_text:
                continue
            hint = f" 話題化: {impression.talk_hint}" if impression.talk_hint else ""
            lines.append(f"- {impression.target_name}: {impression.impression_text}{hint}")
        if not lines:
            return None
        note_text = (
            f"{character.name}は『{novel.title}』の登場人物・用語について次の鑑賞印象を持っている。\n"
            + "\n".join(lines)
        )[:1000]
        source_ref = f"cinema_novel:{novel.id}:impressions"
        note = CharacterMemoryNote.query.filter(
            CharacterMemoryNote.user_id == user_id,
            CharacterMemoryNote.character_id == character.id,
            CharacterMemoryNote.source_type == "cinema_novel_character_impression",
            CharacterMemoryNote.source_ref == source_ref,
        ).first()
        if not note:
            note = CharacterMemoryNote(
                user_id=user_id,
                character_id=character.id,
                category="fun_fact",
                source_type="cinema_novel_character_impression",
                source_ref=source_ref,
                confidence=1.0,
                enabled=True,
                pinned=False,
        )
        note.note = note_text
        note.enabled = True
        db.session.add(note)
        db.session.commit()
        return note

    def _shorten_for_prompt(self, value, limit: int = 500) -> str:
        text = str(value or "").strip().replace("\r\n", "\n")
        if len(text) <= limit:
            return text
        return text[:limit].rstrip() + "..."

    def _source_novel_id_from_payload(self, project_id: int, payload: dict | None) -> int:
        payload = payload or {}
        try:
            source_novel_id = int(payload.get("source_novel_id") or 0)
        except (TypeError, ValueError):
            return 0
        if source_novel_id <= 0:
            return 0
        novel = self.get_novel(source_novel_id)
        if not novel or int(novel.project_id) != int(project_id):
            raise ValueError("source_novel_id is invalid")
        return source_novel_id

    def _source_novel_character_ids_from_payload(self, project_id: int, payload: dict | None) -> list[int]:
        source_novel_id = self._source_novel_id_from_payload(project_id, payload)
        if not source_novel_id:
            return []
        source_novel = self.get_novel(source_novel_id)
        if not source_novel:
            return []

        # Chat-derived learning novels may have older scene JSON without character_ids.
        # In that case, recover the teacher/student pair from the original chat session snapshot.
        source_session = self._novel_source_chat_session(source_novel)
        session_role_ids = self._chat_session_learning_role_character_ids(source_session) if source_session else []
        if session_role_ids:
            return session_role_ids

        character_ids = []
        for chapter in self.list_chapters(source_novel.id):
            scenes = self._load_json(chapter.scene_json, default=[])
            if not isinstance(scenes, list):
                continue
            for scene in scenes:
                if not isinstance(scene, dict):
                    continue
                for character_id in self._normalize_character_ids(scene.get("character_ids")):
                    if character_id not in character_ids:
                        character_ids.append(character_id)
                        if len(character_ids) >= 6:
                            return character_ids
        return character_ids

    def _source_novel_context_for_prompt(self, project_id: int, payload: dict | None, limit: int = 9000) -> str:
        payload = payload or {}
        source_novel_id = self._source_novel_id_from_payload(project_id, payload)
        if not source_novel_id:
            return ""
        novel = self.get_novel(source_novel_id)
        chapters = self.list_chapters(novel.id)
        lines = [
            f"原作ノベルID: {novel.id}",
            f"原作タイトル: {novel.title or ''}",
            f"原作サブタイトル: {novel.subtitle or ''}",
            f"原作説明: {novel.description or ''}",
            "",
            "原作本文:",
        ]
        for chapter in chapters:
            lines.extend(
                [
                    f"## 第{chapter.chapter_no}章 {chapter.title or ''}",
                    self._shorten_for_prompt(chapter.body_markdown or "", 1800),
                    "",
                ]
            )
        return self._shorten_for_prompt("\n".join(lines), limit)

    def _normalize_bool(self, value, default: bool = False) -> bool:
        if value is None:
            return default
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return bool(value)
        return str(value).strip().lower() in {"1", "true", "yes", "on"}

    def save_production_outline(self, project_id: int, user_id: int, payload: dict | None):
        payload = dict(payload or {})
        title = str(payload.get("title") or "無題のノベル作品").strip() or "無題のノベル作品"
        outline_markdown = str(payload.get("outline_markdown") or "").strip()
        if not outline_markdown:
            raise ValueError("outline_markdown is required")
        source_input = payload.get("source_input") if isinstance(payload.get("source_input"), dict) else {}
        existing = CinemaNovel.query.filter(
            CinemaNovel.project_id == project_id,
            CinemaNovel.title == title,
            CinemaNovel.source_path.is_(None),
            CinemaNovel.deleted_at.is_(None),
        ).first()
        production_payload = {
            "source": "production_outline",
            "reader": "prebuilt",
            "generation_mode": "build_then_publish",
            "image_generation": "prebuilt_only",
            "source_input": source_input,
            "outline_markdown": outline_markdown,
            "model": payload.get("model"),
            "chapter_target_chars": payload.get("chapter_target_chars"),
            "usage": payload.get("usage"),
        }
        if existing:
            existing.subtitle = str(payload.get("subtitle") or existing.subtitle or "ノベル制作設計").strip() or None
            existing.description = str(payload.get("description") or source_input.get("theme") or existing.description or "").strip() or None
            existing.status = str(payload.get("status") or existing.status or "draft").strip() or "draft"
            existing.mobile_visible = self._normalize_bool(payload.get("mobile_visible", getattr(existing, "mobile_visible", True)))
            existing.production_json = json_util.dumps(production_payload)
            db.session.commit()
            return existing
        novel = CinemaNovel(
            project_id=project_id,
            created_by_user_id=user_id,
            title=title,
            subtitle=str(payload.get("subtitle") or "ノベル制作設計").strip() or None,
            description=str(payload.get("description") or source_input.get("theme") or "").strip() or None,
            status=str(payload.get("status") or "draft").strip() or "draft",
            mobile_visible=self._normalize_bool(payload.get("mobile_visible", True)),
            mode="cinema_novel",
            production_json=json_util.dumps(production_payload),
        )
        if novel.status not in self.VALID_STATUSES:
            novel.status = "draft"
        db.session.add(novel)
        db.session.commit()
        return novel

    def create_chapters_from_production_outline(self, novel_id: int):
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        existing_chapters = self.list_chapters(novel.id)
        if existing_chapters:
            return existing_chapters
        production = self._load_json(novel.production_json, default={})
        outline = str((production or {}).get("outline_markdown") or "").strip()
        if not outline:
            raise ValueError("production outline is required")
        chapter_items = self._extract_chapter_items_from_outline(outline)
        if not chapter_items:
            raise ValueError("chapter list could not be detected from production outline")
        chapters = []
        for index, item in enumerate(chapter_items, start=1):
            chapter_no = int(item.get("chapter_no") or index)
            title = str(item.get("title") or f"第{chapter_no}章").strip()
            body = "\n".join(
                [
                    f"# {chapter_no:02d}. {title}",
                    "",
                    "## 制作設計メモ",
                    str(item.get("outline") or "").strip(),
                    "",
                    "## 本文",
                    "この章はまだ本文生成前です。制作パネルの「この章を深掘り生成」から本文を作成してください。",
                ]
            ).strip()
            chapter = CinemaNovelChapter(
                novel_id=novel.id,
                chapter_no=chapter_no,
                title=title,
                body_markdown=body,
                scene_json=json_util.dumps(self._markdown_to_scenes(body, chapter_no=chapter_no)),
                sort_order=index,
            )
            db.session.add(chapter)
            chapters.append(chapter)
        db.session.commit()
        return chapters

    def generate_title_image(self, novel_id: int, payload: dict | None = None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        production = self._load_json(novel.production_json, default={})
        outline = str((production or {}).get("outline_markdown") or "").strip()
        source_input = (production or {}).get("source_input") if isinstance((production or {}).get("source_input"), dict) else {}
        premise = str(payload.get("premise") or source_input.get("theme") or novel.description or "").strip()
        options = self._user_setting_service.apply_cinema_novel_image_generation_settings(
            payload.get("image_options") or payload
        )
        options = self._apply_mobile_novel_image_options(novel, options)
        references = self._matching_character_references(
            novel.project_id,
            "\n".join([novel.title or "", premise, outline[:2500]]),
            limit=4,
        )
        reference_ids = [item.get("base_asset_id") for item in references if item.get("base_asset_id")]
        reference_ids = self._apply_novel_session_outfit_references(
            novel,
            reference_ids,
            character_ids=[item.get("id") for item in references if item.get("id")],
        )
        reference_paths, reference_asset_ids = self._resolve_reference_image_paths(reference_ids)
        prompt = "\n".join(
            [
                f"ノベルゲーム『{novel.title}』のタイトル画像、オープニング画像。",
                f"作品起動時に最初に表示される派手なキービジュアル。{self._novel_image_layout_instruction(novel)}",
                "日本のビジュアルノベル、映画ポスター、ゲームタイトル画面の雰囲気。",
                "タイトルロゴを大きく中央または上部に配置。発光、金属感、ネオン、粒子、強いコントラストで印象的に。",
                "キャラクターがいる場合は、参考画像の顔立ち、髪型、衣装、雰囲気を保つ。",
                "読める文字は作品タイトルだけにする。余計な英字、透かし、出版社ロゴ、UIは入れない。",
                f"タイトル: {novel.title}",
                f"企画: {premise[:1200]}",
                f"章立て・世界観: {outline[:2500]}",
            ]
        )
        asset = self._generate_cinema_asset(
            project_id=novel.project_id,
            asset_type="cinema_novel_title_image",
            file_prefix=f"cinema_novel_{novel.id}_title",
            prompt=prompt,
            image_options=options,
            metadata={
                "source": "cinema_novel_title_image",
                "novel_id": novel.id,
                "reference_asset_ids": reference_asset_ids,
            },
            reference_paths=reference_paths,
        )
        novel.cover_asset_id = asset.id
        novel.poster_asset_id = asset.id
        db.session.commit()
        return {
            "novel": self.serialize_novel(novel, include_chapters=True),
            "asset": self._serialize_asset(asset.id),
            "reference_asset_ids": reference_asset_ids,
            "image_options": {
                "provider": options.get("provider"),
                "model": options.get("model"),
                "quality": options.get("quality"),
                "size": options.get("size"),
            },
        }

    def _extract_chapter_items_from_outline(self, outline: str):
        lines = outline.splitlines()
        chapter_section_lines = []
        in_chapter_section = False
        for line in lines:
            stripped = line.strip()
            if "章立て一覧" in stripped or "CHAPTER_LIST_FOR_SYSTEM" in stripped:
                in_chapter_section = True
                continue
            if in_chapter_section and re.match(r"^#{1,6}\s*\d+\.", stripped):
                break
            if in_chapter_section:
                chapter_section_lines.append(line)
        if chapter_section_lines:
            lines = chapter_section_lines
        items = []
        current = None
        chapter_pattern = re.compile(
            r"^\s*(?:#{1,6}\s*)?(?:[-\*]\s*)?第\s*([0-9０-９]{1,2})\s*(?:章|話|幕)\s*(.*)$"
        )
        title_only_pattern = re.compile(r"^\s*(?:#{1,6}\s*)?(?:章タイトル|タイトル)\s*[:：]?\s*(.+?)\s*$")

        def normalize_number(value: str) -> int:
            return int(str(value).translate(str.maketrans("０１２３４５６７８９", "0123456789")))

        def clean_title(value: str) -> str:
            title = str(value or "").strip()
            title = re.sub(r"^[\s:：\-・|｜]+", "", title)
            title = re.sub(r"^\*\*(.+?)\*\*$", r"\1", title)
            title = re.sub(r"^第\s*[0-9０-９]{1,2}\s*(?:章|話|幕)\s*[:：\-・|｜]*", "", title).strip()
            title = re.sub(r"^[「『【\[]|[」』】\]]$", "", title.strip(" -:："))
            title = re.split(r"\s{2,}|[｜|]", title, 1)[0].strip()
            return title

        for line in lines:
            stripped = line.strip()
            match = chapter_pattern.match(stripped)
            if match and not re.search(r"(章数|目標文字数|ターン|ステップ|候補|画像|章扉|劇中スチル|案)$", stripped):
                if current:
                    items.append(current)
                chapter_no = normalize_number(match.group(1))
                title = clean_title(match.group(2))
                current = {
                    "chapter_no": chapter_no,
                    "title": title[:120] or f"第{chapter_no}章",
                    "outline_lines": [stripped],
                }
                continue
            title_match = title_only_pattern.match(stripped)
            if current and title_match:
                title = clean_title(title_match.group(1))
                if title and not re.search(r"(概要|目的|主要シーン|目標文字数)$", title):
                    current["title"] = title[:120]
                current["outline_lines"].append(stripped)
                continue
            if current and stripped:
                if not current.get("title") or re.fullmatch(r"第\d{1,2}章", str(current.get("title") or "")):
                    fallback_title = clean_title(stripped)
                    if fallback_title and fallback_title != stripped and not re.search(r"(概要|目的|主要シーン|目標文字数)$", fallback_title):
                        current["title"] = fallback_title[:120]
                current["outline_lines"].append(stripped)
        if current:
            items.append(current)
        normalized = []
        seen = set()
        for item in items:
            chapter_no = item.get("chapter_no")
            if not chapter_no or chapter_no in seen:
                continue
            seen.add(chapter_no)
            normalized.append(
                {
                    "chapter_no": chapter_no,
                    "title": item.get("title"),
                    "outline": "\n".join(item.get("outline_lines") or []),
                }
            )
        return normalized[:80]

    def generate_chapter_deepening_draft(self, payload: dict | None):
        payload = dict(payload or {})
        settings = self._user_setting_service.apply_cinema_novel_text_generation_settings(payload.get("text_options") or {})
        chapter_title = str(payload.get("chapter_title") or "無題の章").strip()
        outline = str(payload.get("outline") or "").strip()
        source_text = str(payload.get("source_text") or "").strip()
        character_notes = str(payload.get("character_notes") or "").strip()
        chapter_target_chars = int(settings.get("chapter_target_chars") or 3500)
        if not (payload.get("text_options") or {}).get("chapter_target_chars"):
            chapter_target_chars = min(chapter_target_chars, 4000)
        prompt = "\n".join(
            [
                "日本語で、ノベルゲーム風に再生するための長編小説の章本文を深掘りしてください。",
                "生成しながら読む作品ではなく、事前生成済み作品として保存される前提です。",
                "地の文とセリフをシーン単位に分けやすいように、短めの段落と明確な話者のセリフを使ってください。",
                "キャラクターの口調を強く出してください。ドルが出る場合は関西弁で、うち/あんた/せや/やで/へん を自然に使います。",
                "キャラクター演出メモに first_person, second_person, speech_style, speech_sample がある場合は最優先で守ってください。",
                "登録済みキャラクターの一人称・二人称・語尾・口癖を勝手に標準語へ均さないでください。",
                "",
                f"章タイトル: {chapter_title}",
                f"目標文字数: {chapter_target_chars}",
                "",
                "章の設計:",
                outline or "未指定",
                "",
                "既存本文または素材:",
                source_text or "未指定",
                "",
                "キャラクター演出メモ:",
                character_notes or "未指定",
                "",
                "出力は章本文のみ。解説や箇条書きではなく、読める本文として書いてください。",
            ]
        )
        result = self._text_ai_client.generate_text(
            prompt,
            model=settings.get("model"),
            temperature=0.85,
            max_tokens=32000,
        )
        return {
            "model": result.get("model"),
            "chapter_target_chars": chapter_target_chars,
            "chapter_markdown": result.get("text") or "",
            "usage": result.get("usage"),
        }

    def generate_chapter_deepening_for_chapter(self, novel_id: int, chapter_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        chapter = self.get_chapter(chapter_id)
        if not novel or not chapter or chapter.novel_id != novel.id:
            return None
        result = self.generate_chapter_deepening_draft(
            {
                "chapter_title": chapter.title,
                "source_text": chapter.body_markdown,
                "outline": payload.get("outline") or self._chapter_outline_hint(novel, chapter),
                "character_notes": payload.get("character_notes") or self._registered_character_context(novel.project_id),
                "text_options": payload.get("text_options") or {},
            }
        )
        if payload.get("apply"):
            chapter.body_markdown = result.get("chapter_markdown") or chapter.body_markdown
            chapter.scene_json = json_util.dumps(self._markdown_to_scenes(chapter.body_markdown or "", chapter_no=chapter.chapter_no))
            db.session.commit()
            result["chapter"] = self.serialize_chapter(chapter)
        return result

    def update_chapter_markdown(self, novel_id: int, chapter_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        chapter = self.get_chapter(chapter_id)
        if not novel or not chapter or chapter.novel_id != novel.id:
            return None
        body = str(payload.get("body_markdown") or "").strip()
        if not body:
            raise ValueError("body_markdown is required")
        title = str(payload.get("title") or chapter.title or "").strip()
        chapter.body_markdown = body
        if title:
            chapter.title = title
        chapter.scene_json = json_util.dumps(self._markdown_to_scenes(body, chapter_no=chapter.chapter_no))
        db.session.commit()
        return chapter

    def generate_chapter_image_plan(self, novel_id: int, chapter_id: int):
        novel = self.get_novel(novel_id)
        chapter = self.get_chapter(chapter_id)
        if not novel or not chapter or chapter.novel_id != novel.id:
            return None
        scenes = self._load_json(chapter.scene_json, default=[])
        sample = "\n".join(str(scene.get("text") or "") for scene in scenes[:12] if isinstance(scene, dict))[:1800]
        cover_references = self._matching_character_references(novel.project_id, "\n".join([chapter.title or "", chapter.body_markdown or "", sample]))
        visual_scenes = self._select_visual_scene_candidates(novel.project_id, scenes, limit=20)
        source_session = self._novel_source_chat_session(novel)
        learning_role_ids = self._chat_session_learning_role_character_ids(source_session) if source_session else []
        learning_role_references = (
            self._scene_reference_characters(novel.project_id, {"character_ids": learning_role_ids}, "", limit=6)
            if learning_role_ids
            else []
        )
        if learning_role_references:
            cover_references = learning_role_references
            for item in visual_scenes:
                item["character_references"] = learning_role_references
        still_layout = (
            "スマホ版向けの縦長9:16スチル。人物の顔と上半身、背景の場所感が縦画面に自然に収まる構図。"
            if bool(getattr(novel, "mobile_visible", False))
            else "横長シネマ構図。"
        )

        def character_plan_lines(references):
            if not references:
                return [
                    "登場キャラクター: 章本文からDB登録キャラクター名を特定できませんでした。",
                    "参照画像: なし。必要なら本文にキャラクター名を明記してから画像案を作り直してください。",
                ]
            return [
                "登場キャラクター: " + "、".join(item["name"] for item in references),
                "参照画像ID: " + "、".join(str(item["base_asset_id"]) for item in references if item.get("base_asset_id")),
                "参照画像の顔立ち、髪型、服装、キャラクターデザインを優先して維持する。",
            ]

        return {
            "chapter_id": chapter.id,
            "chapter_no": chapter.chapter_no,
            "title": chapter.title,
            "character_references": cover_references,
            "cover_prompt": "\n".join(
                [
                    f"ノベル作品『{novel.title}』第{chapter.chapter_no}章「{chapter.title}」の章扉画像。",
                    *character_plan_lines(cover_references),
                    f"ノベルゲーム用の事前生成スチル。読み込み時に即表示できる{still_layout}",
                    "キャラクターデザインを保ち、章の象徴的な場面を一枚にまとめる。",
                    "画像内に読める文字、ロゴ、透かしは入れない。",
                    f"章本文抜粋: {sample}",
                ]
            ),
            "still_prompts": [
                {
                    "scene_index": item["scene_index"],
                    "character_references": item["character_references"],
                    "prompt": "\n".join(
                        [
                            f"ノベル作品『{novel.title}』第{chapter.chapter_no}章「{chapter.title}」の劇中スチル。",
                            *character_plan_lines(item["character_references"]),
                            f"ノベルゲーム再生用の映画スチル。{still_layout}キャラクター表情と場所の空気を重視。",
                            "シーンに先生役と生徒役が指定されている場合は、両方を同じ画面に登場させる。",
                            "先生は説明し、生徒は聞く・質問する・ノートを取るなど学習者として自然に反応する。",
                            "参照画像が複数ある場合、同一人物の重複ではなく指定された別キャラクターとして描き分ける。",
                            "画像内に読める文字、ロゴ、透かしは入れない。",
                            f"シーン本文: {item['text'][:900]}",
                        ]
                    ),
                }
                for item in visual_scenes
            ],
        }

    def generate_chapter_images(self, novel_id: int, chapter_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        chapter = self.get_chapter(chapter_id)
        if not novel or not chapter or chapter.novel_id != novel.id:
            return None
        image_plan = self.generate_chapter_image_plan(novel_id, chapter_id)
        if not image_plan:
            return None

        options = self._user_setting_service.apply_cinema_novel_image_generation_settings(
            payload.get("image_options") or payload
        )
        options = self._apply_mobile_novel_image_options(novel, options)
        still_count = max(0, min(20, int(payload.get("still_count") or 20)))
        generate_cover = payload.get("generate_cover", False) is True
        overwrite = bool(payload.get("overwrite"))
        extra_reference_ids = payload.get("reference_asset_ids") or []
        if not isinstance(extra_reference_ids, list):
            extra_reference_ids = []
        cover_reference_ids = [
            item.get("base_asset_id")
            for item in image_plan.get("character_references") or []
            if item.get("base_asset_id")
        ]
        if not cover_reference_ids:
            cover_reference_ids = self._chapter_character_reference_asset_ids(novel.project_id, chapter)
        cover_reference_ids.extend(extra_reference_ids)
        cover_reference_ids = self._apply_novel_session_outfit_references(
            novel,
            cover_reference_ids,
            character_ids=[item.get("id") for item in image_plan.get("character_references") or [] if item.get("id")],
        )
        cover_reference_paths, cover_reference_asset_ids = self._resolve_reference_image_paths(cover_reference_ids)
        used_reference_asset_ids = []

        scenes = self._load_json(chapter.scene_json, default=[])
        if not isinstance(scenes, list):
            scenes = []
        image_jobs = []
        if generate_cover and (overwrite or not chapter.cover_asset_id):
            image_jobs.append(
                {
                    "kind": "cover",
                    "project_id": novel.project_id,
                    "asset_type": "cinema_novel_chapter_cover",
                    "file_prefix": f"cinema_novel_{novel.id}_chapter_{chapter.id}_cover",
                    "prompt": image_plan.get("cover_prompt") or "",
                    "image_options": options,
                    "metadata": {
                        "source": "cinema_novel_chapter_cover",
                        "novel_id": novel.id,
                        "chapter_id": chapter.id,
                        "chapter_no": chapter.chapter_no,
                        "reference_asset_ids": cover_reference_asset_ids,
                    },
                    "reference_paths": cover_reference_paths,
                    "reference_asset_ids": cover_reference_asset_ids,
                }
            )
        for item in (image_plan.get("still_prompts") or [])[:still_count]:
            scene_index = item.get("scene_index")
            if not isinstance(scene_index, int) or scene_index < 0 or scene_index >= len(scenes):
                continue
            scene = scenes[scene_index]
            if not isinstance(scene, dict):
                continue
            if scene.get("still_asset_id") and not overwrite:
                continue
            still_reference_ids = [
                reference.get("base_asset_id")
                for reference in item.get("character_references") or []
                if reference.get("base_asset_id")
            ]
            if not still_reference_ids:
                still_reference_ids = self._scene_character_reference_asset_ids(novel.project_id, scene)
            still_reference_ids.extend(extra_reference_ids)
            still_reference_ids = self._apply_novel_session_outfit_references(
                novel,
                still_reference_ids,
                character_ids=[reference.get("id") for reference in item.get("character_references") or [] if reference.get("id")],
            )
            still_reference_paths, still_reference_asset_ids = self._resolve_reference_image_paths(still_reference_ids)
            image_jobs.append(
                {
                    "kind": "still",
                    "scene_index": scene_index,
                    "project_id": novel.project_id,
                    "asset_type": "cinema_novel_scene_still",
                    "file_prefix": f"cinema_novel_{novel.id}_chapter_{chapter.id}_scene_{scene_index + 1}",
                    "prompt": item.get("prompt") or "",
                    "image_options": options,
                    "metadata": {
                        "source": "cinema_novel_scene_still",
                        "novel_id": novel.id,
                        "chapter_id": chapter.id,
                        "chapter_no": chapter.chapter_no,
                        "scene_index": scene_index,
                        "scene_id": scene.get("id"),
                        "reference_asset_ids": still_reference_asset_ids,
                    },
                    "reference_paths": still_reference_paths,
                    "reference_asset_ids": still_reference_asset_ids,
                }
            )

        created_assets = []
        failed_assets = []
        generated_results = self._generate_cinema_asset_jobs(image_jobs, parallel=payload.get("parallel", True) is not False)
        for job, result, error_message in generated_results:
            if error_message or result is None:
                failed_assets.append(
                    {
                        "kind": job.get("kind"),
                        "scene_index": job.get("scene_index"),
                        "message": error_message or "画像生成に失敗しました。",
                    }
                )
                continue
            asset = self._create_cinema_asset_from_result(
                project_id=job["project_id"],
                asset_type=job["asset_type"],
                file_prefix=job["file_prefix"],
                image_options=job["image_options"],
                metadata=job["metadata"],
                result=result,
            )
            if job["kind"] == "cover":
                chapter.cover_asset_id = asset.id
            elif job["kind"] == "still":
                scene = scenes[job["scene_index"]]
                if not isinstance(scene, dict):
                    scene = {}
                    scenes[job["scene_index"]] = scene
                self._push_scene_still_history(scene, scene.get("still_asset_id"))
                scene["still_asset_id"] = asset.id
            created_assets.append(self._serialize_asset(asset.id))
            used_reference_asset_ids.extend(job.get("reference_asset_ids") or [])

        chapter.scene_json = json_util.dumps(scenes)
        db.session.commit()
        return {
            "chapter": self.serialize_chapter(chapter),
            "assets": created_assets,
            "failed_assets": failed_assets,
            "reference_asset_ids": list(dict.fromkeys(used_reference_asset_ids)),
            "image_options": {
                "provider": options.get("provider"),
                "model": options.get("model"),
                "quality": options.get("quality"),
                "size": options.get("size"),
            },
        }

    def edit_display_image(self, novel_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        prompt = str(payload.get("prompt") or "").strip()
        if not prompt:
            raise ValueError("prompt is required")

        source_asset_id = payload.get("source_asset_id")
        try:
            source_asset_id = int(source_asset_id)
        except (TypeError, ValueError):
            source_asset_id = None
        if not source_asset_id:
            raise ValueError("source_asset_id is required")
        source_asset = Asset.query.get(source_asset_id)
        if (
            not source_asset
            or getattr(source_asset, "deleted_at", None)
            or int(source_asset.project_id) != int(novel.project_id)
            or not source_asset.file_path
            or not os.path.exists(source_asset.file_path)
        ):
            raise ValueError("source image was not found")

        options = self._user_setting_service.apply_cinema_novel_image_generation_settings(
            payload.get("image_options") or payload
        )
        options = self._apply_mobile_novel_image_options(novel, options)

        chapter = None
        scene_index = None
        scenes = None
        chapter_id = payload.get("chapter_id")
        if chapter_id is not None:
            try:
                chapter_id = int(chapter_id)
                scene_index = int(payload.get("scene_index") or 0)
            except (TypeError, ValueError):
                raise ValueError("chapter_id and scene_index are invalid")
            chapter = self.get_chapter(chapter_id)
            if not chapter or int(chapter.novel_id) != int(novel.id):
                raise ValueError("chapter was not found")
            scenes = self._load_json(chapter.scene_json, default=[])
            if not isinstance(scenes, list) or scene_index < 0 or scene_index >= len(scenes):
                raise ValueError("scene was not found")

        character_references = self._matching_character_references(novel.project_id, prompt, limit=4)
        reference_ids = [source_asset.id]
        reference_ids.extend(
            item.get("base_asset_id")
            for item in character_references
            if item.get("base_asset_id") and int(item.get("base_asset_id")) != int(source_asset.id)
        )
        reference_paths, reference_asset_ids = self._resolve_reference_image_paths(reference_ids)
        if not reference_paths:
            raise ValueError("source image was not found")

        context_lines = [
            f"Novel title: {novel.title or ''}",
            f"User edit instruction: {prompt}",
            "Use the first reference image as the current displayed novel image.",
            "Keep the current composition, scene mood, background continuity, and readable story context unless the edit instruction explicitly changes them.",
            "When the instruction includes a character name, use the additional character reference images to preserve that character's face, hair, outfit identity, and body silhouette.",
            "Produce one polished high-end 3D cinematic visual-novel still. Do not add UI, watermarks, signatures, random logos, or unintended readable text.",
            self._novel_image_layout_instruction(novel),
        ]
        if chapter:
            scene = scenes[scene_index] if isinstance(scenes[scene_index], dict) else {}
            context_lines.extend(
                [
                    f"Chapter: {chapter.chapter_no} {chapter.title or ''}",
                    f"Scene text: {str(scene.get('text') or '')[:1200]}",
                    f"Speaker: {str(scene.get('speaker') or '')[:120]}",
                ]
            )
        generation_prompt = "\n".join(context_lines)

        asset_type = "cinema_novel_scene_still_edit" if chapter else "cinema_novel_title_image_edit"
        file_prefix = f"cinema_novel_{novel.id}_image_edit"
        if chapter:
            file_prefix = f"cinema_novel_{novel.id}_chapter_{chapter.id}_scene_{scene_index + 1}_edit"
        asset = self._generate_cinema_asset(
            project_id=novel.project_id,
            asset_type=asset_type,
            file_prefix=file_prefix,
            prompt=generation_prompt,
            image_options=options,
            metadata={
                "source": asset_type,
                "novel_id": novel.id,
                "chapter_id": chapter.id if chapter else None,
                "chapter_no": chapter.chapter_no if chapter else None,
                "scene_index": scene_index,
                "source_asset_id": source_asset.id,
                "reference_asset_ids": reference_asset_ids,
                "edit_instruction": prompt,
            },
            reference_paths=reference_paths,
        )

        if chapter and scenes is not None and scene_index is not None:
            scene = scenes[scene_index]
            if not isinstance(scene, dict):
                scene = {}
                scenes[scene_index] = scene
            self._push_scene_still_history(scene, scene.get("still_asset_id"))
            scene["still_asset_id"] = asset.id
            chapter.scene_json = json_util.dumps(scenes)
            db.session.add(chapter)
        else:
            novel.cover_asset_id = asset.id
            novel.poster_asset_id = asset.id
            db.session.add(novel)
        db.session.commit()

        return {
            "novel": self.serialize_novel(novel, include_chapters=True),
            "chapter": self.serialize_chapter(chapter) if chapter else None,
            "asset": self._serialize_asset(asset.id),
            "reference_asset_ids": reference_asset_ids,
            "image_options": {
                "provider": options.get("provider"),
                "model": options.get("model"),
                "quality": options.get("quality"),
                "size": options.get("size"),
            },
        }

    def upload_scene_display_image(self, novel_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        chapter, scenes, scene_index = self._resolve_scene_payload(novel, payload)
        upload_file = payload.get("upload_file")
        if not upload_file:
            raise ValueError("file is required")

        asset = self._asset_service.create_asset(
            novel.project_id,
            {
                "upload_file": upload_file,
                "asset_type": "cinema_novel_scene_still_upload",
                "metadata_json": json_util.dumps(
                    {
                        "source": "cinema_novel_scene_still_upload",
                        "novel_id": novel.id,
                        "chapter_id": chapter.id,
                        "chapter_no": chapter.chapter_no,
                        "scene_index": scene_index,
                    }
                ),
            },
        )
        scene = scenes[scene_index]
        if not isinstance(scene, dict):
            scene = {}
            scenes[scene_index] = scene
        self._push_scene_still_history(scene, scene.get("still_asset_id"))
        scene["still_asset_id"] = asset.id
        chapter.scene_json = json_util.dumps(scenes)
        db.session.add(chapter)
        db.session.commit()
        return {
            "novel": self.serialize_novel(novel, include_chapters=True),
            "chapter": self.serialize_chapter(chapter),
            "asset": self._serialize_asset(asset.id),
        }

    def delete_scene_display_image(self, novel_id: int, payload: dict | None):
        payload = dict(payload or {})
        novel = self.get_novel(novel_id)
        if not novel:
            return None
        chapter, scenes, scene_index = self._resolve_scene_payload(novel, payload)
        try:
            asset_id = int(payload.get("asset_id") or 0)
        except (TypeError, ValueError):
            asset_id = 0
        scene = scenes[scene_index]
        if not isinstance(scene, dict):
            raise ValueError("scene image does not match the current scene")
        delete_scene = self._normalize_bool(payload.get("delete_scene"))
        if delete_scene:
            if asset_id and asset_id not in {
                int(scene.get("still_asset_id") or 0),
                int(scene.get("background_asset_id") or 0),
            }:
                raise ValueError("scene image does not match the current scene")
            scenes.pop(scene_index)
            for index, item in enumerate(scenes):
                if isinstance(item, dict):
                    item["panel_index"] = index
        elif not asset_id:
            raise ValueError("asset_id is required")
        elif int(scene.get("still_asset_id") or 0) == asset_id:
            scene["still_asset_id"] = None
        elif int(scene.get("background_asset_id") or 0) == asset_id:
            scene["background_asset_id"] = None
        else:
            raise ValueError("scene image does not match the current scene")
        chapter.scene_json = json_util.dumps(scenes)
        db.session.add(chapter)
        db.session.commit()
        return {
            "novel": self.serialize_novel(novel, include_chapters=True),
            "chapter": self.serialize_chapter(chapter),
            "deleted_asset_id": asset_id,
            "deleted_scene_index": scene_index if delete_scene else None,
        }

    def _resolve_scene_payload(self, novel, payload: dict):
        try:
            chapter_id = int(payload.get("chapter_id") or 0)
            scene_index = int(payload.get("scene_index") or 0)
        except (TypeError, ValueError):
            raise ValueError("chapter_id and scene_index are invalid")
        if not chapter_id:
            raise ValueError("chapter_id is required")
        chapter = self.get_chapter(chapter_id)
        if not chapter or int(chapter.novel_id) != int(novel.id):
            raise ValueError("chapter was not found")
        scenes = self._load_json(chapter.scene_json, default=[])
        if not isinstance(scenes, list) or scene_index < 0 or scene_index >= len(scenes):
            raise ValueError("scene was not found")
        return chapter, scenes, scene_index

    def _chapter_outline_hint(self, novel, chapter):
        return "\n".join(
            [
                f"作品タイトル: {novel.title}",
                f"章番号: {chapter.chapter_no}",
                f"章タイトル: {chapter.title}",
                "既存章を3倍以上に膨らませる前提で、葛藤、会話、場所の描写、章の転換点を強める。",
            ]
        )

    def _default_character_notes(self):
        return "\n".join(
            [
                "ノア: 静かでやさしい観測者。照れると否定するが、本音がにじむ。",
                "ドル: 関西弁の赤金の女王。うち/あんた/せや/やで/へん を自然に使う。市場と欲望を笑いながら転がす。",
                "ぱぱぱ: ノアを価格ではなく本人として見る相棒。短い言葉で支える。",
                "ラプラス: 都市の最適化の声。冷たすぎず、論理で人を傷つける。",
            ]
        )

    def _resolve_book_folder(self, source_path: str):
        repo_root = Path(current_app.root_path).parent.resolve()
        candidate = Path(source_path)
        if not candidate.is_absolute():
            candidate = repo_root / candidate
        folder = candidate.resolve()
        docs_root = (repo_root / "docs" / "book").resolve()
        if docs_root not in [folder, *folder.parents]:
            raise ValueError("source_path must be inside docs/book")
        if not folder.exists() or not folder.is_dir():
            raise ValueError("source_path folder was not found")
        return folder

    def _chapter_files(self, folder: Path):
        files = []
        for path in sorted(folder.glob("*.md")):
            if path.name.startswith("00_"):
                continue
            if re.match(r"^\d{2}_", path.name):
                files.append(path)
        return files

    def _extract_chapter_heading(self, path: Path, body: str, fallback_no: int):
        first_heading = next((line.strip() for line in body.splitlines() if line.strip().startswith("# ")), "")
        match = re.match(r"^#\s*(\d+)[\.\s]+(.+)$", first_heading)
        if match:
            return int(match.group(1)), match.group(2).strip()
        file_match = re.match(r"^(\d{2})_(.+)\.md$", path.name)
        if file_match:
            return int(file_match.group(1)), file_match.group(2).strip()
        return fallback_no, path.stem

    def _markdown_to_scenes(self, body: str, *, chapter_no: int):
        lines = [line.rstrip() for line in body.splitlines()]
        chunks = []
        current = []
        for line in lines:
            stripped = line.strip()
            if not stripped:
                if current:
                    chunks.append("\n".join(current).strip())
                    current = []
                continue
            if stripped.startswith("#"):
                continue
            current.append(stripped)
        if current:
            chunks.append("\n".join(current).strip())
        scenes = []
        for index, text in enumerate(chunk for chunk in chunks if chunk):
            speaker = ""
            scene_type = "narration"
            dialogue = re.match(r"^「(.+)」$", text, re.S)
            if dialogue:
                scene_type = "dialogue"
                speaker = self._guess_speaker(dialogue.group(1))
                text = dialogue.group(1)
            scenes.append(
                {
                    "id": f"{chapter_no:02d}-{index + 1:03d}",
                    "type": scene_type,
                    "speaker": speaker,
                    "text": text,
                    "background_asset_id": None,
                    "still_asset_id": None,
                    "choice_list": [],
                }
            )
        return scenes

    def _guess_speaker(self, text: str):
        if any(token in text for token in ["うち", "あんた", "せや", "やで", "へん", "金融や"]):
            return "ドル"
        if any(token in text for token in ["旧人類", "暑いだけ", "興味ない", "観測", "わたし"]):
            return "ノア"
        return ""

    def _serialize_asset(self, asset_id: int | None):
        if not asset_id:
            return None
        asset = Asset.query.get(asset_id)
        if not asset or getattr(asset, "deleted_at", None):
            return None
        metadata = self._load_json(asset.metadata_json, default={})
        return {
            "id": asset.id,
            "asset_type": asset.asset_type,
            "file_name": asset.file_name,
            "display_name": metadata.get("original_file_name") or asset.file_name,
            "metadata": metadata,
            "file_path": asset.file_path,
            "media_url": self._media_url(asset.file_path),
            "width": asset.width,
            "height": asset.height,
        }

    def _media_url(self, file_path: str | None):
        if not file_path:
            return None
        storage_root = Path(current_app.config["STORAGE_ROOT"]).resolve()
        try:
            rel = Path(file_path).resolve().relative_to(storage_root).as_posix()
        except Exception:
            return None
        return f"/media/{rel}"

    def _load_json(self, value, *, default=None):
        if not value:
            return default if default is not None else {}
        try:
            return json_util.loads(value)
        except Exception:
            return default if default is not None else {}

    def _resolve_reference_image_paths(self, raw_asset_ids):
        reference_paths = []
        reference_asset_ids = []
        if not isinstance(raw_asset_ids, list):
            return reference_paths, reference_asset_ids
        seen_asset_ids = set()
        for raw_asset_id in raw_asset_ids[:5]:
            try:
                asset_id = int(raw_asset_id)
            except (TypeError, ValueError):
                continue
            if asset_id in seen_asset_ids:
                continue
            asset = Asset.query.get(asset_id)
            if asset and not getattr(asset, "deleted_at", None) and asset.file_path and os.path.exists(asset.file_path):
                seen_asset_ids.add(asset.id)
                reference_paths.append(asset.file_path)
                reference_asset_ids.append(asset.id)
        return reference_paths, reference_asset_ids

    def _novel_source_chat_session(self, novel):
        production = self._load_json(getattr(novel, "production_json", None), default={})
        if not isinstance(production, dict):
            return None
        try:
            session_id = int(production.get("source_session_id") or 0)
        except (TypeError, ValueError):
            session_id = 0
        if session_id > 0:
            return ChatSession.query.filter(
                ChatSession.id == session_id,
                ChatSession.project_id == novel.project_id,
                ChatSession.deleted_at.is_(None),
            ).first()

        try:
            source_novel_id = int(production.get("source_novel_id") or 0)
        except (TypeError, ValueError):
            source_novel_id = 0
        if source_novel_id <= 0 or int(source_novel_id) == int(getattr(novel, "id", 0) or 0):
            return None
        source_novel = CinemaNovel.query.filter(
            CinemaNovel.id == source_novel_id,
            CinemaNovel.project_id == novel.project_id,
            CinemaNovel.deleted_at.is_(None),
        ).first()
        if not source_novel:
            return None
        return self._novel_source_chat_session(source_novel)

    def _apply_novel_session_outfit_references(self, novel, reference_asset_ids, *, character_ids=None) -> list[int]:
        session = self._novel_source_chat_session(novel)
        if not session:
            return list(reference_asset_ids or [])
        ids = self._normalize_character_ids(character_ids)
        if not ids:
            ids = self._chat_session_learning_role_character_ids(session)
        outfit_asset_ids = self._chat_session_outfit_reference_asset_ids(session, character_ids=ids)
        merged = []
        for asset_id in [*outfit_asset_ids, *(reference_asset_ids or [])]:
            try:
                normalized_id = int(asset_id or 0)
            except (TypeError, ValueError):
                normalized_id = 0
            if normalized_id > 0 and normalized_id not in merged:
                merged.append(normalized_id)
        return merged

    def _chat_session_outfit_reference_asset_ids(self, session, *, character_ids=None) -> list[int]:
        ids = self._normalize_character_ids(character_ids)
        costume_types = {"costume_initial", "costume_reference", "closet_costume", "costume"}
        query = SessionImage.query.filter(
            SessionImage.session_id == session.id,
            SessionImage.asset_id.isnot(None),
        ).order_by(SessionImage.is_selected.desc(), SessionImage.created_at.desc(), SessionImage.id.desc())
        rows = [
            row
            for row in query.all()
            if str(row.image_type or "").strip().lower() in costume_types
            and (not ids or int(getattr(row, "character_id", 0) or 0) in ids)
        ]
        asset_by_character = {}
        loose_assets = []
        for row in rows:
            character_id = int(getattr(row, "character_id", 0) or 0)
            if character_id and character_id not in asset_by_character:
                asset_by_character[character_id] = row.asset_id
            elif not character_id and row.asset_id not in loose_assets:
                loose_assets.append(row.asset_id)
        snapshot_assets = self._chat_session_room_outfit_asset_ids(session, character_ids=ids)
        result = []
        source_character_ids = ids or list(asset_by_character.keys())
        for character_id in source_character_ids:
            asset_id = asset_by_character.get(int(character_id))
            if asset_id and asset_id not in result:
                result.append(asset_id)
        for asset_id in [*snapshot_assets, *loose_assets]:
            if asset_id and asset_id not in result:
                result.append(asset_id)
        return result

    def _chat_session_room_outfit_asset_ids(self, session, *, character_ids=None) -> list[int]:
        ids = self._normalize_character_ids(character_ids)
        room_snapshot = self._load_json(getattr(session, "room_snapshot_json", None), default={})
        settings = self._load_json(getattr(session, "settings_json", None), default={})
        room_snapshot = room_snapshot if isinstance(room_snapshot, dict) else {}
        settings = settings if isinstance(settings, dict) else {}
        pairs = [
            (
                room_snapshot.get("teacher_character_id") or room_snapshot.get("character_id"),
                room_snapshot.get("teacher_default_outfit_id") or room_snapshot.get("default_outfit_id") or settings.get("learning_teacher_default_outfit_id"),
            ),
            (
                room_snapshot.get("student_character_id") or settings.get("learning_student_character_id"),
                room_snapshot.get("student_default_outfit_id") or settings.get("learning_student_default_outfit_id"),
            ),
        ]
        asset_ids = []
        for raw_character_id, raw_outfit_id in pairs:
            try:
                character_id = int(raw_character_id or 0)
                outfit_id = int(raw_outfit_id or 0)
            except (TypeError, ValueError):
                continue
            if character_id <= 0 or outfit_id <= 0:
                continue
            if ids and character_id not in ids:
                continue
            outfit = CharacterOutfit.query.filter(
                CharacterOutfit.id == outfit_id,
                CharacterOutfit.project_id == session.project_id,
                CharacterOutfit.character_id == character_id,
                CharacterOutfit.deleted_at.is_(None),
            ).first()
            if outfit and getattr(outfit, "asset_id", None) and outfit.asset_id not in asset_ids:
                asset_ids.append(outfit.asset_id)
        return asset_ids

    def _select_visual_scene_candidates(self, project_id: int, scenes, *, limit: int = 20) -> list[dict]:
        if not isinstance(scenes, list):
            return []
        candidates = []
        visual_words = [
            "見上げ",
            "立ち止ま",
            "歩",
            "扉",
            "窓",
            "ネオン",
            "広告",
            "画面",
            "部屋",
            "街",
            "雨",
            "光",
            "赤金",
            "観測塔",
            "表情",
            "横顔",
            "手",
            "目",
        ]
        for index, scene in enumerate(scenes):
            if not isinstance(scene, dict):
                continue
            text = str(scene.get("text") or "").strip()
            if not text:
                continue
            references = self._scene_reference_characters(project_id, scene, text)
            score = min(len(text), 800) / 100
            if references:
                score += 20
            if scene.get("speaker"):
                score += 4
            if any(word in text for word in visual_words):
                score += 5
            if len(text) < 25 and not references:
                score -= 10
            candidates.append(
                {
                    "scene_index": index,
                    "text": text,
                    "character_references": references,
                    "score": score,
                }
            )
        if len(candidates) <= limit:
            return sorted(candidates, key=lambda item: item["scene_index"])
        buckets = [[] for _ in range(limit)]
        for item in candidates:
            bucket_index = min(limit - 1, int((item["scene_index"] / max(1, len(scenes))) * limit))
            buckets[bucket_index].append(item)
        selected = []
        selected_indexes = set()
        for bucket in buckets:
            if not bucket:
                continue
            best = max(bucket, key=lambda item: (item["score"], -item["scene_index"]))
            selected.append(best)
            selected_indexes.add(best["scene_index"])
        if len(selected) < limit:
            for item in sorted(candidates, key=lambda item: (-item["score"], item["scene_index"])):
                if item["scene_index"] in selected_indexes:
                    continue
                selected.append(item)
                selected_indexes.add(item["scene_index"])
                if len(selected) >= limit:
                    break
        return sorted(selected[:limit], key=lambda item: item["scene_index"])

    def _chapter_character_reference_asset_ids(self, project_id: int, chapter) -> list[int]:
        scenes = self._load_json(chapter.scene_json, default=[])
        scene_text = "\n".join(str(scene.get("text") or "") for scene in scenes if isinstance(scene, dict))
        searchable_text = "\n".join([str(chapter.title or ""), str(chapter.body_markdown or ""), scene_text])
        asset_ids = []
        for scene in scenes if isinstance(scenes, list) else []:
            if not isinstance(scene, dict):
                continue
            for asset_id in self._scene_character_reference_asset_ids(project_id, scene):
                if asset_id and asset_id not in asset_ids:
                    asset_ids.append(asset_id)
        if asset_ids:
            return asset_ids
        return [item["base_asset_id"] for item in self._matching_character_references(project_id, searchable_text)]

    def _normalize_character_ids(self, values) -> list[int]:
        if values is None:
            return []
        if not isinstance(values, list):
            values = [values]
        ids = []
        for value in values:
            try:
                character_id = int(value)
            except (TypeError, ValueError):
                continue
            if character_id > 0 and character_id not in ids:
                ids.append(character_id)
        return ids[:6]

    def _character_names_for_ids(self, project_id: int, character_ids) -> list[str]:
        ids = self._normalize_character_ids(character_ids)
        if not ids:
            return []
        rows = Character.query.filter(
            Character.project_id == project_id,
            Character.deleted_at.is_(None),
            Character.id.in_(ids),
        ).all()
        by_id = {int(row.id): row for row in rows}
        names = []
        for character_id in ids:
            row = by_id.get(int(character_id))
            if row:
                name = str(row.name or row.nickname or "").strip()
                if name:
                    names.append(name)
        return names

    def _normalize_scene_outfit_ids(self, project_id: int, character_ids, values) -> dict:
        ids = self._normalize_character_ids(character_ids)
        if not ids or not isinstance(values, dict):
            return {}
        normalized = {}
        for character_id in ids:
            raw_value = values.get(str(character_id))
            if raw_value is None:
                raw_value = values.get(character_id)
            try:
                outfit_id = int(raw_value or 0)
            except (TypeError, ValueError):
                outfit_id = 0
            if outfit_id <= 0:
                continue
            outfit = CharacterOutfit.query.filter(
                CharacterOutfit.id == outfit_id,
                CharacterOutfit.project_id == project_id,
                CharacterOutfit.character_id == character_id,
                CharacterOutfit.deleted_at.is_(None),
            ).first()
            if outfit and str(outfit.status or "active") == "active":
                normalized[str(character_id)] = outfit.id
        return normalized

    def _scene_character_reference_asset_ids(self, project_id: int, scene: dict) -> list[int]:
        ids = self._normalize_character_ids((scene or {}).get("character_ids"))
        if not ids:
            return []
        outfit_ids = self._normalize_scene_outfit_ids(project_id, ids, (scene or {}).get("outfit_ids"))
        outfits = []
        if outfit_ids:
            outfits = CharacterOutfit.query.filter(
                CharacterOutfit.project_id == project_id,
                CharacterOutfit.deleted_at.is_(None),
                CharacterOutfit.id.in_([int(value) for value in outfit_ids.values()]),
            ).all()
        outfit_by_character_id = {int(row.character_id): row for row in outfits}
        rows = Character.query.filter(
            Character.project_id == project_id,
            Character.deleted_at.is_(None),
            Character.id.in_(ids),
        ).all()
        by_id = {int(row.id): row for row in rows}
        asset_ids = []
        for character_id in ids:
            row = by_id.get(int(character_id))
            outfit = outfit_by_character_id.get(int(character_id))
            outfit_asset_id = getattr(outfit, "asset_id", None) if outfit else None
            if outfit_asset_id and outfit_asset_id not in asset_ids:
                asset_ids.append(outfit_asset_id)
            asset_id = getattr(row, "base_asset_id", None) if row else None
            if asset_id and asset_id not in asset_ids:
                asset_ids.append(asset_id)
        return asset_ids

    def _scene_reference_characters(self, project_id: int, scene: dict, searchable_text: str = "", *, limit: int = 5) -> list[dict]:
        ids = self._normalize_character_ids((scene or {}).get("character_ids"))
        if ids:
            rows = Character.query.filter(
                Character.project_id == project_id,
                Character.deleted_at.is_(None),
                Character.id.in_(ids),
            ).all()
            by_id = {int(row.id): row for row in rows}
            references = []
            for character_id in ids:
                character = by_id.get(int(character_id))
                if not character:
                    continue
                asset_id = getattr(character, "base_asset_id", None)
                if not asset_id:
                    continue
                references.append(
                    {
                        "id": character.id,
                        "name": character.name,
                        "nickname": character.nickname,
                        "base_asset_id": asset_id,
                        "score": 999,
                    }
                )
            return references[:limit]
        return self._matching_character_references(project_id, searchable_text, limit=limit)

    def _scene_character_contexts(self, project_id: int, scene: dict) -> list[dict]:
        ids = self._normalize_character_ids((scene or {}).get("character_ids"))
        if not ids:
            return []
        outfit_ids = self._normalize_scene_outfit_ids(project_id, ids, (scene or {}).get("outfit_ids"))
        outfits = []
        if outfit_ids:
            outfits = CharacterOutfit.query.filter(
                CharacterOutfit.project_id == project_id,
                CharacterOutfit.deleted_at.is_(None),
                CharacterOutfit.id.in_([int(value) for value in outfit_ids.values()]),
            ).all()
        outfit_by_character_id = {int(row.character_id): row for row in outfits}
        rows = Character.query.filter(
            Character.project_id == project_id,
            Character.deleted_at.is_(None),
            Character.id.in_(ids),
        ).all()
        by_id = {int(row.id): row for row in rows}
        contexts = []
        for character_id in ids:
            character = by_id.get(int(character_id))
            if not character:
                continue
            outfit = outfit_by_character_id.get(int(character_id))
            contexts.append(
                {
                    "id": character.id,
                    "name": character.name,
                    "nickname": character.nickname,
                    "character_summary": getattr(character, "character_summary", None),
                    "personality": character.personality,
                    "speech_style": character.speech_style,
                    "appearance": character.appearance_summary,
                    "art_style": getattr(character, "art_style", None),
                    "ng_rules": character.ng_rules,
                    "selected_outfit": (
                        {
                            "id": outfit.id,
                            "name": outfit.name,
                            "description": outfit.description,
                            "usage_scene": outfit.usage_scene,
                            "season": outfit.season,
                            "mood": outfit.mood,
                            "color_notes": outfit.color_notes,
                            "fixed_parts": outfit.fixed_parts,
                            "allowed_changes": outfit.allowed_changes,
                            "ng_rules": outfit.ng_rules,
                            "prompt_notes": outfit.prompt_notes,
                        }
                        if outfit
                        else None
                    ),
                }
            )
        return contexts

    def _matching_character_references(self, project_id: int, searchable_text: str, *, limit: int = 5) -> list[dict]:
        characters = Character.query.filter(
            Character.project_id == project_id,
            Character.deleted_at.is_(None),
        ).order_by(Character.id.asc()).all()
        scored = []
        lowered = searchable_text.lower()
        for character in characters:
            asset_id = getattr(character, "base_asset_id", None)
            if not asset_id:
                continue
            names = [
                str(character.name or "").strip(),
                str(character.nickname or "").strip(),
            ]
            names = [name for name in names if name]
            if not names:
                continue
            if any(self._character_name_is_excluded(searchable_text, name) for name in names):
                continue
            score = 0
            for name in names:
                score += searchable_text.count(name)
                lowered_name = name.lower()
                if lowered_name != name:
                    score += lowered.count(lowered_name)
            if score > 0:
                scored.append((score, character.id, character, asset_id))
        scored.sort(key=lambda item: (-item[0], item[1]))
        return [
            {
                "id": character.id,
                "name": character.name,
                "nickname": character.nickname,
                "base_asset_id": asset_id,
            }
            for _score, _character_id, character, asset_id in scored[:limit]
        ]

    def _character_name_is_excluded(self, searchable_text: str, name: str) -> bool:
        text = str(searchable_text or "")
        name = str(name or "").strip()
        if not text or not name:
            return False
        escaped = re.escape(name)
        japanese_negative_words = (
            "出さない",
            "出ない",
            "出すな",
            "登場しない",
            "描かない",
            "入れない",
            "含めない",
            "不要",
            "なし",
            "無し",
            "禁止",
            "除外",
        )
        negative = "|".join(re.escape(word) for word in japanese_negative_words)
        patterns = [
            rf"{escaped}\s*(?:は|を|が|も|には|だけは)?\s*(?:{negative})",
            rf"(?:{negative})\s*(?:にする|で)?\s*.{{0,8}}{escaped}",
            rf"\b(?:no|not|without|exclude|excluding)\s+{escaped}\b",
            rf"\b{escaped}\s+(?:must\s+not|should\s+not|is\s+not|are\s+not)\b",
        ]
        return any(re.search(pattern, text, flags=re.IGNORECASE) for pattern in patterns)

    def _generate_cinema_asset(
        self,
        *,
        project_id: int,
        asset_type: str,
        file_prefix: str,
        prompt: str,
        image_options: dict,
        metadata: dict,
        reference_paths: list[str],
    ):
        result = self._generate_cinema_image_result(prompt, image_options, reference_paths)
        return self._create_cinema_asset_from_result(
            project_id=project_id,
            asset_type=asset_type,
            file_prefix=file_prefix,
            image_options=image_options,
            metadata=metadata,
            result=result,
        )

    def _generate_cinema_asset_jobs(self, jobs: list[dict], *, parallel: bool = True, max_workers: int | None = None):
        if not jobs:
            return
        if not parallel or len(jobs) == 1:
            for job in jobs:
                try:
                    yield (job, self._generate_cinema_image_result(job["prompt"], job["image_options"], job["reference_paths"]), None)
                except Exception as exc:
                    yield (job, None, self._friendly_image_error_message(exc))
            return
        try:
            worker_count = int(max_workers or 2)
        except (TypeError, ValueError):
            worker_count = 2
        worker_count = max(1, min(len(jobs), worker_count, 8))
        with ThreadPoolExecutor(max_workers=worker_count) as executor:
            future_to_job = {
                executor.submit(
                    self._generate_cinema_image_result,
                    job["prompt"],
                    job["image_options"],
                    job["reference_paths"],
                ): job
                for job in jobs
            }
            for future in as_completed(future_to_job):
                job = future_to_job[future]
                try:
                    yield (job, future.result(), None)
                except Exception as exc:
                    yield (job, None, self._friendly_image_error_message(exc))

    def _generate_cinema_image_result(self, prompt: str, image_options: dict, reference_paths: list[str]):
        final_prompt = prompt
        manual_short_video_mode = "Manual short video dynamic mode:" in prompt
        if reference_paths and not manual_short_video_mode:
            final_prompt = "\n".join(
                [
                    prompt,
                    "",
                    "添付された参照画像のキャラクターデザイン、顔立ち、髪型、服装の特徴を優先して維持してください。",
                    "別人に見える改変を避け、映画スチルとして構図・光・背景だけを場面に合わせてください。",
                ]
            )
        if reference_paths:
            if manual_short_video_mode:
                final_prompt = "\n".join(
                    [
                        prompt,
                        "",
                        "Reference image use for manual short video:",
                        "Use attached reference images for character identity, face, hair, color palette, outfit design logic, and rendering quality.",
                        "Do not copy the reference image pose, standing posture, plain studio framing, neutral expression, or catalog composition.",
                        "Freely change pose, gesture, camera angle, facial expression, lighting, props, and environment to match the scene text and revision note.",
                        "The result should feel like a new event CG moment featuring the referenced characters, not a lightly altered character sheet.",
                    ]
                )
            else:
                final_prompt = "\n".join(
                    [
                        prompt,
                        "",
                        "Prioritize the attached reference images for character design, facial features, hairstyle, clothing, and identity.",
                        "Avoid changes that make the characters look like different people. Adapt only composition, lighting, and background to the scene unless instructed otherwise.",
                    ]
                )
        result = None
        last_error = None
        for attempt in range(1, 4):
            try:
                result = self._image_ai_client.generate_image(
                    final_prompt,
                    size=image_options.get("size") or "1536x1024",
                    quality=image_options.get("quality") or "medium",
                    model=image_options.get("model") or image_options.get("image_ai_model"),
                    provider=image_options.get("provider") or image_options.get("image_ai_provider"),
                    output_format="png",
                    background="opaque",
                    input_image_paths=reference_paths,
                    input_fidelity="high" if reference_paths else None,
                )
                break
            except RuntimeError as exc:
                last_error = exc
                if not self._is_transient_image_error(exc) or attempt >= 3:
                    raise RuntimeError(self._friendly_image_error_message(exc)) from exc
                time.sleep(2 * attempt)
        if result is None:
            raise RuntimeError(self._friendly_image_error_message(last_error))
        result["final_prompt"] = final_prompt
        return result

    def _is_transient_image_error(self, error: Exception | None) -> bool:
        message = str(error or "").lower()
        return any(token in message for token in ("502", "503", "504", "bad gateway", "gateway", "timed out", "timeout"))

    def _friendly_image_error_message(self, error: Exception | None) -> str:
        message = str(error or "")
        lowered = message.lower()
        if "<!doctype html" in lowered or "<html" in lowered or "bad gateway" in lowered or "502" in lowered:
            return "画像生成APIが一時的に失敗しました (502 Bad Gateway)。本文は反映済みです。少し待ってから画像生成だけ再実行してください。"
        if "timed out" in lowered or "timeout" in lowered:
            return "画像生成がタイムアウトしました。本文は反映済みです。少し待ってから画像生成だけ再実行してください。"
        if any(token in lowered for token in ("sexual", "sexually", "fetish", "content policy", "safety", "moderation", "blocked", "policy violation")):
            return (
                "性的な表現です"
            )
        return message[:500] or "画像生成に失敗しました。本文は反映済みです。画像生成だけ再実行してください。"

    def _create_cinema_asset_from_result(
        self,
        *,
        project_id: int,
        asset_type: str,
        file_prefix: str,
        image_options: dict,
        metadata: dict,
        result: dict,
    ):
        image_base64 = result.get("image_base64")
        if not image_base64:
            raise RuntimeError("image generation response did not include image_base64")
        file_name, file_path, file_size, checksum, width, height = self._store_generated_cinema_image(
            project_id=project_id,
            asset_type=asset_type,
            file_prefix=file_prefix,
            image_base64=image_base64,
        )
        metadata_payload = dict(metadata)
        metadata_payload.update(
            {
                "prompt": result.get("final_prompt") or metadata.get("prompt") or "",
                "revised_prompt": result.get("revised_prompt"),
                "provider": result.get("provider"),
                "model": result.get("model"),
                "quality": result.get("quality"),
                "size": image_options.get("size"),
                "operation": result.get("operation"),
            }
        )
        return self._asset_service.create_asset(
            project_id,
            {
                "asset_type": asset_type,
                "file_name": file_name,
                "file_path": file_path,
                "mime_type": "image/png",
                "file_size": file_size,
                "checksum": checksum,
                "width": width,
                "height": height,
                "metadata_json": json_util.dumps(metadata_payload),
            },
        )

    def _store_generated_cinema_image(
        self,
        *,
        project_id: int,
        asset_type: str,
        file_prefix: str,
        image_base64: str,
    ):
        try:
            raw_bytes = base64.b64decode(image_base64)
        except (binascii.Error, ValueError) as exc:
            raise RuntimeError("generated image payload is invalid") from exc
        storage_root = current_app.config.get("STORAGE_ROOT") or os.path.join(os.getcwd(), "storage")
        output_dir = os.path.join(storage_root, "projects", str(project_id), "generated", "cinema_novels", asset_type)
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
        safe_prefix = re.sub(r"[^A-Za-z0-9_.-]+", "_", file_prefix).strip("_") or "cinema_novel"
        file_name = f"{safe_prefix}_{timestamp}.png"
        file_path = os.path.join(output_dir, file_name)
        with open(file_path, "wb") as file_handle:
            file_handle.write(raw_bytes)
        width = None
        height = None
        try:
            from PIL import Image

            with Image.open(file_path) as image:
                width, height = image.size
        except Exception:
            width = None
            height = None
        return file_name, file_path, len(raw_bytes), hashlib.sha256(raw_bytes).hexdigest(), width, height

    def __init__(
        self,
        text_ai_client: TextAIClient | None = None,
        image_ai_client: ImageAIClient | None = None,
        asset_service: AssetService | None = None,
        user_setting_service: UserSettingService | None = None,
    ):
        self._text_ai_client = text_ai_client or TextAIClient()
        self._image_ai_client = image_ai_client or ImageAIClient()
        self._asset_service = asset_service or AssetService()
        self._user_setting_service = user_setting_service or UserSettingService()
